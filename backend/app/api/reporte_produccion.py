"""
Reporte de producción y ventas de Carnes y Mariscos — router. Captura diaria
de los formatos Porcionador y Cocina, catálogo editable de renglones, cuadre y
cruce porcionador → cocina. Ver `services/reporte_produccion.py`.
"""

from dataclasses import dataclass
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import get_current_user
from app.db.session import get_db
from app.models import (
    ReporteProduccionCaptura,
    ReporteProduccionDia,
    ReporteProduccionPesaje,
    ReporteProduccionProductoBase,
    ReporteProduccionRenglon,
)
from app.services.reporte_produccion import (
    HOJAS,
    PRECIOS_KG_SEMILLA,
    calcular_cadena_renglon,
    calcular_control_limpieza,
    calcular_inventario_bruto,
    completar_rendimientos,
    construir_clave_cruce,
    construir_entradas_automaticas,
    estado_diferencia,
    factores_por_producto,
    filas_semilla,
    gramaje_efectivo_kg,
    merma_kg,
    renglones_sin_contraparte,
    resolver_crudos,
    salidas_auto_crudo,
)

router = APIRouter()


def _validar_hoja(hoja: str) -> str:
    if hoja not in HOJAS:
        raise HTTPException(status_code=400, detail=f"hoja debe ser una de {list(HOJAS)}")
    return hoja


def _renglon_a_dict(r: ReporteProduccionRenglon) -> dict:
    return {
        "id": r.id, "hoja": r.hoja, "seccion": r.seccion, "producto": r.producto,
        "presentacion": r.presentacion, "gramaje": r.gramaje, "unidad": r.unidad,
        "orden": r.orden, "activo": r.activo, "clave_cruce": r.clave_cruce,
        "gramaje_kg": r.gramaje_kg, "prueba_bruto_kg": r.prueba_bruto_kg, "prueba_limpio_kg": r.prueba_limpio_kg,
        "es_crudo": r.es_crudo,
        "es_limpio": r.es_limpio,
        # el número capturado o, si falta, el que dice el texto impreso ("180g" → 0.18); es el que usa el cálculo
        "gramaje_kg_efectivo": gramaje_efectivo_kg(r),
    }


# ── Catálogo de renglones ──────────────────────────────────────────────


class RenglonIn(BaseModel):
    hoja: str
    seccion: str
    producto: str
    presentacion: str | None = None
    gramaje: str | None = None
    unidad: str = "PORCION"
    orden: int | None = None
    clave_cruce: str | None = None
    gramaje_kg: float | None = None
    prueba_bruto_kg: float | None = None
    prueba_limpio_kg: float | None = None
    es_crudo: bool | None = None
    es_limpio: bool | None = None


class RenglonPatch(BaseModel):
    seccion: str | None = None
    producto: str | None = None
    presentacion: str | None = None
    gramaje: str | None = None
    unidad: str | None = None
    orden: int | None = None
    activo: bool | None = None
    clave_cruce: str | None = None
    gramaje_kg: float | None = None
    prueba_bruto_kg: float | None = None
    prueba_limpio_kg: float | None = None
    es_crudo: bool | None = None
    es_limpio: bool | None = None


@router.get("/catalogo")
async def listar_catalogo(
    hoja: str | None = Query(None),
    incluir_inactivos: bool = Query(False),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    q = select(ReporteProduccionRenglon).order_by(ReporteProduccionRenglon.hoja, ReporteProduccionRenglon.orden)
    if hoja is not None:
        q = q.where(ReporteProduccionRenglon.hoja == _validar_hoja(hoja))
    if not incluir_inactivos:
        q = q.where(ReporteProduccionRenglon.activo.is_(True))
    return [_renglon_a_dict(r) for r in (await db.execute(q)).scalars().all()]


@router.post("/catalogo/sembrar")
async def sembrar_catalogo(
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Inserta los renglones del formato original en las hojas que aún no tienen ninguno y rellena (solo vacíos)
    gramaje, pruebas de rendimiento y precios por kg tomados de la hoja PORCIONADOR. Idempotente: nunca pisa
    ni vuelve a crear lo que el usuario editó."""
    # Los renglones solo se siembran en una hoja VACÍA. El catálogo es editable: si el usuario renombró, movió o
    # desactivó renglones, buscar "los que faltan" por nombre los volvía a crear duplicados.
    hojas_con_renglones = {
        h for h in (await db.execute(select(ReporteProduccionRenglon.hoja).distinct())).scalars().all()
    }
    creados = 0
    for hoja in HOJAS:
        if hoja in hojas_con_renglones:
            continue
        for fila in filas_semilla(hoja):
            db.add(ReporteProduccionRenglon(**fila))
            creados += 1
    await db.flush()

    todos = list((await db.execute(select(ReporteProduccionRenglon))).scalars().all())
    rendimientos = completar_rendimientos(todos)
    precios_actuales = {
        p.producto: p for p in (await db.execute(select(ReporteProduccionProductoBase))).scalars().all()
    }
    precios = 0
    for producto in dict.fromkeys(r.producto for r in todos):
        semilla = PRECIOS_KG_SEMILLA.get(producto)
        if semilla is None:
            continue
        existente = precios_actuales.get(producto)
        if existente is None:
            db.add(ReporteProduccionProductoBase(producto=producto, precio_kg=semilla))
            precios += 1
        elif existente.precio_kg is None:
            existente.precio_kg = semilla
            precios += 1
    await db.flush()
    return {"creados": creados, "rendimientos": rendimientos, "precios": precios}


@router.post("/catalogo")
async def crear_renglon(
    datos: RenglonIn,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    _validar_hoja(datos.hoja)
    orden = datos.orden
    if orden is None:
        ultimos = (
            await db.execute(select(ReporteProduccionRenglon.orden).where(ReporteProduccionRenglon.hoja == datos.hoja))
        ).scalars().all()
        orden = (max(ultimos) if ultimos else 0) + 10
    renglon = ReporteProduccionRenglon(
        hoja=datos.hoja, seccion=datos.seccion.strip(), producto=datos.producto.strip(),
        presentacion=(datos.presentacion or "").strip() or None,
        gramaje=(datos.gramaje or "").strip() or None, unidad=datos.unidad, orden=orden, activo=True,
        clave_cruce=datos.clave_cruce or construir_clave_cruce(datos.producto, datos.presentacion),
        gramaje_kg=datos.gramaje_kg, prueba_bruto_kg=datos.prueba_bruto_kg, prueba_limpio_kg=datos.prueba_limpio_kg,
        es_crudo=datos.es_crudo, es_limpio=datos.es_limpio,
    )
    db.add(renglon)
    await db.flush()
    return _renglon_a_dict(renglon)


@router.patch("/catalogo/{renglon_id}")
async def editar_renglon(
    renglon_id: int,
    datos: RenglonPatch,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    renglon = await db.get(ReporteProduccionRenglon, renglon_id)
    if renglon is None:
        raise HTTPException(status_code=404, detail="Renglón no encontrado")
    cambios = datos.model_dump(exclude_unset=True)
    for campo, valor in cambios.items():
        if isinstance(valor, str):
            valor = valor.strip() or None
        setattr(renglon, campo, valor)
    # si cambió producto/presentación y no se dio clave explícita, se recalcula
    if ("producto" in cambios or "presentacion" in cambios) and "clave_cruce" not in cambios:
        renglon.clave_cruce = construir_clave_cruce(renglon.producto, renglon.presentacion)
    await db.flush()
    return _renglon_a_dict(renglon)


@router.delete("/catalogo/{renglon_id}")
async def desactivar_renglon(
    renglon_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Baja lógica: conserva el historial de capturas del renglón."""
    renglon = await db.get(ReporteProduccionRenglon, renglon_id)
    if renglon is None:
        raise HTTPException(status_code=404, detail="Renglón no encontrado")
    renglon.activo = False
    await db.flush()
    return {"ok": True}


# ── Captura diaria ─────────────────────────────────────────────────────
#
# La existencia final no se captura: se calcula en cadena por renglón (ver
# `services/reporte_produccion.py::calcular_cadena_renglon`). Se captura
# entradas y salidas cada día. `inicial` y la existencia final (`conteo_fisico`)
# son editables cualquier día para corregir fallas u olvidos: solo se guardan
# cuando se capturan a mano; si no, mandan los valores calculados en cadena.


class CapturaRenglonIn(BaseModel):
    renglon_id: int
    inicial: float | None = None
    entradas: float | None = None
    salidas: float | None = None
    conteo_fisico: float | None = None


class CapturaIn(BaseModel):
    fecha: date
    renglones: list[CapturaRenglonIn]
    # Al guardar el día, registra TODOS los productos de esa hoja (los que no se tocaron quedan sin movimiento:
    # su inicial y su existencia final salen de la cadena). Sin esto solo se guardan los renglones enviados.
    completar_hoja: str | None = None


class DiaIn(BaseModel):
    fecha: date
    responsable_porcionador: str | None = None
    responsable_cocina: str | None = None
    notas: str | None = None


@dataclass
class _CapturaCadena:
    fecha: date
    inicial: float | None
    entradas: float | None
    salidas: float | None
    conteo_fisico: float | None


async def _renglones_activos(db: AsyncSession) -> list[ReporteProduccionRenglon]:
    return list(
        (
            await db.execute(
                select(ReporteProduccionRenglon)
                .where(ReporteProduccionRenglon.activo.is_(True))
                .order_by(ReporteProduccionRenglon.hoja, ReporteProduccionRenglon.orden)
            )
        ).scalars().all()
    )


async def _calcular_dia(db: AsyncSession, fecha: date, hoja: str) -> tuple[list[dict], date | None]:
    """Filas de la hoja para `fecha` con la cadena de saldos resuelta (y fecha del último día previo con captura)."""
    todos = await _renglones_activos(db)
    ids = [r.id for r in todos]
    capturas = (
        list(
            (
                await db.execute(
                    select(ReporteProduccionCaptura).where(
                        ReporteProduccionCaptura.fecha <= fecha, ReporteProduccionCaptura.renglon_id.in_(ids)
                    )
                )
            ).scalars().all()
        )
        if ids
        else []
    )
    entradas_auto = construir_entradas_automaticas(todos, capturas)
    fechas_con_datos = {c.fecha for c in capturas}
    # crudo que se limpia: su salida es el bruto pesado ese día ("usado en limpieza"); lo capturado a mano manda
    pesajes_hasta = list(
        (await db.execute(select(ReporteProduccionPesaje).where(ReporteProduccionPesaje.fecha <= fecha))).scalars().all()
    )
    crudos = resolver_crudos(todos)
    salidas_auto = salidas_auto_crudo([r for r in todos if r.hoja == "porcionador"], pesajes_hasta)
    por_renglon: dict[int, dict[date, ReporteProduccionCaptura]] = {}
    for c in capturas:
        por_renglon.setdefault(c.renglon_id, {})[c.fecha] = c

    filas: list[dict] = []
    fecha_previa: date | None = None
    for r in (r for r in todos if r.hoja == hoja):
        guardadas = por_renglon.get(r.id, {})
        dias = set(guardadas)
        # cocina: un día con salidas del porcionador aporta entradas aunque cocina aún no lo haya guardado
        if hoja == "cocina" and entradas_auto(r, fecha) is not None:
            dias |= {d for d in fechas_con_datos if (entradas_auto(r, d) or 0) > 0}
        # crudo: un día con pesaje descuenta aunque ese renglón aún no se haya guardado
        if hoja == "porcionador" and r.id in crudos:
            dias |= {d for (rid, d) in salidas_auto if rid == r.id}

        vistas: list[_CapturaCadena] = []
        for d in dias:
            c = guardadas.get(d)
            auto = entradas_auto(r, d)
            vistas.append(
                _CapturaCadena(
                    fecha=d,
                    inicial=c.inicial if c else None,
                    # lo capturado a mano manda; si no, lo que el porcionador mandó a cocina
                    entradas=c.entradas if (c is not None and c.entradas is not None) else auto,
                    salidas=c.salidas if (c is not None and c.salidas is not None) else salidas_auto.get((r.id, d)),
                    conteo_fisico=c.conteo_fisico if c else None,
                )
            )
        cadena = calcular_cadena_renglon(vistas)
        del_dia = next((f for f in cadena if f["fecha"] == fecha), None)
        previas = [f for f in cadena if f["fecha"] < fecha]
        if previas and (fecha_previa is None or previas[-1]["fecha"] > fecha_previa):
            fecha_previa = previas[-1]["fecha"]

        auto_hoy = entradas_auto(r, fecha)
        base = {
            **_renglon_a_dict(r),
            "entradas_es_auto": auto_hoy is not None,
            "entradas_auto": auto_hoy,
            "entradas_manual": guardadas[fecha].entradas if fecha in guardadas else None,
            "capturado": fecha in guardadas,
            # crudo: sugerencia de salida = bruto pesado ese día (la columna "salidas" sigue siendo lo capturado a mano)
            "salidas_es_auto": hoja == "porcionador" and r.id in crudos,
            "salidas_auto": salidas_auto.get((r.id, fecha)) if (hoja == "porcionador" and r.id in crudos) else None,
        }
        if del_dia is not None:
            filas.append({
                **base,
                "inicial": del_dia["inicial"],
                "inicial_manual": del_dia["inicial_manual"],
                "inicial_arrastrado": del_dia["inicial_arrastrado"],
                "es_apertura": del_dia["es_apertura"],
                "entradas": del_dia["entradas"] if (fecha in guardadas or auto_hoy is not None) else None,
                "salidas": guardadas[fecha].salidas if fecha in guardadas else None,
                "conteo_fisico": del_dia["conteo_fisico"],
                "final_calculado": del_dia["final_calculado"],
                "existencia_final": del_dia["existencia_final"],
                "diferencia": del_dia["diferencia"],
                "estado_diferencia": estado_diferencia(del_dia["diferencia"]),
                "alerta_negativo": del_dia["alerta_negativo"],
            })
        else:
            # día aún sin captura de este renglón: se muestra el saldo que arrastra
            filas.append({
                **base,
                "inicial": previas[-1]["existencia_final"] if previas else None,
                "inicial_manual": None,
                "inicial_arrastrado": previas[-1]["existencia_final"] if previas else None,
                "es_apertura": not previas,
                "entradas": auto_hoy,
                "salidas": None,
                "conteo_fisico": None,
                "final_calculado": None,
                "existencia_final": None,
                "diferencia": None,
                "estado_diferencia": "sin_conteo",
                "alerta_negativo": False,
            })
    return filas, fecha_previa


def _dia_a_dict(dia: ReporteProduccionDia | None) -> dict:
    return {
        "responsable_porcionador": dia.responsable_porcionador if dia else None,
        "responsable_cocina": dia.responsable_cocina if dia else None,
        "notas": dia.notas if dia else None,
    }


@router.get("/fechas")
async def listar_fechas(
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Días con alguna captura, con cuántos renglones se capturaron."""
    filas = (await db.execute(select(ReporteProduccionCaptura))).scalars().all()
    por_fecha: dict[date, int] = {}
    for c in filas:
        por_fecha[c.fecha] = por_fecha.get(c.fecha, 0) + 1
    return [{"fecha": f.isoformat(), "total": n} for f, n in sorted(por_fecha.items(), reverse=True)]


@router.get("/cruce")
async def obtener_cruce(
    fecha: date = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Renglones sin equivalente en la otra hoja (los de cocina capturan sus entradas a mano)."""
    return {"fecha": fecha.isoformat(), "sin_contraparte": renglones_sin_contraparte(await _renglones_activos(db))}


@router.get("/export")
async def exportar_hoja(
    fecha: date = Query(...),
    hoja: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    _validar_hoja(hoja)
    filas, _ = await _calcular_dia(db, fecha, hoja)
    dia = (await db.execute(select(ReporteProduccionDia).where(ReporteProduccionDia.fecha == fecha))).scalar_one_or_none()

    es_porcionador = hoja == "porcionador"
    titulo = "ÁREA DE PRODUCCIÓN" if es_porcionador else "ÁREA DE COCINA"
    cols = ["PRODUCTO", "PRESENTACIÓN / VARIANTE", "UNIDAD", "GRAMAJE"]
    cols += (
        ["INVENTARIO INICIAL", "ENTRADAS Y PRODUCCIÓN", "SALIDAS A COCINA", "EXISTENCIA FINAL"]
        if es_porcionador
        else ["EXISTENCIA INICIAL", "ENTRADAS", "VENTAS", "EXISTENCIA FINAL"]
    )
    cols += ["CONTEO FÍSICO", "DIFERENCIA"]

    wb = Workbook()
    ws = wb.active
    ws.title = f"formato {hoja}"
    ws.cell(row=1, column=1, value=f"EXISTENCIAS DE PORCIONES DE {titulo}        DÍA: {fecha.strftime('%d/%m/%Y')}").font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    encabezado_fill = PatternFill("solid", fgColor="D9D9D9")
    for i, nombre in enumerate(cols, start=1):
        c = ws.cell(row=3, column=i, value=nombre)
        c.font = Font(bold=True)
        c.fill = encabezado_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fila = 4
    seccion_actual = None
    for r in filas:
        if r["seccion"] != seccion_actual:
            seccion_actual = r["seccion"]
            c = ws.cell(row=fila, column=1, value=r["seccion"])
            c.font = Font(bold=True)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(cols))
            fila += 1
        ws.cell(row=fila, column=1, value=r["producto"])
        ws.cell(row=fila, column=2, value=r["presentacion"])
        ws.cell(row=fila, column=3, value=r["unidad"])
        ws.cell(row=fila, column=4, value=r["gramaje"])
        ws.cell(row=fila, column=5, value=r["inicial"])
        ws.cell(row=fila, column=6, value=r["entradas"])
        ws.cell(row=fila, column=7, value=r["salidas"])
        # existencia final: fórmula viva (el sistema también la calcula en cadena)
        ws.cell(row=fila, column=8, value=f"=E{fila}+F{fila}-G{fila}")
        ws.cell(row=fila, column=9, value=r["conteo_fisico"])
        ws.cell(row=fila, column=10, value=f'=IF(I{fila}="","",I{fila}-H{fila})')
        fila += 1

    fila += 1
    ws.cell(row=fila, column=1, value="Este inventario se encuentra bajo mi resguardo y custodia, mismo que debo cuidar y manejar bajo control.")
    fila += 2
    nombre_responsable = (dia.responsable_porcionador if es_porcionador else dia.responsable_cocina) if dia else None
    ws.cell(row=fila, column=3, value=nombre_responsable or "")
    ws.cell(row=fila + 1, column=3, value="Nombre y firma del porcionador" if es_porcionador else "Nombre y firma del responsable")

    for letra, ancho in zip("ABCDEFGHIJ", (38, 30, 10, 12, 14, 14, 14, 14, 12, 12)):
        ws.column_dimensions[letra].width = ancho

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=reporte_{hoja}_{fecha.isoformat()}.xlsx"},
    )


@router.put("/captura")
async def guardar_captura(
    datos: CapturaIn,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Upsert por (fecha, renglón). Un renglón con todo vacío borra su captura (no hay movimiento: el saldo se arrastra).

    `inicial` y `conteo_fisico` (existencia final) son correcciones manuales opcionales: si vienen, mandan sobre el arrastre.
    Con `completar_hoja` el día queda registrado para todos los productos de la hoja; esas filas sin datos NO fijan
    inicial ni final (se calculan en cadena), así corregir un día anterior sigue actualizando los siguientes.
    Las entradas de cocina con equivalente en porcionador se sugieren solas; si se envían, mandan las capturadas.
    """
    ids = [r.renglon_id for r in datos.renglones]
    renglones = {
        r.id: r
        for r in (await db.execute(select(ReporteProduccionRenglon).where(ReporteProduccionRenglon.id.in_(ids)))).scalars().all()
    }
    desconocidos = sorted(set(ids) - set(renglones))
    if desconocidos:
        raise HTTPException(status_code=400, detail=f"Renglones inexistentes: {desconocidos}")

    existentes = {
        c.renglon_id: c
        for c in (await db.execute(select(ReporteProduccionCaptura).where(ReporteProduccionCaptura.fecha == datos.fecha))).scalars().all()
    }
    guardados = borrados = 0
    for item in datos.renglones:
        entradas = item.entradas
        vacio = all(v is None for v in (item.inicial, entradas, item.salidas, item.conteo_fisico))
        actual = existentes.get(item.renglon_id)
        if vacio:
            if actual is not None:
                await db.execute(delete(ReporteProduccionCaptura).where(ReporteProduccionCaptura.id == actual.id))
                borrados += 1
            continue
        if actual is None:
            db.add(
                ReporteProduccionCaptura(
                    fecha=datos.fecha, renglon_id=item.renglon_id, inicial=item.inicial, entradas=entradas,
                    salidas=item.salidas, conteo_fisico=item.conteo_fisico,
                )
            )
        else:
            actual.inicial, actual.entradas, actual.salidas, actual.conteo_fisico = (
                item.inicial, entradas, item.salidas, item.conteo_fisico,
            )
        guardados += 1
    await db.flush()

    completados = 0
    if datos.completar_hoja is not None:
        _validar_hoja(datos.completar_hoja)
        ids_hoja = [
            r.id
            for r in (await _renglones_activos(db))
            if r.hoja == datos.completar_hoja
        ]
        con_fila = set(
            (
                await db.execute(
                    select(ReporteProduccionCaptura.renglon_id).where(
                        ReporteProduccionCaptura.fecha == datos.fecha, ReporteProduccionCaptura.renglon_id.in_(ids_hoja)
                    )
                )
            ).scalars().all()
        ) if ids_hoja else set()
        for rid in ids_hoja:
            if rid not in con_fila:
                db.add(ReporteProduccionCaptura(fecha=datos.fecha, renglon_id=rid))
                completados += 1
        await db.flush()
    return {"guardados": guardados, "borrados": borrados, "completados": completados}


@router.put("/dia")
async def guardar_dia(
    datos: DiaIn,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    dia = (await db.execute(select(ReporteProduccionDia).where(ReporteProduccionDia.fecha == datos.fecha))).scalar_one_or_none()
    if dia is None:
        dia = ReporteProduccionDia(fecha=datos.fecha)
        db.add(dia)
    # solo se tocan los campos enviados: cada hoja guarda su propio responsable sin pisar el de la otra
    for campo in datos.model_fields_set & {"responsable_porcionador", "responsable_cocina", "notas"}:
        setattr(dia, campo, (getattr(datos, campo) or "").strip() or None)
    await db.flush()
    return {"ok": True, **_dia_a_dict(dia)}


# ── Pesaje de limpieza ─────────────────────────────────────────────────
# El porcionador pesa lo que recibe (bruto) y lo que le queda limpio; la merma es la diferencia y el factor
# bruto/limpio sale de ahí. Alimenta el inventario en bruto y el control contra las porciones producidas.


class PesajeIn(BaseModel):
    fecha: date
    producto: str
    bruto_kg: float
    limpio_kg: float


class PesajePatch(BaseModel):
    bruto_kg: float | None = None
    limpio_kg: float | None = None


def _validar_pesaje(bruto: float, limpio: float) -> None:
    if bruto <= 0 or limpio <= 0:
        raise HTTPException(status_code=400, detail="El bruto y el limpio deben ser mayores a 0")
    if limpio > bruto:
        raise HTTPException(status_code=400, detail="El limpio no puede ser mayor que el bruto (la diferencia es la merma)")


def _pesaje_a_dict(p: ReporteProduccionPesaje) -> dict:
    return {
        "id": p.id,
        "fecha": p.fecha.isoformat(),
        "producto": p.producto,
        "bruto_kg": p.bruto_kg,
        "limpio_kg": p.limpio_kg,
        "merma_kg": merma_kg(p.bruto_kg, p.limpio_kg),
        "factor": round(p.bruto_kg / p.limpio_kg, 4) if p.limpio_kg > 0 else None,
    }


@router.get("/pesajes")
async def listar_pesajes(
    fecha: date = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Pesajes del día y, por producto base procesable, el control limpio pesado vs porciones producidas."""
    pesajes = list(
        (
            await db.execute(
                select(ReporteProduccionPesaje).where(ReporteProduccionPesaje.fecha == fecha).order_by(ReporteProduccionPesaje.id)
            )
        ).scalars().all()
    )
    porcionador = [r for r in await _renglones_activos(db) if r.hoja == "porcionador"]
    filas, _ = await _calcular_dia(db, fecha, "porcionador")
    entradas = {f["id"]: f["entradas"] for f in filas}  # "Entradas y producción" efectivas del día
    salidas = {f["id"]: f["salidas"] for f in filas}  # lo usado (para el limpio intermedio: filetitos)

    por_producto: dict[str, list] = {}
    for r in porcionador:
        por_producto.setdefault(r.producto, []).append(r)
    productos = []
    for producto, rs in por_producto.items():
        if all((r.unidad or "").upper() == "KG" for r in rs):
            continue  # sin presentaciones procesadas: no hay limpieza que pesar
        ps = [p for p in pesajes if p.producto == producto]
        control = calcular_control_limpieza(ps, rs, entradas, salidas)
        productos.append({
            "producto": producto,
            "activo_hoy": bool(ps) or control["kg_porciones"] > 0 or bool(control["sin_gramaje"]),
            "pesajes": [_pesaje_a_dict(p) for p in ps],
            "control": control,
        })
    return {"fecha": fecha.isoformat(), "productos": productos}


@router.post("/pesajes")
async def crear_pesaje(
    datos: PesajeIn,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    _validar_pesaje(datos.bruto_kg, datos.limpio_kg)
    producto = datos.producto.strip()
    if producto not in {r.producto for r in await _renglones_activos(db)}:
        raise HTTPException(status_code=400, detail="Ese producto no está en el catálogo")
    pesaje = ReporteProduccionPesaje(fecha=datos.fecha, producto=producto, bruto_kg=datos.bruto_kg, limpio_kg=datos.limpio_kg)
    db.add(pesaje)
    await db.flush()
    return _pesaje_a_dict(pesaje)


@router.patch("/pesajes/{pesaje_id}")
async def editar_pesaje(
    pesaje_id: int,
    datos: PesajePatch,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    pesaje = await db.get(ReporteProduccionPesaje, pesaje_id)
    if pesaje is None:
        raise HTTPException(status_code=404, detail="Pesaje no encontrado")
    bruto = datos.bruto_kg if datos.bruto_kg is not None else pesaje.bruto_kg
    limpio = datos.limpio_kg if datos.limpio_kg is not None else pesaje.limpio_kg
    _validar_pesaje(bruto, limpio)
    pesaje.bruto_kg, pesaje.limpio_kg = bruto, limpio
    await db.flush()
    return _pesaje_a_dict(pesaje)


@router.delete("/pesajes/{pesaje_id}")
async def borrar_pesaje(
    pesaje_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    pesaje = await db.get(ReporteProduccionPesaje, pesaje_id)
    if pesaje is None:
        raise HTTPException(status_code=404, detail="Pesaje no encontrado")
    await db.delete(pesaje)
    await db.flush()
    return {"ok": True}


# ── Inventario en kg bruto y en pesos ──────────────────────────────────


class PrecioProductoIn(BaseModel):
    producto: str
    precio_kg: float | None = None


async def _precios_por_producto(db: AsyncSession) -> dict[str, float | None]:
    return {p.producto: p.precio_kg for p in (await db.execute(select(ReporteProduccionProductoBase))).scalars().all()}


BASES_INVENTARIO = ("existencia", "produccion")


def _validar_base(base: str) -> str:
    if base not in BASES_INVENTARIO:
        raise HTTPException(status_code=400, detail=f"base debe ser una de {list(BASES_INVENTARIO)}")
    return base


async def _inventario_bruto(db: AsyncSession, fecha: date, hoja: str, base: str = "existencia") -> dict:
    """Kg limpio, kg bruto y pesos de la hoja.

    - `existencia`: todo lo que hay (existencia del día o el saldo que arrastra).
    - `produccion`: SOLO "Entradas y producción" del día en presentaciones procesadas (unidad distinta de KG):
      lo que hizo el porcionador. Los productos sin producción no se calculan.
    """
    filas, _ = await _calcular_dia(db, fecha, hoja)
    if base == "produccion":
        cantidades = {}
        for f in filas:
            if (f["unidad"] or "").upper() != "KG":
                cantidades[f["id"]] = f["entradas"]
            elif f.get("es_limpio"):
                # sobrante ya limpio (filetitos): lo que sobró suma y lo que se usó para hacer porciones resta,
                # porque esas porciones ya están contadas (si no, se contaría dos veces)
                cantidades[f["id"]] = (
                    round((f["entradas"] or 0.0) - (f["salidas"] or 0.0), 4)
                    if (f["entradas"] is not None or f["salidas"] is not None)
                    else None
                )
            else:
                cantidades[f["id"]] = None
    else:
        cantidades = {
            f["id"]: (f["existencia_final"] if f["existencia_final"] is not None else f["inicial"]) for f in filas
        }
    todos = await _renglones_activos(db)
    pesajes = list(
        (await db.execute(select(ReporteProduccionPesaje).where(ReporteProduccionPesaje.fecha <= fecha))).scalars().all()
    )
    return calcular_inventario_bruto(
        [r for r in todos if r.hoja == hoja],
        cantidades,
        await _precios_por_producto(db),
        catalogo_referencia=[r for r in todos if r.hoja != hoja],
        # el factor del último pesaje de limpieza (≤ fecha) manda sobre la prueba fija del catálogo
        factores_pesaje=factores_por_producto(pesajes, fecha),
    )


@router.get("/productos-base")
async def listar_productos_base(
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Productos base del catálogo con su precio por kg (None si aún no se captura)."""
    precios = await _precios_por_producto(db)
    productos = dict.fromkeys(r.producto for r in await _renglones_activos(db))
    return [{"producto": p, "precio_kg": precios.get(p)} for p in productos]


@router.put("/productos-base")
async def guardar_precio_producto_base(
    datos: PrecioProductoIn,
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    if datos.precio_kg is not None and datos.precio_kg < 0:
        raise HTTPException(status_code=400, detail="El precio por kg no puede ser negativo")
    nombre = datos.producto.strip()
    fila = (
        await db.execute(select(ReporteProduccionProductoBase).where(ReporteProduccionProductoBase.producto == nombre))
    ).scalar_one_or_none()
    if fila is None:
        db.add(ReporteProduccionProductoBase(producto=nombre, precio_kg=datos.precio_kg))
    else:
        fila.precio_kg = datos.precio_kg
    await db.flush()
    return {"producto": nombre, "precio_kg": datos.precio_kg}


@router.get("/inventario-bruto")
async def obtener_inventario_bruto(
    fecha: date = Query(...),
    hoja: str = Query("porcionador"),
    base: str = Query("existencia"),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    _validar_hoja(hoja)
    _validar_base(base)
    return {"fecha": fecha.isoformat(), "hoja": hoja, "base": base, **(await _inventario_bruto(db, fecha, hoja, base))}


@router.get("/inventario-bruto/export")
async def exportar_inventario_bruto(
    fecha: date = Query(...),
    hoja: str = Query("porcionador"),
    base: str = Query("existencia"),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Excel con el esquema de la hoja PORCIONADOR y fórmulas vivas: kg limpio = cantidad × gramaje,
    factor = bruto/limpio de la prueba, kg bruto = kg limpio × factor, valor = kg bruto × precio por kg."""
    _validar_hoja(hoja)
    _validar_base(base)
    inv = await _inventario_bruto(db, fecha, hoja, base)
    producido = base == "produccion"

    wb = Workbook()
    ws = wb.active
    ws.title = f"inventario {hoja}"
    ws.cell(row=1, column=1, value=f"{'PRODUCCIÓN' if producido else 'INVENTARIO'} EN KG BRUTO Y VALOR — {hoja.upper()}        DÍA: {fecha.strftime('%d/%m/%Y')}").font = Font(bold=True, size=12)
    encabezados = ["PRODUCTO / PRESENTACIÓN", "UNIDAD", "ENTRADAS Y PRODUCCIÓN" if producido else "CANTIDAD", "GRAMAJE (KG)", "KG LIMPIO", "PRUEBA BRUTO", "PRUEBA LIMPIO", "FACTOR", "KG BRUTO", "PRECIO / KG", "VALOR $"]
    fill = PatternFill("solid", fgColor="D9D9D9")
    for i, nombre in enumerate(encabezados, start=1):
        c = ws.cell(row=3, column=i, value=nombre)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fila = 4
    filas_grupo: list[int] = []
    gris = PatternFill("solid", fgColor="EDEDED")
    for grupo in inv["grupos"]:
        fila_grupo = fila
        for col in range(1, 12):
            ws.cell(row=fila_grupo, column=col).fill = gris
        ws.cell(row=fila_grupo, column=1, value=grupo["seccion"]).font = Font(bold=True, size=11)
        fila += 1
        filas_totales: list[int] = []
        for prod in grupo["productos"]:
            fila_base = fila
            ws.cell(row=fila_base, column=1, value=prod["producto"]).font = Font(bold=True)
            hijas = list(range(fila_base + 1, fila_base + 1 + len(prod["filas"])))
            for r_hija, f in zip(hijas, prod["filas"]):
                ws.cell(row=r_hija, column=1, value=f"   {f['presentacion'] or '(sin presentación)'}")
                ws.cell(row=r_hija, column=2, value=f["unidad"])
                ws.cell(row=r_hija, column=3, value=f["cantidad"])
                ws.cell(row=r_hija, column=4, value=f["gramaje_kg"])
                ws.cell(row=r_hija, column=5, value=f'=IF(OR(C{r_hija}="",D{r_hija}=""),"",C{r_hija}*D{r_hija})')
                ws.cell(row=r_hija, column=6, value=f["prueba_bruto_kg"])
                ws.cell(row=r_hija, column=7, value=f["prueba_limpio_kg"])
                ws.cell(row=r_hija, column=8, value=f'=IF(AND(ISNUMBER(F{r_hija}),ISNUMBER(G{r_hija})),IF(G{r_hija}>0,F{r_hija}/G{r_hija},""),"")')
                if (f["unidad"] or "").upper() == "KG":  # ya viene en bruto si no hay prueba
                    ws.cell(row=r_hija, column=9, value=f'=IF(E{r_hija}="","",IF(H{r_hija}="",E{r_hija},E{r_hija}*H{r_hija}))')
                else:
                    ws.cell(row=r_hija, column=9, value=f'=IF(OR(E{r_hija}="",H{r_hija}=""),"",E{r_hija}*H{r_hija})')
            if hijas:
                ws.cell(row=fila_base, column=9, value=f"=SUM(I{hijas[0]}:I{hijas[-1]})")
            ws.cell(row=fila_base, column=10, value=prod["precio_kg"])
            ws.cell(row=fila_base, column=11, value=f'=IF(J{fila_base}="","",I{fila_base}*J{fila_base})')
            filas_totales.append(fila_base)
            fila = (hijas[-1] if hijas else fila_base) + 1
        # subtotal del grupo: suma de sus productos (fórmulas vivas)
        if filas_totales:
            ws.cell(row=fila_grupo, column=9, value="=" + "+".join(f"I{r}" for r in filas_totales))
            ws.cell(row=fila_grupo, column=11, value="=" + "+".join(f"N(K{r})" for r in filas_totales))
        for col in (9, 11):
            ws.cell(row=fila_grupo, column=col).font = Font(bold=True)
        filas_grupo.append(fila_grupo)
        fila += 1  # renglón en blanco entre grupos

    ws.cell(row=fila, column=1, value="TOTAL PRODUCIDO" if producido else "TOTAL INVENTARIO").font = Font(bold=True)
    ws.cell(row=fila, column=9, value="=" + "+".join(f"I{r}" for r in filas_grupo) if filas_grupo else 0)
    ws.cell(row=fila, column=11, value="=" + "+".join(f"N(K{r})" for r in filas_grupo) if filas_grupo else 0)
    for col in (9, 11):
        ws.cell(row=fila, column=col).font = Font(bold=True)

    for letra, ancho in zip("ABCDEFGHIJK", (46, 10, 11, 12, 11, 12, 13, 9, 11, 12, 14)):
        ws.column_dimensions[letra].width = ancho

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={'produccion' if producido else 'inventario'}_bruto_{hoja}_{fecha.isoformat()}.xlsx"},
    )


@router.delete("/dia")
async def borrar_dia(
    fecha: date = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Borra TODA la captura de un día (ambas hojas, responsables y notas). No toca el catálogo ni otros días."""
    capturas = (await db.execute(delete(ReporteProduccionCaptura).where(ReporteProduccionCaptura.fecha == fecha))).rowcount
    await db.execute(delete(ReporteProduccionDia).where(ReporteProduccionDia.fecha == fecha))
    return {"borrados": capturas or 0}


# Ruta con query params, definida al final del grupo de rutas literales.
@router.get("")
async def obtener_dia(
    fecha: date = Query(...),
    hoja: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Renglones activos de la hoja con la cadena de saldos resuelta para ese día."""
    _validar_hoja(hoja)
    filas, fecha_prev = await _calcular_dia(db, fecha, hoja)
    dia = (await db.execute(select(ReporteProduccionDia).where(ReporteProduccionDia.fecha == fecha))).scalar_one_or_none()
    return {
        "fecha": fecha.isoformat(),
        "hoja": hoja,
        "fecha_dia_anterior": fecha_prev.isoformat() if fecha_prev else None,
        "dia": _dia_a_dict(dia),
        "renglones": filas,
    }
