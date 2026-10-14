"""
Pagos de Recepción — gastos que hace un huésped durante su estancia.
Módulo independiente del flujo de Ingresos/Pagos/Conciliación y de
Facturación: por ahora no lee ni escribe la base de datos.

El archivo origen está desordenado (sin fila de encabezados confiable), así
que se lee crudo (`header=None`) y las columnas de interés se identifican
por posición fija de letra, igual que `cfdi_analysis.py` hace para
Pagos/Ingresos.
"""

import re
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

COLUMNAS_PAGOS_RECEPCION = ["B", "F", "L", "N", "S", "U", "AD", "AF", "AJ", "AP", "BD"]


async def leer_excel_pagos_recepcion(file_content: bytes) -> pd.DataFrame:
    """
    Leer un archivo Excel de Pagos de Recepción crudo (sin encabezados),
    con fallback a xlrd para .xls.
    """
    import zipfile

    try:
        df = pd.read_excel(BytesIO(file_content), header=None)
    except zipfile.BadZipFile:
        df = pd.read_excel(BytesIO(file_content), header=None, engine="xlrd")
    return df


def extraer_columnas_pagos_recepcion(df: pd.DataFrame) -> pd.DataFrame:
    """
    Selecciona únicamente las columnas fijas en COLUMNAS_PAGOS_RECEPCION
    (por letra de columna Excel) y las etiqueta con esa misma letra. Las
    celdas vacías se normalizan a `None` (pandas las da como NaN, que es
    "truthy" en Python y rompe cualquier `fila.get(...) or ""`).
    """
    indices = [column_index_from_string(letra) - 1 for letra in COLUMNAS_PAGOS_RECEPCION]
    df_extraido = df.iloc[:, indices].copy()
    df_extraido.columns = COLUMNAS_PAGOS_RECEPCION
    return df_extraido.astype(object).where(pd.notna(df_extraido), None)


# ── Archivo de Movimiento (conciliación de forma de pago) ───────────
# Archivo Excel distinto al de arriba: trae, agrupadas por método de pago
# (MN. EFECTIVO / VISA / MASTER CARD / etc.), las fechas y montos de cómo
# pagó cada huésped. Por ahora solo se extraen sus columnas fijas para
# previsualizarlas — la conciliación contra las cuentas de Tabla 2 se
# diseña en una siguiente iteración.

COLUMNAS_MOVIMIENTO_PAGOS_RECEPCION = ["C", "H", "M", "Y"]


async def leer_excel_movimiento_pagos_recepcion(file_content: bytes) -> pd.DataFrame:
    """
    Leer un archivo Excel de Movimiento (Pagos de Recepción) crudo (sin
    encabezados), con fallback a xlrd para .xls.
    """
    import zipfile

    try:
        df = pd.read_excel(BytesIO(file_content), header=None)
    except zipfile.BadZipFile:
        df = pd.read_excel(BytesIO(file_content), header=None, engine="xlrd")
    return df


def extraer_columnas_movimiento_pagos_recepcion(df: pd.DataFrame) -> pd.DataFrame:
    """
    Selecciona únicamente las columnas fijas en
    COLUMNAS_MOVIMIENTO_PAGOS_RECEPCION (por letra de columna Excel) y las
    etiqueta con esa misma letra. Las celdas vacías se normalizan a `None`:
    pandas las da como NaN, que es "truthy" en Python, así que
    `fila.get("H") or ""` NUNCA daba `""` para una celda vacía — y
    `_es_fila_titulo_movimiento`/`_es_fila_dato_movimiento` dependen
    exactamente de ese `== ""` para distinguir título/dato/vacío. Sin este
    tratamiento ninguna fila título se reconocía jamás y la conciliación de
    Movimiento no agrupaba nada (confirmado con archivos reales).
    """
    indices = [column_index_from_string(letra) - 1 for letra in COLUMNAS_MOVIMIENTO_PAGOS_RECEPCION]
    df_extraido = df.iloc[:, indices].copy()
    df_extraido.columns = COLUMNAS_MOVIMIENTO_PAGOS_RECEPCION
    return df_extraido.astype(object).where(pd.notna(df_extraido), None)


COLUMNAS_MOVIMIENTOS_PERIODO_PAGOS_RECEPCION = ["A", "B", "C", "G"]


async def leer_excel_movimientos_periodo_pagos_recepcion(file_content: bytes) -> pd.DataFrame:
    """
    Leer un archivo Excel de Movimientos por Periodo (Pagos de Recepción)
    crudo (sin encabezados), con fallback a xlrd para .xls.
    """
    import zipfile

    try:
        df = pd.read_excel(BytesIO(file_content), header=None)
    except zipfile.BadZipFile:
        df = pd.read_excel(BytesIO(file_content), header=None, engine="xlrd")
    return df


def extraer_columnas_movimientos_periodo_pagos_recepcion(df: pd.DataFrame) -> pd.DataFrame:
    """
    Selecciona únicamente las columnas fijas en
    COLUMNAS_MOVIMIENTOS_PERIODO_PAGOS_RECEPCION (por letra de columna
    Excel) y las etiqueta con esa misma letra. Celdas vacías -> `None`
    (ver `extraer_columnas_movimiento_pagos_recepcion`); `
    extraer_pagos_amex_de_movimientos_periodo` hacía este mismo tratamiento
    por su cuenta justo antes de usar este resultado, ya no hace falta ahí.
    """
    indices = [column_index_from_string(letra) - 1 for letra in COLUMNAS_MOVIMIENTOS_PERIODO_PAGOS_RECEPCION]
    df_extraido = df.iloc[:, indices].copy()
    df_extraido.columns = COLUMNAS_MOVIMIENTOS_PERIODO_PAGOS_RECEPCION
    return df_extraido.astype(object).where(pd.notna(df_extraido), None)


def extraer_pagos_amex_de_movimientos_periodo(df_extraido: pd.DataFrame) -> list[dict]:
    """
    Recorre el archivo de Movimientos por Periodo llevando un flag "dentro
    del segmento American Express": se activa cuando la columna A (título
    de segmento) empieza con "SEGMENTO" y contiene "AMERICAN EXPRESS", y se
    desactiva cuando la columna C (fila de cierre/total del segmento)
    contiene "AMERICAN EXPRESS" — mismo convenio abre-en-A/cierra-en-C que
    el archivo homónimo de Facturación (`extraer_cuentas_por_segmento`).
    Mientras el flag está activo, cada fila con Descripción (columna C) no
    vacía y Monto (columna G) distinto de 0 se toma como un pago real de
    Amex; esto descarta solo automáticamente la fila de encabezado
    ("Descripción"), la fila de subtotal (C vacía) y la fila de cierre.
    """
    dentro_de_amex = False
    resultado: list[dict] = []

    for fila in df_extraido.to_dict(orient="records"):
        col_a = str(fila.get("A") or "").strip()
        col_c = str(fila.get("C") or "").strip()

        if col_a.upper().startswith("SEGMENTO") and "AMERICAN EXPRESS" in col_a.upper():
            dentro_de_amex = True
            continue
        if "AMERICAN EXPRESS" in col_c.upper():
            dentro_de_amex = False
            continue
        if not dentro_de_amex:
            continue

        descripcion = col_c
        if descripcion == "" or descripcion.upper() in ("DESCRIPCIÓN", "DESCRIPCION"):
            continue
        monto = _parse_monto(fila.get("G"))
        if monto == 0.0:
            continue

        resultado.append(
            {
                "folio": str(fila.get("A") or "").strip(),
                "fecha": str(fila.get("B") or "").strip(),
                "huesped": descripcion,
                "monto": monto,
            }
        )

    return resultado


def _es_folio(valor) -> bool:
    """
    Una fila abre un bloque nuevo cuando su columna B es puramente un
    número de folio (ej. 38785). Pandas puede leerlo como int/float
    (38785.0) o como texto "38785" — cualquiera de las dos formas cuenta,
    siempre que no tenga letras ni otros caracteres.
    """
    if valor is None:
        return False
    if isinstance(valor, bool):
        return False
    if isinstance(valor, (int, float)):
        return not pd.isna(valor)
    return bool(re.fullmatch(r"\d+", str(valor).strip()))


def _truncar_en_primer_saldo(filas: list[dict]) -> list[dict]:
    """
    Cada cuenta real termina en una sola fila "SALDO". En el último bloque
    del reporte, el resumen con la suma de todas las cuentas viene pegado
    (sin folio propio) después de ese "SALDO" real, así que trae un segundo
    "SALDO" propio al final. Cortamos el bloque en la primera "SALDO" para
    quedarnos solo con la cuenta real y descartar cualquier cola pegada
    (no afecta a los bloques normales, que ya terminan justo ahí).
    """
    for i, fila in enumerate(filas):
        if str(fila.get("B") or "").strip().upper() == "SALDO":
            return filas[: i + 1]
    return filas


def agrupar_bloques_pagos_recepcion(df_extraido: pd.DataFrame) -> list[dict]:
    """
    Agrupa las filas ya extraídas en bloques por huésped: cada fila cuya
    columna B es un número de folio abre un bloque nuevo, que se lleva esa
    fila y todas las siguientes hasta el próximo folio. Filas previas al
    primer folio detectado (ruido antes de que empiece el primer bloque
    real) se descartan. Cada bloque se trunca en su primera fila "SALDO"
    (ver `_truncar_en_primer_saldo`).
    """
    df_json = df_extraido.astype(object).where(pd.notna(df_extraido), None)
    bloques: list[dict] = []

    for fila in df_json.to_dict(orient="records"):
        if _es_folio(fila.get("B")):
            consecutivo = int(float(fila["B"]))
            bloques.append({"consecutivo": consecutivo, "filas": [fila]})
        elif bloques:
            bloques[-1]["filas"].append(fila)
        # Filas antes del primer folio: se descartan (sin bloque abierto aún).

    for bloque in bloques:
        bloque["filas"] = _truncar_en_primer_saldo(bloque["filas"])

    return bloques


def _parse_monto(valor) -> float:
    """
    Convierte un valor de celda (float, o texto tipo "1,012.97" / "$1,012.97"
    / "-3,633.00") a float. None/vacío/no numérico -> 0.0.
    """
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return 0.0 if pd.isna(valor) else float(valor)
    texto = str(valor).strip().replace("$", "").replace(",", "")
    if texto == "":
        return 0.0
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _es_iva(etiqueta_norm: str) -> bool:
    """
    Reconoce una línea de IVA sin importar cómo esté escrita: "IVA PALAPA",
    "IVA CARRETA", pero también "I.V.A." (con puntos, sin sufijo) — se
    quitan los puntos antes de comparar el prefijo.
    """
    return etiqueta_norm.replace(".", "").startswith("IVA")


def _indice_primer_subtotal(filas: list[dict]) -> int | None:
    for i, fila in enumerate(filas):
        if str(fila.get("B") or "").strip().upper() == "SUBTOTAL":
            return i
    return None


def _subtotal1_bloque(filas: list[dict]) -> float:
    """
    SUBTOTAL1 (neto, sin impuestos ni propina) = suma de los renglones de
    "Concepto" reales entre la fila del folio y la primera fila "Subtotal",
    excluyendo PROPINAS/IVA/ISH (esas se calculan y muestran aparte) — más
    cualquier cargo suelto que aparezca después de esa primera "Subtotal"
    (ver `_cargos_extra_bloque`), ya que algunos reportes mezclan cargos
    reales (no impuestos) dentro de lo que normalmente es el desglose de
    IVA — ej. "AGENCIA DE VIAJES", "LLAVE EXTRAVIADA".

    Esto cubre dos formatos de reporte distintos vistos en archivos reales:
    - Cuentas de huésped normales: los cargos (incluida a veces la propina)
      van todos antes de la primera "Subtotal"; el IVA va después, en un
      desglose aparte.
    - Cuentas de agencia (ej. RENTALS-EXPEDIA): el IVA ("I.V.A.") viene
      mezclado como un renglón de cargo más, antes de la primera "Subtotal"
      — por eso hay que excluirlo explícitamente en vez de asumir que todo
      lo anterior a la primera "Subtotal" es neto.
    """
    idx_subtotal = _indice_primer_subtotal(filas)
    fin = idx_subtotal if idx_subtotal is not None else len(filas)
    total = 0.0
    for fila in filas[1:fin]:  # [1:] excluye la fila del folio (índice 0)
        etiqueta_norm = str(fila.get("B") or "").strip().upper()
        if etiqueta_norm in ("PROPINAS", "IMPTO S/HOSPEDAJE") or _es_iva(etiqueta_norm):
            continue
        total += _parse_monto(fila.get("S"))
    return total + _cargos_extra_bloque(filas)


def _iva_bloque(filas: list[dict]) -> float:
    """
    Suma el monto (columna S) de todas las filas de IVA (ver `_es_iva`),
    sin importar en qué parte del bloque aparezcan: "IVA PALAPA",
    "IVA CARRETA", "IVA AGENCIA", "IVA TIENDA", "I.V.A.", etc.

    OJO: la última fila "Subtotal" del desglose de IVA NO sirve para esto —
    es un acumulado que el propio reporte imprime, y puede incluir de más
    una línea "AGENCIA DE VIAJES" (sin el prefijo "IVA") que es un cargo/
    comisión aparte, no un impuesto. Sumar solo las líneas de IVA evita que
    ese cargo se cuele en el total de IVA.
    """
    return sum(
        _parse_monto(fila.get("S"))
        for fila in filas
        if _es_iva(str(fila.get("B") or "").strip().upper())
    )


_ETIQUETAS_ADMINISTRATIVAS = {"SUBTOTAL", "ABONO A CTA.", "SALDO", "IMPTO S/HOSPEDAJE"}


def _cargos_extra_bloque(filas: list[dict]) -> float:
    """
    Después de la primera fila "Subtotal", el reporte normalmente solo trae
    el desglose de IVA ("IVA PALAPA", "IVA CARRETA", ...), pero a veces
    mete ahí mismo cargos reales sueltos que no son impuestos — ej.
    "AGENCIA DE VIAJES" (comisión de agencia), "LLAVE EXTRAVIADA" (cargo
    por llave perdida). Cualquier línea en esa zona que no sea una etiqueta
    administrativa (Subtotal/ABONO A CTA./SALDO), ni ISH, ni una línea de
    IVA (ver `_es_iva`) es un cargo real y se suma aquí aparte — ya que
    `_iva_bloque` solo suma las líneas de IVA y el resto se perdía.
    """
    idx_subtotal = _indice_primer_subtotal(filas)
    if idx_subtotal is None:
        return 0.0
    total = 0.0
    for fila in filas[idx_subtotal + 1 :]:
        etiqueta_norm = str(fila.get("B") or "").strip().upper()
        if etiqueta_norm == "" or etiqueta_norm in _ETIQUETAS_ADMINISTRATIVAS or _es_iva(etiqueta_norm):
            continue
        total += _parse_monto(fila.get("S"))
    return total


def _propina_bloque(filas: list[dict]) -> float:
    """
    Busca la fila cuya columna B sea literalmente "PROPINAS" y regresa el
    monto de su columna S. 0.0 si no aparece.
    """
    for fila in filas:
        etiqueta = fila.get("B")
        if etiqueta is not None and str(etiqueta).strip().upper() == "PROPINAS":
            return _parse_monto(fila.get("S"))
    return 0.0


def _ish_bloque(filas: list[dict]) -> float:
    """
    Busca la fila cuya columna B sea literalmente "IMPTO S/HOSPEDAJE"
    (Impuesto Sobre Hospedaje) y regresa el monto de su columna S. 0.0 si
    no aparece (la mayoría de los archivos del usuario no lo traen).
    """
    for fila in filas:
        etiqueta = fila.get("B")
        if etiqueta is not None and str(etiqueta).strip().upper() == "IMPTO S/HOSPEDAJE":
            return _parse_monto(fila.get("S"))
    return 0.0


def calcular_metricas_desde_valores(subtotal1: float, iva16: float, ish3: float, propina: float) -> dict:
    """
    Las 4 fórmulas derivadas (Subtotal2/Gravado/Tasa 0%/Total), operando
    directo sobre valores ya conocidos en vez de un bloque crudo — usada
    tanto por `calcular_metricas_cuenta` (al procesar un archivo) como por
    la edición manual de "Pagos del Mes" (al corregir un campo base).
    """
    subtotal2 = subtotal1 + iva16
    gravado = iva16 / 0.16 + iva16
    tasa0 = subtotal2 - gravado
    total = subtotal1 + iva16 + ish3 + propina

    return {
        "subtotal1": subtotal1,
        "iva16": iva16,
        "ish3": ish3,
        "subtotal2": subtotal2,
        "gravado": gravado,
        "tasa0": tasa0,
        "propina": propina,
        "total": total,
    }


def calcular_metricas_cuenta(bloque: dict) -> dict:
    """
    Calcula las 8 métricas financieras de una cuenta a partir de su bloque
    ya agrupado: Subtotal1, IVA 16%, ISH 3%, Subtotal2, Gravado, Tasa 0%,
    Propina y Total.
    """
    filas = bloque["filas"]
    subtotal1 = _subtotal1_bloque(filas)
    iva16 = _iva_bloque(filas)
    ish3 = _ish_bloque(filas)
    propina = _propina_bloque(filas)

    return calcular_metricas_desde_valores(subtotal1, iva16, ish3, propina)


def construir_tabla1_pagos_recepcion(bloques: list[dict]) -> list[dict]:
    """Una fila por cuenta con las 8 métricas calculadas."""
    return [{"consecutivo": bloque["consecutivo"], **calcular_metricas_cuenta(bloque)} for bloque in bloques]


def _primer_valor_no_vacio(filas: list[dict], columna: str):
    """
    Regresa el primer valor no vacío de `columna` recorriendo las filas del
    bloque de arriba hacia abajo. Habitación/Huésped suelen venir en la
    fila del folio, pero la Fecha de Emisión no siempre — a veces solo
    aparece más abajo, en una fila posterior del mismo bloque.

    Se saltan las filas de encabezado de sub-sección (columna B empieza con
    "#", ej. "#Reserva", "#Factura") — sus demás columnas traen etiquetas
    de campo ("Titular", "Fecha-Hora Emisión", "Cajero", "SubTotal", "Tax",
    "TOTAL"), no datos reales, y si no se saltan se toman por error como si
    fueran el valor buscado.
    """
    for fila in filas:
        etiqueta_b = str(fila.get("B") or "").strip()
        if etiqueta_b.startswith("#"):
            continue
        valor = fila.get(columna)
        if valor is not None and str(valor).strip() != "":
            return valor
    return None


def construir_tabla2_pagos_recepcion(bloques: list[dict]) -> list[dict]:
    """
    Una fila por cuenta con identificación (Cuenta/Habitación/Huésped/Fecha
    de emisión — cada una tomada del primer valor no vacío de su columna
    dentro del bloque, ya que no siempre están en la fila del folio),
    desglose de forma de pago (placeholder, aún no identificado) y las
    mismas 8 métricas calculadas.
    """
    tabla2 = []
    for bloque in bloques:
        filas = bloque["filas"]
        efectivo = amex = visa = mcard = credito = 0.0  # Placeholder: se llena al conciliar.
        tabla2.append(
            {
                "cuenta": bloque["consecutivo"],
                "habitacion": _primer_valor_no_vacio(filas, "F"),
                "huesped": _primer_valor_no_vacio(filas, "L"),
                "fecha_emision": _primer_valor_no_vacio(filas, "N"),
                "efectivo": efectivo,
                "amex": amex,
                "visa": visa,
                "mcard": mcard,
                "credito": credito,
                "totales": efectivo + amex + visa + mcard + credito,
                **calcular_metricas_cuenta(bloque),
            }
        )
    return tabla2


# ── Conciliación de forma de pago (Archivo de Movimiento vs Tabla 2) ─

FORMAS_PAGO_RECONOCIDAS = ("efectivo", "amex", "visa", "mcard", "credito")


def _es_fila_titulo_movimiento(fila: dict) -> bool:
    """
    Fila de título de sección de forma de pago (ej. "MN. EFECTIVO", "VISA",
    "MASTER CARD"): trae texto en C y H/M/Y vacíos. Distinta de la fila de
    encabezado real del archivo ("# Hab. Huésped Monto", que sí trae texto
    en H/M/Y) y de las filas de subtotal (que sí traen Y).
    """
    c = str(fila.get("C") or "").strip()
    h = str(fila.get("H") or "").strip()
    m = str(fila.get("M") or "").strip()
    y = str(fila.get("Y") or "").strip()
    return c != "" and h == "" and m == "" and y == ""


def _es_fila_dato_movimiento(fila: dict) -> bool:
    """
    Fila de dato real (un pago de un huésped): H no vacío, H no es la
    etiqueta de encabezado "Hab.", y Y parsea como número. Esto descarta de
    un solo golpe la fila de encabezado ("# Hab. Huésped Monto"), las filas
    de subtotal (H vacío) y la fila final "Hora / Monto".
    """
    h = str(fila.get("H") or "").strip()
    if h == "" or h.upper() == "HAB.":
        return False
    y = fila.get("Y")
    if y is None or str(y).strip() == "":
        return False
    # _parse_monto nunca truena (regresa 0.0 en vez de fallar); confirmamos
    # aquí que el texto realmente se parece a un número antes de aceptarlo,
    # para no tomar filas de texto suelto (ej. "Monto") como si fueran 0.0.
    texto = str(y).strip().replace("$", "").replace(",", "")
    try:
        float(texto)
        return True
    except ValueError:
        return False


def _forma_pago_desde_titulo(texto: str) -> str | None:
    """
    Clasifica el título de una sección del archivo de Movimiento a una de
    las formas de pago reconocidas en Tabla 2 (ver FORMAS_PAGO_RECONOCIDAS),
    por palabra clave. Ninguna coincidencia -> None (sección no reconocida,
    se ignora).
    """
    texto_norm = texto.strip().upper()
    if "EFECTIVO" in texto_norm:
        return "efectivo"
    if "AMEX" in texto_norm or "AMERICAN EXPRESS" in texto_norm:
        return "amex"
    if "VISA" in texto_norm:
        return "visa"
    if "MASTER" in texto_norm:
        return "mcard"
    if "CREDITO" in texto_norm:
        return "credito"
    return None


def agrupar_movimiento_por_forma_pago(filas: list[dict]) -> dict[str, list[dict]]:
    """
    Recorre las filas planas del archivo de Movimiento llevando la sección
    (forma de pago) actual, actualizada en cada fila título. Cada fila de
    dato se agrega a la lista de su sección actual. Filas título no
    reconocidas, o filas de dato antes de la primera sección reconocida, se
    ignoran.
    """
    resultado: dict[str, list[dict]] = {forma: [] for forma in FORMAS_PAGO_RECONOCIDAS}
    forma_actual: str | None = None

    for fila in filas:
        if _es_fila_titulo_movimiento(fila):
            forma_actual = _forma_pago_desde_titulo(str(fila.get("C") or ""))
            continue
        if forma_actual is not None and _es_fila_dato_movimiento(fila):
            resultado[forma_actual].append(
                {
                    "habitacion": str(fila.get("H") or "").strip(),
                    "huesped": str(fila.get("M") or "").strip(),
                    "monto": _parse_monto(fila.get("Y")),
                }
            )

    return resultado


def _cancelar_pares_duplicados(filas_forma_pago: list[dict]) -> list[dict]:
    """
    Agrupa por (habitacion, huesped, |monto|) y cancela pares +/- dentro de
    cada grupo (ej. dos cargos de +2,595.10 y uno de -2,595.10 -> queda
    solo un +2,595.10, el reverso canceló uno de los duplicados). Regresa
    las filas que sobran sin pareja, en su orden original.
    """
    grupos: dict[tuple[str, str, float], list[dict]] = {}
    for fila in filas_forma_pago:
        clave = (fila["habitacion"].upper(), fila["huesped"].upper(), round(abs(fila["monto"]), 2))
        grupos.setdefault(clave, []).append(fila)

    sobrantes: list[dict] = []
    for grupo in grupos.values():
        positivos = [f for f in grupo if f["monto"] >= 0]
        negativos = [f for f in grupo if f["monto"] < 0]
        num_cancelar = min(len(positivos), len(negativos))
        sobrantes.extend(positivos[num_cancelar:])
        sobrantes.extend(negativos[num_cancelar:])
    return sobrantes


TOLERANCIA_CONCILIACION = 0.05


def _buscar_cuenta_tabla2(
    habitacion: str, huesped_movimiento: str, monto: float, tabla2: list[dict]
) -> dict | None:
    """
    Busca en tabla2 la(s) cuenta(s) cuya habitación coincida exactamente
    (trim/mayúsculas) y cuyo huésped empiece con `huesped_movimiento` (los
    nombres no vienen idénticos entre archivos).

    Una misma habitación puede repetirse entre varias cuentas/reservas
    distintas con el mismo apellido (vistas en archivos reales: 3 folios
    separados para "D601 PIERCY", cada uno con su propio Total exacto).
    Cuando hay más de un candidato, se prefiere aquel cuyo acumulado actual
    más este `monto` cuadre (dentro de tolerancia) con su propio `total` —
    así cada pago se empareja con la cuenta correcta en vez de apilarse
    todo sobre la primera que aparezca.
    """
    habitacion_norm = habitacion.strip().upper()
    huesped_norm = huesped_movimiento.strip().upper()
    candidatos = [
        cuenta
        for cuenta in tabla2
        if str(cuenta.get("habitacion") or "").strip().upper() == habitacion_norm
        and str(cuenta.get("huesped") or "").strip().upper().startswith(huesped_norm)
    ]
    if not candidatos:
        return None
    if len(candidatos) == 1:
        return candidatos[0]

    def _acumulado(cuenta: dict) -> float:
        return sum(cuenta.get(forma, 0.0) for forma in FORMAS_PAGO_RECONOCIDAS)

    for cuenta in candidatos:
        if abs((_acumulado(cuenta) + monto) - cuenta.get("total", 0.0)) <= TOLERANCIA_CONCILIACION:
            return cuenta
    for cuenta in candidatos:
        if _acumulado(cuenta) < cuenta.get("total", 0.0) - TOLERANCIA_CONCILIACION:
            return cuenta
    return candidatos[0]


def _reclasificar_credito_como_amex(
    filas_credito: list[dict], filas_amex_periodo: list[dict]
) -> tuple[list[dict], list[dict], list[dict]]:
    """
    El archivo de Movimiento agrupa todo lo que no es Efectivo/Visa/Master
    Card bajo "CREDITO", sin distinguir Crédito real de American Express.
    Compara cada fila ya clasificada como "credito" contra las filas del
    segmento Amex del archivo de Movimientos por Periodo (Huésped por
    prefijo en cualquier dirección, ya que el formato del nombre no es
    idéntico entre archivos, + Monto exacto dentro de tolerancia) — cada
    match se consume (no se reusa dos veces) y esa fila pasa a ser Amex.

    Regresa (credito_restante, reclasificadas_a_amex, amex_sin_match).
    """
    disponibles = list(filas_amex_periodo)
    credito_restante: list[dict] = []
    reclasificadas: list[dict] = []

    for fila in filas_credito:
        huesped_norm = fila["huesped"].strip().upper()
        indice_match = None
        for i, amex in enumerate(disponibles):
            amex_huesped_norm = amex["huesped"].strip().upper()
            coincide_nombre = huesped_norm.startswith(amex_huesped_norm) or amex_huesped_norm.startswith(
                huesped_norm
            )
            coincide_monto = abs(abs(fila["monto"]) - amex["monto"]) <= TOLERANCIA_CONCILIACION
            if coincide_nombre and coincide_monto:
                indice_match = i
                break
        if indice_match is not None:
            disponibles.pop(indice_match)
            reclasificadas.append(fila)
        else:
            credito_restante.append(fila)

    return credito_restante, reclasificadas, disponibles


def conciliar_forma_pago_pagos_recepcion(
    tabla2: list[dict], movimiento_filas: list[dict], filas_amex_periodo: list[dict] | None = None
) -> dict:
    """
    Cruza el archivo de Movimiento contra Tabla 2: por cada forma de pago,
    cancela duplicados/reversos y suma cada monto restante en la columna
    correspondiente (efectivo/amex/visa/mcard/credito) de la cuenta que le
    corresponda por Habitación + prefijo de Huésped, recalculando
    "totales". Regresa `{"tabla2": <copia actualizada>, "no_conciliados":
    [...]}` (montos de Movimiento que no encontraron ninguna cuenta).

    Si se pasa `filas_amex_periodo` (del archivo opcional de Movimientos
    por Periodo), primero reclasifica de "credito" a "amex" las filas que
    coincidan con ese segmento (ver `_reclasificar_credito_como_amex`)
    antes de hacer el resto del cruce contra Tabla 2 — el resultado agrega
    `"credito_reclasificado_a_amex"` (conteo) y `"amex_sin_match"` (filas
    del periodo que no encontraron ningún crédito correspondiente).

    Una habitación puede reutilizarse entre varias reservas/huéspedes con
    el mismo apellido durante el periodo del archivo, así que Habitación +
    prefijo de Huésped a veces calza con varias filas de Movimiento — el
    sistema no tiene forma de saber si esas filas de más sí le pertenecen a
    la misma cuenta (ej. una estadía larga con varios pagos) o son de otra
    reserva no incluida en esta Tabla 2 (visto en archivos reales de
    ambos tipos, sin una señal confiable para distinguirlos: el archivo de
    Movimiento no trae número de reserva). Por eso aquí se suma todo lo que
    empate, sin rechazar nada — el frontend marca visualmente cuando
    "Totales" no cuadra con el "Total" de la cuenta, para que el usuario
    decida caso por caso y corrija a mano si hace falta.
    """
    tabla2_actualizada = [dict(cuenta) for cuenta in tabla2]
    no_conciliados: list[dict] = []

    movimiento_por_forma = agrupar_movimiento_por_forma_pago(movimiento_filas)

    amex_sin_match: list[dict] = []
    credito_reclasificado_a_amex = 0
    if filas_amex_periodo:
        credito_restante, reclasificadas, amex_sin_match = _reclasificar_credito_como_amex(
            movimiento_por_forma["credito"], filas_amex_periodo
        )
        movimiento_por_forma["credito"] = credito_restante
        movimiento_por_forma["amex"] = movimiento_por_forma["amex"] + reclasificadas
        credito_reclasificado_a_amex = len(reclasificadas)

    for forma_pago, filas_forma_pago in movimiento_por_forma.items():
        for fila in _cancelar_pares_duplicados(filas_forma_pago):
            cuenta = _buscar_cuenta_tabla2(
                fila["habitacion"], fila["huesped"], fila["monto"], tabla2_actualizada
            )
            if cuenta is None:
                no_conciliados.append({"forma_pago": forma_pago, **fila})
                continue
            cuenta[forma_pago] = cuenta.get(forma_pago, 0.0) + fila["monto"]

    for cuenta in tabla2_actualizada:
        cuenta["totales"] = sum(cuenta.get(forma, 0.0) for forma in FORMAS_PAGO_RECONOCIDAS)

    return {
        "tabla2": tabla2_actualizada,
        "no_conciliados": no_conciliados,
        "amex_sin_match": amex_sin_match,
        "credito_reclasificado_a_amex": credito_reclasificado_a_amex,
    }


def filtrar_cuentas_conciliadas(tabla2: list[dict]) -> list[dict]:
    """
    Regresa solo las cuentas de Tabla 2 que sí encontraron pago en el
    archivo de Movimiento (Totales > 0 tras conciliar) — el "resultado
    limpio" de un día, para guardar en Pagos del Mes.
    """
    return [cuenta for cuenta in tabla2 if cuenta.get("totales", 0.0) > 0]


_COLUMNAS_REPORTE_PAGOS_DEL_MES: list[tuple[str, str]] = [
    ("fecha", "Fecha"),
    ("cuenta", "Cuenta"),
    ("habitacion", "Habitación"),
    ("huesped", "Huésped"),
    ("efectivo", "Efectivo"),
    ("amex", "Amex"),
    ("visa", "Visa"),
    ("mcard", "M.Card"),
    ("credito", "Crédito"),
    ("totales", "Totales"),
    ("subtotal1", "SUBTOTAL1"),
    ("iva16", "IVA 16%"),
    ("ish3", "ISH 3%"),
    ("subtotal2", "SUBTOTAL2"),
    ("gravado", "GRAVADO"),
    ("tasa0", "TASA 0%"),
    ("propina", "PROPINA"),
    ("total", "TOTAL"),
]
_PRIMERA_COLUMNA_NUMERICA = 5  # "Efectivo"
_ULTIMA_COLUMNA = len(_COLUMNAS_REPORTE_PAGOS_DEL_MES)

_FILL_GRUPO = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_FILL_SUBTOTAL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_FILL_TOTAL_GENERAL = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
_FONT_NEGRITA = Font(bold=True)


def generar_reporte_pagos_del_mes_excel(filas: list[dict]) -> bytes:
    """
    Genera un .xlsx con las cuentas de Pagos del Mes (mismas columnas que
    la tabla "Pagos del Mes" del frontend), agrupadas por fecha: cada día
    lleva una fila de subtotal al final de su grupo y, al cierre del
    archivo, una fila "TOTAL GENERAL" con la suma de todo — así el reporte
    ya deja claro qué días incluye sin anotarlo aparte.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Pagos del Mes"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.cell(row=1, column=1, value="Identificación")
    ws.merge_cells(start_row=1, start_column=5, end_row=1, end_column=9)
    ws.cell(row=1, column=5, value="Forma de Pago")
    ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=_ULTIMA_COLUMNA)
    ws.cell(row=1, column=10, value="Cálculos")
    for columna in (1, 5, 10):
        celda = ws.cell(row=1, column=columna)
        celda.font = _FONT_NEGRITA
        celda.alignment = Alignment(horizontal="center")
        celda.fill = _FILL_GRUPO

    for indice, (_clave, etiqueta) in enumerate(_COLUMNAS_REPORTE_PAGOS_DEL_MES, start=1):
        celda = ws.cell(row=2, column=indice, value=etiqueta)
        celda.font = _FONT_NEGRITA
        celda.fill = _FILL_GRUPO

    filas_por_fecha: dict[str, list[dict]] = {}
    for fila in filas:
        filas_por_fecha.setdefault(fila.get("fecha", ""), []).append(fila)

    fila_actual = 3
    totales_generales = [0.0] * (_ULTIMA_COLUMNA - _PRIMERA_COLUMNA_NUMERICA + 1)

    for fecha in sorted(filas_por_fecha):
        filas_dia = filas_por_fecha[fecha]
        subtotales_dia = [0.0] * (_ULTIMA_COLUMNA - _PRIMERA_COLUMNA_NUMERICA + 1)

        for fila in filas_dia:
            for indice, (clave, _etiqueta) in enumerate(_COLUMNAS_REPORTE_PAGOS_DEL_MES, start=1):
                ws.cell(row=fila_actual, column=indice, value=fila.get(clave))
                if indice >= _PRIMERA_COLUMNA_NUMERICA:
                    valor = float(fila.get(clave) or 0.0)
                    posicion = indice - _PRIMERA_COLUMNA_NUMERICA
                    subtotales_dia[posicion] += valor
                    totales_generales[posicion] += valor
            fila_actual += 1

        ws.cell(row=fila_actual, column=1, value=f"Subtotal {fecha}").font = _FONT_NEGRITA
        for indice in range(_PRIMERA_COLUMNA_NUMERICA, _ULTIMA_COLUMNA + 1):
            posicion = indice - _PRIMERA_COLUMNA_NUMERICA
            celda = ws.cell(row=fila_actual, column=indice, value=round(subtotales_dia[posicion], 2))
            celda.font = _FONT_NEGRITA
        for indice in range(1, _ULTIMA_COLUMNA + 1):
            ws.cell(row=fila_actual, column=indice).fill = _FILL_SUBTOTAL
        fila_actual += 1

    ws.cell(row=fila_actual, column=1, value="TOTAL GENERAL").font = _FONT_NEGRITA
    for indice in range(_PRIMERA_COLUMNA_NUMERICA, _ULTIMA_COLUMNA + 1):
        posicion = indice - _PRIMERA_COLUMNA_NUMERICA
        celda = ws.cell(row=fila_actual, column=indice, value=round(totales_generales[posicion], 2))
        celda.font = _FONT_NEGRITA
    for indice in range(1, _ULTIMA_COLUMNA + 1):
        ws.cell(row=fila_actual, column=indice).fill = _FILL_TOTAL_GENERAL

    for indice in range(1, _ULTIMA_COLUMNA + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 14

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
