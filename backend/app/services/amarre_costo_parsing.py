"""
Amarre de Costo de Ventas — parsing de los 12 reportes crudos que Novohit
exporta cada mes (balanza, 4 conteos físicos de almacén, 2 reportes de
salidas/despachos, 5 de entradas/compras).

Decisión de diseño clave (ver plan): estos reportes son jerárquicos y de
longitud variable — cuántas filas ocupan cambia cada mes según cuántos
artículos/movimientos hubo. El libro maestro que el usuario arma a mano hoy
referencia números de fila fijos (`=AD1494+AD1794+...`), que hay que
reajustar cada mes — exactamente el trabajo manual que este módulo elimina.
Por eso NINGÚN parser de este archivo usa un número de fila fijo: todos
ubican sus totales buscando el TEXTO de la etiqueta (confirmado contra
archivos reales en `INVETARIOS/`), sin importar en qué fila caiga.

Los reportes de conteo físico (`INV030`) tienen 3 niveles de indentación:
categoría → subcategoría → artículos, cada nivel cerrado por su propia fila
"TOTAL <nombre>". La fila de categoría (nivel 1, la que interesa) SIEMPRE
empieza sin espacio-no-rompible (`\xa0`); las de subcategoría casi siempre
lo traen al inicio/final (confirmado con datos reales de agosto 2026) — eso
es lo único que distingue un nivel de otro, el nombre solo no alcanza
("TOTAL ALIMENTOS" existe como wrapper de nivel 1 en Almacén General, pero
"TOTAL ALIMENTOS TASA 0%" cerca del final del mismo archivo es una fila de
verificación/checksum, no una categoría real — hay que confirmar cada mapeo
contra los valores ya conocidos del libro maestro, ver Amarre COSTO real).
"""

from io import BytesIO

import xlrd
from openpyxl import load_workbook

Fila = list
Grid = list[Fila]


# ── Lectura cruda (grid uniforme, .xlsx vía openpyxl, .xls vía xlrd) ────


def leer_filas_crudas(contenido: bytes, nombre_archivo: str) -> Grid:
    """
    Lee la primera hoja del archivo como una matriz de filas (0-indexed),
    sin encabezados ni tipos especiales — celdas vacías quedan como "".
    Detecta el formato por extensión: `.xls` viejo requiere `xlrd`
    (confirmado con archivos reales: `openpyxl.load_workbook` falla en
    ellos con `InvalidFileException`), `.xlsx` usa `openpyxl`.
    """
    if nombre_archivo.lower().endswith(".xls"):
        libro = xlrd.open_workbook(file_contents=contenido)
        hoja = libro.sheet_by_index(0)
        return [[hoja.cell_value(fila, col) for col in range(hoja.ncols)] for fila in range(hoja.nrows)]

    wb = load_workbook(BytesIO(contenido), data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas: Grid = []
    for fila_celdas in ws.iter_rows(values_only=True):
        filas.append(["" if v is None else v for v in fila_celdas])
    return filas


def _texto(valor) -> str:
    return valor if isinstance(valor, str) else ""


def _a_numero(valor) -> float:
    """
    Convierte una celda (float, o texto tipo "1,234.56" / "$1,234.56") a
    float. None/vacío/no numérico -> 0.0. Mismo criterio que
    `pagos_recepcion.py::_parse_monto`.
    """
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("$", "").replace(",", "")
    if texto == "":
        return 0.0
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _es_total_nivel1(texto_crudo: str) -> bool:
    """
    Una fila "TOTAL X" de nivel 1 (categoría) nunca trae espacio-no-rompible
    al inicio; las de subcategoría (nivel 2/3) sí — confirmado con datos
    reales: `'\\xa0TOTAL CIGARROS\\xa0'` (subcategoría) vs `'TOTAL TABAQUERIA'`
    (categoría).

    IMPORTANTE: `texto_crudo` debe venir SIN `.strip()` aplicado. Python
    considera `\xa0` como espacio en blanco, así que `.strip()` lo borra
    igual que un espacio normal — si se llama con texto ya "stripeado", esta
    función nunca detecta subcategorías y termina sumando de más (bug real
    confirmado: "TOTAL EQUIPO DE OPERACION" aparece 7 veces en
    `almacen general agosto 2026.xls`, 6 de ellas subcategoría — sin este
    chequeo intacto se sumaban las 7).
    """
    return texto_crudo.startswith("TOTAL ") and not texto_crudo.startswith("\xa0")


def _es_celda_numerica(valor) -> bool:
    """
    True para números reales o texto que representa un número (ej.
    "16,189.85") — confirmado con datos reales: los reportes Novohit
    guardan el monto de artículos individuales como float, pero el de las
    filas "TOTAL X"/"Sub Total:" como texto con comas.
    """
    if isinstance(valor, (int, float)):
        return True
    if isinstance(valor, str):
        texto = valor.strip().replace("$", "").replace(",", "")
        if texto == "":
            return False
        try:
            float(texto)
            return True
        except ValueError:
            return False
    return False


def _numeros_en_fila(fila: Fila) -> list[float]:
    return [_a_numero(v) for v in fila if _es_celda_numerica(v)]


def _valor_numerico_en_fila(fila: Fila, ultimo: bool = True) -> float:
    """Primer (o último) valor numérico de la fila (texto con comas incluido)."""
    numeros = _numeros_en_fila(fila)
    if not numeros:
        return 0.0
    return numeros[-1] if ultimo else numeros[0]


def _ultimos_n_numeros_en_fila(fila: Fila, n: int) -> list[float]:
    """Los últimos `n` valores numéricos de la fila, en orden de aparición."""
    numeros = _numeros_en_fila(fila)
    return numeros[-n:] if len(numeros) >= n else numeros + [0.0] * (n - len(numeros))


# ── Variantes "por coordenada" (fila, columna) ──────────────────────────
# Mismas reglas de localización que arriba (reusan _texto/_es_total_nivel1/
# _es_celda_numerica, no las duplican), pero regresan DÓNDE está el dato en
# vez de su valor — usadas por `amarre_costo_workbook.py` para escribir
# fórmulas de Excel que apunten a esa celda exacta (recalculan solas si el
# usuario edita el archivo pegado), en vez de un número ya sumado por
# Python. Las funciones "por valor" de arriba siguen siendo las que usa la
# vista web (`amarre_costo_reportes.py`) — no se tocan.


def _indices_numericos_en_fila(fila: Fila) -> list[int]:
    return [i for i, v in enumerate(fila) if _es_celda_numerica(v)]


def _columna_numerica_en_fila(fila: Fila, ultimo: bool = True) -> int | None:
    indices = _indices_numericos_en_fila(fila)
    if not indices:
        return None
    return indices[-1] if ultimo else indices[0]


def _ultimas_n_columnas_en_fila(fila: Fila, n: int) -> list[int | None]:
    indices = _indices_numericos_en_fila(fila)
    if len(indices) >= n:
        return indices[-n:]
    return indices + [None] * (n - len(indices))


# ── Balanza (1BALANZA / 1balanza.xlsx) ──────────────────────────────────


def parsear_balanza(contenido: bytes, nombre_archivo: str = "1balanza.xlsx") -> dict[str, dict]:
    """
    Código de cuenta (columna A, tal cual con puntos) -> {descripcion,
    saldo_anterior, saldo_actual}. Header en la fila que empieza con
    "Código Cuenta" (confirmado fila 4, pero se ubica por texto, no por
    número fijo); cuentas con saldo $0 en el mes simplemente no aparecen en
    el reporte — el dict resultante no las trae, y cualquier consulta
    posterior debe manejar la ausencia sin lanzar excepción.
    """
    filas = leer_filas_crudas(contenido, nombre_archivo)
    fila_header = None
    for i, fila in enumerate(filas):
        if _texto(fila[0]).strip() == "Código Cuenta":
            fila_header = i
            break
    if fila_header is None:
        raise ValueError("No se encontró el encabezado 'Código Cuenta' en el archivo de balanza")

    resultado: dict[str, dict] = {}
    for fila in filas[fila_header + 1 :]:
        codigo = _texto(fila[0]).strip()
        if codigo == "":
            continue
        resultado[codigo] = {
            "descripcion": _texto(fila[1]).strip(),
            "saldo_anterior": _a_numero(fila[2]),
            "saldo_actual": _a_numero(fila[6]),
        }
    return resultado


def buscar_en_balanza(codigo: str, balanza: dict[str, dict], campo: str, respaldo: float = 0.0) -> float:
    """IFERROR(VLOOKUP(...), respaldo) — nunca lanza si el código no está (cuenta en $0)."""
    cuenta = balanza.get(codigo)
    if cuenta is None:
        return respaldo
    return cuenta.get(campo, respaldo)


# ── Conteo físico de almacén (INV030) ───────────────────────────────────

# Categorías objetivo por almacén: nombre EXACTO de la etiqueta "TOTAL
# <categoria>" tal como aparece en el reporte Novohit (confirmado contra
# `INVETARIOS/*.xls` reales de agosto 2026) -> nombre del campo del modelo.
# Alimentos/Bebidas quedan marcados aparte (ver `_sumar_categorias_compuestas`)
# porque el libro maestro real las arma sumando varias líneas dispersas, no
# un solo total limpio — deben calibrarse contra los valores ya conocidos
# del libro maestro antes de confiar en ellas en producción.
CATEGORIAS_ALM_GRAL: dict[str, str] = {
    "TABAQUERIA": "tabaqueria",
    "PAPELERIA": "papeleria",
    "STOCK MTO.": "stock_mantenimiento",
    # "Ropa y Uniformes" (la cuenta contable) en realidad se alimenta de la
    # etiqueta Novohit "SOUVENIRS  Y CAFE EN GRANO" (doble espacio literal,
    # confirmado carácter por carácter contra el archivo real — el nombre
    # de la subcategoría en Novohit no corresponde al de la cuenta
    # contable, pero es la celda que el libro maestro real referencia:
    # `AD2180=AD2162`, 27,003.39 exacto). NO usar "ROPA  ALMACEN": esa
    # etiqueta sí existe en el reporte pero es un total distinto (126,191.58,
    # confirmado que no es la fuente real de esta cuenta).
    "SOUVENIRS  Y CAFE EN GRANO": "ropa_uniformes",
    # "Material para Mantto" se alimenta de "TOTAL EQUIPO DE OPERACION"
    # (confirmado por el usuario, y verificado exacto contra el real:
    # $9,603.45 — la etiqueta aparece 7 veces en el archivo, 6 son
    # subcategoría con espacio-no-rompible y no cuentan, solo 1 es la
    # categoría real; ver el fix de `_es_total_nivel1` más arriba).
    "EQUIPO DE OPERACION": "material_mantto",
    # "MEDICINAS Y MEDICAMENTOS" y "SERVICIOS": la cuenta contable está en
    # $0 el mes validado (agosto 2026) y el reporte no trae una fila propia
    # para ninguna — quedan en 0.0 por defecto, consistente con el dato real.
}
CATEGORIAS_COMPUESTAS_ALM_GRAL: dict[str, tuple[str, ...]] = {
    # Cada tupla es la lista de nombres de categoría (SIN el prefijo
    # "TOTAL ", igual que las claves que arma `_totales_nivel1`) a sumar
    # para ese campo. Confirmado en archivo real: la fila "TOTAL
    # ALIMENTOS"/"TOTAL BEBIDAS" existe como wrapper limpio y suma todas
    # sus ocurrencias correctamente.
    "alimentos": ("ALIMENTOS",),
    "bebidas": ("BEBIDAS",),
    # "Higiénicos y Suministros" = TOTAL SUMINISTROS + TOTAL ARTESANIAS
    # (confirmado por el usuario, verificado exacto contra el real:
    # $471,912.35 + $2,795.60 = $474,707.95).
    "higienicos": ("SUMINISTROS", "ARTESANIAS"),
}

CATEGORIAS_TIENDA: dict[str, str] = {
    "TABAQUERIA": "tabaqueria",
    "MEDICINAS Y ARTS DE CURACION": "medicinas",
    "ROPA": "ropa",
    # "Higiénicos y Sum" (la cuenta contable) se alimenta de la etiqueta
    # Novohit "SUMINISTROS" (única en el archivo, confirmado: `AD3999=AD1755`,
    # 12,438.05 exacto) — no de un nombre que mencione "higiénicos".
    "SUMINISTROS": "higienicos",
    # "ARTESANIAS" a secas NO se usa: ese texto aparece más de una vez en
    # el reporte (una ocurrencia fuera de la sección que realmente
    # pertenece a esta cuenta), así que sumar "todas las ocurrencias"
    # suma de más (39,752.68 vs. el real 24,969.34). La etiqueta específica
    # "ARTESANIAS DIVERSAS GRAVADOS" sí es única y reproduce el monto real
    # exacto (confirmado: `AD4002=AD2013`, misma cifra).
    "ARTESANIAS DIVERSAS GRAVADOS": "artesanias",
}
CATEGORIAS_COMPUESTAS_TIENDA: dict[str, tuple[str, ...]] = {
    "alimentos": ("ALIMENTOS",),
    "bebidas": ("BEBIDAS",),
}

CATEGORIAS_REST_BAR: dict[str, str] = {
    "ALIMENTOS": "alimentos",
    "BEBIDAS": "bebidas",
}

CATEGORIAS_CAFETERIA: dict[str, str] = {
    "ROPA": "souvenirs",  # el rubro "Souvenirs" del amarre se registra como "Ropa" en el reporte Novohit
}
CATEGORIAS_COMPUESTAS_CAFETERIA: dict[str, tuple[str, ...]] = {
    # Cafetería no tiene cuenta contable propia de "Bebidas" — ese monto
    # se registra dentro de "Alimentos" (confirmado contra el libro real:
    # `AD426=AD321+AD391+AD417+0.01`, donde AD391 es justamente la única
    # ocurrencia de "TOTAL BEBIDAS" del reporte). Sumar todas las
    # ocurrencias de ambas etiquetas + el ajuste de redondeo del archivo
    # original reproduce el monto real exacto: 19,034.90 + 1,954.56 + 0.01
    # = 20,989.47.
    "alimentos": ("ALIMENTOS", "BEBIDAS"),
}
AJUSTE_REDONDEO_CAFETERIA_ALIMENTOS = 0.01


def _totales_nivel1(filas: Grid) -> dict[str, float]:
    """
    {"TABAQUERIA": monto, "ALIMENTOS": monto, ...} para toda fila 'TOTAL X'
    de nivel 1 — SUMA todas las ocurrencias de una misma categoría, no solo
    la última. Confirmado con datos reales: "Alimentos"/"Bebidas" aparecen
    repetidas varias veces en el mismo archivo (distintas áreas/menús del
    almacén, cada una con su propia fila "TOTAL ALIMENTOS"/"TOTAL BEBIDAS")
    y el libro maestro real las suma todas (`=AD684+AD966+AD1272`, 3 celdas
    para "Alimentos" de Restaurant-Bar) — quedarse solo con la última
    ocurrencia perdía las otras dos.
    """
    resultado: dict[str, float] = {}
    for fila in filas:
        for valor in fila:
            texto_crudo = _texto(valor)
            if _es_total_nivel1(texto_crudo):
                nombre = texto_crudo.strip()[len("TOTAL ") :].strip()
                resultado[nombre] = resultado.get(nombre, 0.0) + _valor_numerico_en_fila(fila)
                break
    return resultado


def parsear_conteo_fisico(
    contenido: bytes,
    nombre_archivo: str,
    categorias_simples: dict[str, str],
    categorias_compuestas: dict[str, tuple[str, ...]] | None = None,
) -> dict[str, float]:
    """
    Regresa {campo_modelo: monto} para un conteo físico (INV030), ubicando
    cada categoría por el texto de su fila "TOTAL <categoria>" (nivel 1),
    sin importar en qué fila del archivo caiga ese mes.
    """
    filas = leer_filas_crudas(contenido, nombre_archivo)
    totales = _totales_nivel1(filas)

    resultado: dict[str, float] = {}
    for etiqueta, campo in categorias_simples.items():
        resultado[campo] = totales.get(etiqueta, 0.0)
    for campo, etiquetas in (categorias_compuestas or {}).items():
        resultado[campo] = sum(totales.get(etiqueta, 0.0) for etiqueta in etiquetas)
    return resultado


# ── Salidas / despachos (INV110) ────────────────────────────────────────

# Encabezado real de cada sección de departamento en "SAL ALMACEN GRAL"
# (columna C), confirmado contra `salidas de almacen general.xls`.
DEPARTAMENTOS_SALIDAS_ALM_GRAL: dict[str, str] = {
    "REST. ALIMENTOS": "entrysal_restaurant",
    "LA PALAPA BEBIDAS": "entrysal_palapa_bar",
    "LA TIENDA": "entrysal_tienda",
    "CAFETERIA ALIMENTOS": "entrysal_cafeteria_alimentos",
    "LA CARRETA ALIMENTOS": "entrysal_carreta_alimentos",
    "LA CARRETA BEBIDAS": "entrysal_carreta_bebidas",
    "CAFETERIA SOUVENIRS Y CAFE EN GRANO": "entrysal_souvenirs_cafe",
}


def parsear_salidas_almacen_general(contenido: bytes, nombre_archivo: str = "salidas de almacen general.xls") -> dict:
    """
    Regresa {"despacho_gasto", "despacho_centro_costo", <7 campos de
    departamento>} — ambos valores CRUDOS de sus etiquetas ("Total General
    DESPACHOS a Centros de Costos y Gastos (Mermas y Otros):" y "Total de
    DESPACHOS a CENTRO DE COSTO:"). OJO: el "Despacho a Gasto" que muestra
    2DATOS ENTRYSAL!D8 NO es `despacho_gasto` solo — es
    `despacho_gasto - despacho_centro_costo` (confirmado contra el libro
    real: `D8=VALUE(U1645)-VALUE(U1644)`, valores agosto 2026: 1,708,188.40
    − 761,863.73 = 946,324.67). Quien consuma este resultado debe hacer esa
    resta explícitamente — `despacho_gasto` aquí es la etiqueta cruda, no
    el D8 final.

    Cada departamento se ubica buscando su encabezado (columna C) y tomando
    el ÚLTIMO "Sub Total:" que aparece antes del siguiente encabezado de
    departamento (o del final de la sección) — el reporte anida varios
    sub-totales por rubro de gasto dentro de cada bloque de departamento; el
    último de ellos es el acumulado real del bloque completo (confirmado:
    coincide exactamente con los valores ya conocidos de 2DATOS ENTRYSAL).
    """
    filas = leer_filas_crudas(contenido, nombre_archivo)

    despacho_gasto = 0.0
    despacho_centro_costo = 0.0
    for fila in filas:
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip()
        if etiqueta.startswith("Total General DESPACHOS a Centros de Costos y Gastos"):
            despacho_gasto = _valor_numerico_en_fila(fila)
        elif etiqueta.startswith("Total de DESPACHOS a CENTRO DE COSTO"):
            despacho_centro_costo = _valor_numerico_en_fila(fila)

    # Filas donde aparece un encabezado de departamento reconocido, en orden.
    filas_encabezado = [
        (i, DEPARTAMENTOS_SALIDAS_ALM_GRAL[_texto(fila[2]).strip()])
        for i, fila in enumerate(filas)
        if len(fila) > 2 and _texto(fila[2]).strip() in DEPARTAMENTOS_SALIDAS_ALM_GRAL
    ]

    resultado: dict = {"despacho_gasto": despacho_gasto, "despacho_centro_costo": despacho_centro_costo}
    for campo in DEPARTAMENTOS_SALIDAS_ALM_GRAL.values():
        resultado[campo] = 0.0

    for idx, (fila_inicio, campo) in enumerate(filas_encabezado):
        fila_fin = filas_encabezado[idx + 1][0] if idx + 1 < len(filas_encabezado) else len(filas)
        ultimo_subtotal = 0.0
        for fila in filas[fila_inicio:fila_fin]:
            if any("SUB TOTAL" in _texto(v).upper() for v in fila):
                ultimo_subtotal = _valor_numerico_en_fila(fila)
        resultado[campo] = ultimo_subtotal

    return resultado


def parsear_salidas_mantenimiento(contenido: bytes, nombre_archivo: str = "salidas de mantenimiento.xls") -> float:
    """Despacho a Gasto de Mantenimiento = 2DATOS ENTRYSAL D9."""
    filas = leer_filas_crudas(contenido, nombre_archivo)
    for fila in filas:
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip()
        if etiqueta.startswith("Total General DESPACHOS a Centros de Costos y Gastos"):
            return _valor_numerico_en_fila(fila)
    return 0.0


# ── Entradas / compras (INV160) ─────────────────────────────────────────


def parsear_entradas(contenido: bytes, nombre_archivo: str) -> dict:
    """
    {"subtotal", "descuento", "iva", "total"} desde la fila cuya etiqueta es
    "TOTAL GENERAL DE ENTRADAS" — los 4 valores son las últimas 4 columnas
    numéricas de esa fila, en orden (Subtotal, Descuento, IVA, Total con
    IVA), confirmado contra `entradas de almacen general(para
    expottacion).xls`.
    """
    filas = leer_filas_crudas(contenido, nombre_archivo)
    for fila in filas:
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip().upper()
        if "TOTAL GENERAL DE ENTRADAS" in etiqueta:
            subtotal, descuento, iva, total = _ultimos_n_numeros_en_fila(fila, 4)
            return {"subtotal": subtotal, "descuento": descuento, "iva": iva, "total": total}
    return {"subtotal": 0.0, "descuento": 0.0, "iva": 0.0, "total": 0.0}


# ── Localización por coordenada (para `amarre_costo_workbook.py`) ──────

Celda = tuple[int, int]  # (fila, columna), 0-indexed


def localizar_ocurrencias_nivel1(filas: Grid) -> dict[str, list[Celda]]:
    """
    Como `_totales_nivel1`, pero regresa la lista de celdas (fila, columna)
    de cada ocurrencia de "TOTAL <categoria>" de nivel 1, no la suma — para
    poder escribir `='Hoja'!AD38+'Hoja'!AD1495+...` en vez de un número ya
    sumado por Python.
    """
    resultado: dict[str, list[Celda]] = {}
    for i, fila in enumerate(filas):
        for valor in fila:
            texto_crudo = _texto(valor)
            if _es_total_nivel1(texto_crudo):
                nombre = texto_crudo.strip()[len("TOTAL ") :].strip()
                col = _columna_numerica_en_fila(fila)
                if col is not None:
                    resultado.setdefault(nombre, []).append((i, col))
                break
    return resultado


def localizar_codigo_categoria_nivel1(filas: Grid, columna_texto_categoria: int = 0) -> dict[str, Celda]:
    """
    Para cada categoría de nivel 1, regresa la celda (fila, columna) del
    CÓDIGO DE CUENTA (con guiones) más cercana arriba de su fila "TOTAL X"
    — el layout real de Novohit siempre trae el código de cuenta en una
    columna angosta a la izquierda de la cuadrícula de artículos (columna
    A/B), no en la fila del total. Si no se encuentra un código plausible,
    esa categoría queda ausente del dict (el llamador debe usar un nombre
    de respaldo literal, igual que `IFERROR` en el libro real).
    """
    import re

    patron_codigo = re.compile(r"^\d{4,}[-.]?\d*[-.]?\d*$")
    resultado: dict[str, Celda] = {}
    for i, fila in enumerate(filas):
        for valor in fila:
            texto_crudo = _texto(valor)
            if _es_total_nivel1(texto_crudo):
                nombre = texto_crudo.strip()[len("TOTAL ") :].strip()
                for j in range(i, max(-1, i - 40), -1):
                    candidato = _texto(filas[j][columna_texto_categoria] if columna_texto_categoria < len(filas[j]) else "").strip()
                    if patron_codigo.match(candidato):
                        resultado.setdefault(nombre, (j, columna_texto_categoria))
                        break
                break
    return resultado


def localizar_salidas_almacen_general(filas: Grid) -> dict:
    """
    Como `parsear_salidas_almacen_general`, pero regresa coordenadas de
    celda en vez de valores: `{"despacho_gasto": (fila,col), ...,
    "departamentos": {campo: (fila,col)}}`. `None` cuando no se encontró
    (huecos reales, ej. Cafetería Bebidas no tiene sección en este reporte).
    """
    despacho_gasto: Celda | None = None
    despacho_centro_costo: Celda | None = None
    for i, fila in enumerate(filas):
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip()
        if etiqueta.startswith("Total General DESPACHOS a Centros de Costos y Gastos"):
            col = _columna_numerica_en_fila(fila)
            if col is not None:
                despacho_gasto = (i, col)
        elif etiqueta.startswith("Total de DESPACHOS a CENTRO DE COSTO"):
            col = _columna_numerica_en_fila(fila)
            if col is not None:
                despacho_centro_costo = (i, col)

    filas_encabezado = [
        (i, DEPARTAMENTOS_SALIDAS_ALM_GRAL[_texto(fila[2]).strip()])
        for i, fila in enumerate(filas)
        if len(fila) > 2 and _texto(fila[2]).strip() in DEPARTAMENTOS_SALIDAS_ALM_GRAL
    ]

    departamentos: dict[str, Celda | None] = {campo: None for campo in DEPARTAMENTOS_SALIDAS_ALM_GRAL.values()}
    for idx, (fila_inicio, campo) in enumerate(filas_encabezado):
        fila_fin = filas_encabezado[idx + 1][0] if idx + 1 < len(filas_encabezado) else len(filas)
        ultima_celda: Celda | None = None
        for i in range(fila_inicio, fila_fin):
            fila = filas[i]
            if any("SUB TOTAL" in _texto(v).upper() for v in fila):
                col = _columna_numerica_en_fila(fila)
                if col is not None:
                    ultima_celda = (i, col)
        departamentos[campo] = ultima_celda

    return {"despacho_gasto": despacho_gasto, "despacho_centro_costo": despacho_centro_costo, "departamentos": departamentos}


def localizar_salidas_mantenimiento(filas: Grid) -> Celda | None:
    for i, fila in enumerate(filas):
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip()
        if etiqueta.startswith("Total General DESPACHOS a Centros de Costos y Gastos"):
            col = _columna_numerica_en_fila(fila)
            if col is not None:
                return (i, col)
    return None


def localizar_entradas(filas: Grid) -> dict[str, Celda | None] | None:
    """
    {"subtotal": (fila,col), "descuento": ..., "iva": ..., "total": ...}
    desde la fila "TOTAL GENERAL DE ENTRADAS", o None si no se encontró.
    """
    for i, fila in enumerate(filas):
        etiqueta = " ".join(_texto(v) for v in fila if _texto(v)).strip().upper()
        if "TOTAL GENERAL DE ENTRADAS" in etiqueta:
            col_subtotal, col_descuento, col_iva, col_total = _ultimas_n_columnas_en_fila(fila, 4)
            return {
                "subtotal": (i, col_subtotal) if col_subtotal is not None else None,
                "descuento": (i, col_descuento) if col_descuento is not None else None,
                "iva": (i, col_iva) if col_iva is not None else None,
                "total": (i, col_total) if col_total is not None else None,
            }
    return None


def localizar_balanza_fila(filas: Grid, codigo: str) -> int | None:
    """Fila (0-indexed) donde la columna A de 1BALANZA trae exactamente `codigo`, o None."""
    for i, fila in enumerate(filas):
        if _texto(fila[0]).strip() == codigo:
            return i
    return None
