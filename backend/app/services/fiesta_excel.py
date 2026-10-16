"""
Excel de respaldo de la Fiesta de Fin de Año: toda la información importante, explicada hoja por hoja, para poder
seguir trabajando si el sistema no está disponible ("segunda opción").

Es puro (sin base de datos): recibe diccionarios ya armados por el router y devuelve los bytes del .xlsx.
Los totales, saldos y acumulados son FÓRMULAS vivas: si se corrige un importe a mano, se recalculan.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MORADO = "5F3DC4"
GRIS = "F1F3F5"
DINERO = '"$"#,##0.00'
ENTERO = "#,##0"
FECHA = "dd/mm/yyyy"
PORCENTAJE = "0.0%"
_BORDE = Side(style="thin", color="CED4DA")

ETIQUETA_TIPO = {"adulto": "Adulto", "nino": "Niño", "cortesia": "Cortesía"}
ETIQUETA_ESTADO_SILLA = {
    "libre": "Libre", "sin_pago": "Sin pago", "anticipo": "Con anticipo", "pagada": "Pagada",
    "cortesia": "Cortesía", "bloqueada": "Bloqueada",
}
ETIQUETA_ESTADO_RESERVACION = {
    "sin_sillas": "Sin sillas", "cortesia": "Cortesía", "sin_pago": "Sin pago", "anticipo": "Con anticipo", "pagada": "Pagada",
}


def _fecha(valor: Any) -> Any:
    """Fecha de Excel (no texto) a partir de date/datetime/ISO; vacío si no hay."""
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        return str(valor)


def _hoja(libro: Workbook, titulo: str, que_es: str, columnas: list[tuple[str, float, str | None]], primera: bool = False) -> Worksheet:
    """Hoja con título, explicación y encabezados (fila 4). `columnas` = (nombre, ancho, formato de número)."""
    hoja = libro.active if primera else libro.create_sheet()
    hoja.title = titulo
    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=15, color=MORADO)
    hoja["A2"] = que_es
    hoja["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(columnas), 6))
    hoja.row_dimensions[2].height = 62
    for i, (nombre, ancho, _formato) in enumerate(columnas, start=1):
        c = hoja.cell(row=4, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MORADO)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(top=_BORDE, bottom=_BORDE, left=_BORDE, right=_BORDE)
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.row_dimensions[4].height = 34
    hoja.freeze_panes = "A5"
    return hoja


def _fila(hoja: Worksheet, fila: int, valores: list[Any], formatos: list[str | None]) -> None:
    for i, (v, f) in enumerate(zip(valores, formatos), start=1):
        c = hoja.cell(row=fila, column=i, value=v)
        if f:
            c.number_format = f
        c.border = Border(bottom=_BORDE)


def _total(hoja: Worksheet, fila: int, etiqueta: str, columnas_suma: dict[int, str], desde: int, hasta: int, ancho: int) -> None:
    """Renglón TOTAL con SUM de cada columna indicada (col → formato)."""
    hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
    for col, formato in columnas_suma.items():
        letra = get_column_letter(col)
        c = hoja.cell(row=fila, column=col, value=f"=SUM({letra}{desde}:{letra}{hasta})" if hasta >= desde else 0)
        c.number_format = formato
        c.font = Font(bold=True)
    for col in range(1, ancho + 1):
        hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=GRIS)


def construir_libro(datos: dict) -> bytes:
    evento = datos["evento"]
    resumen = datos["resumen"]
    cobros = datos["cobros"]
    reservaciones = datos["reservaciones"]
    mesas = datos["mesas"]
    historico = datos["historico"]
    pagos_hist = datos["historico_pagos"]
    anio = evento["anio"]
    anio_ant = resumen["anio_comparativo"]

    libro = Workbook()

    # ── 1. LEEME ───────────────────────────────────────────────────────────
    leeme = libro.active
    leeme.title = "LEEME"
    leeme.column_dimensions["A"].width = 26
    leeme.column_dimensions["B"].width = 120
    leeme["A1"] = f"{evento['nombre']} — Respaldo en Excel"
    leeme["A1"].font = Font(bold=True, size=18, color=MORADO)
    filas_leeme: list[tuple[str, str]] = [
        ("Generado", f"{datos['generado_en']} por {datos['usuario']}"),
        ("Para qué sirve", "Es la SEGUNDA OPCIÓN por si el sistema no está disponible: trae todas las reservaciones, sillas, pagos y cobros "
                           "al momento de generarlo, para seguir atendiendo y cobrando en papel/Excel y capturar después en el sistema."),
        ("Cómo usarlo", "Cada hoja tiene un título y una explicación arriba, encabezados fijos y filtros (flechita en cada columna). "
                        "Los totales, saldos y acumulados son FÓRMULAS: si cambias un importe se recalculan solos."),
        ("Cuidado", "Es una FOTO del momento en que se descargó: no se actualiza solo. Contiene datos personales de los clientes (teléfono y correo): "
                    "guárdalo en un lugar seguro y no lo compartas."),
        ("", ""),
        ("HOJAS", ""),
        ("Resumen", "Cifras clave: reservaciones, ocupación, dinero, comparativo contra la edición anterior y boletos por tipo de precio."),
        ("Reservaciones", "UNA FILA POR RESERVACIÓN: cliente, contacto, mesa(s), sillas por tipo, importe, pagado, saldo, estado, cortesías y cancelaciones."),
        ("Sillas", "UNA FILA POR SILLA del salón (también las libres y bloqueadas): quién la tiene, tipo, precio, lo pagado y lo que debe."),
        ("Pagos", "UNA FILA POR PAGO registrado (también los cancelados, marcados): fecha, forma de pago, importe, quién lo capturó y a qué sillas se aplicó."),
        ("Cobros por día", "Lo cobrado cada día (Pago 1/2/3), devoluciones, acumulado, sin IVA, póliza de contabilidad y boletos vendidos por tipo."),
        ("Devoluciones", "Dinero regresado a clientes (reservaciones canceladas o con excedente)."),
        ("Comparativo anterior", f"Capturas diarias de {anio_ant} y las listas de adultos/niños de esa edición, para comparar día contra día."),
        ("Configuración", "Precios de preventa y normal, fecha de corte, IVA, propina y mensaje de confirmación."),
        ("", ""),
        ("GLOSARIO", ""),
        ("Preventa / Normal", f"Una silla se vende a precio de PREVENTA si su fecha de venta es igual o anterior al corte de preventa; después, a precio NORMAL. "
                              "El precio se congela al vender; un supervisor puede autorizar preventa fuera de fecha o un precio especial."),
        ("Estados de pago", "Sin pago = no ha pagado nada · Con anticipo = pagó parte · Pagada = liquidada · Cortesía = sin costo (autorizada por un supervisor)."),
        ("Cortesía", "Silla sin costo. Siempre lleva un motivo y la autorización de un supervisor; mientras no se autoriza queda 'por autorizar'."),
        ("Pago 1 / 2 / 3", "Primer, segundo y tercer abono de cada reservación (el 3.º incluye los siguientes). Un cobro con varias formas de pago cuenta como un solo pago."),
        ("Saldo", "Importe de las sillas − lo pagado. Si es negativo, hay excedente (pagó de más)."),
        ("Sin IVA / Desglose", f"Sin IVA = cobrado ÷ {evento['iva_factor']}. El desglose fiscal divide el total entre (IVA + propina {evento['propina_pct']:.0%}); "
                               "el criterio definitivo lo define contraloría."),
    ]
    fila = 3
    for a, b in filas_leeme:
        leeme.cell(row=fila, column=1, value=a)
        c = leeme.cell(row=fila, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        leeme.cell(row=fila, column=1).font = Font(bold=True, color=MORADO if a.isupper() and not b else "000000")
        if len(b) > 115:
            leeme.row_dimensions[fila].height = 15 * (len(b) // 115 + 1)
        fila += 1

    # ── 2. RESUMEN ─────────────────────────────────────────────────────────
    hoja = _hoja(
        libro, "Resumen",
        f"Cifras clave al {_fecha(resumen['fecha_corte']).strftime('%d/%m/%Y')} (corte del resumen). El comparativo pone este año ({anio}) contra "
        f"el cierre de {anio_ant}; las columnas de variación son fórmulas.",
        [("Concepto", 38, None), (str(anio), 16, ENTERO), (f"{anio_ant} (cierre)", 16, ENTERO), ("Variación", 14, ENTERO), ("%", 11, PORCENTAJE)],
    )
    cierre = historico[-1] if historico else None
    r = resumen["reservas"]
    e = resumen["estado"]
    comparativo = [
        ("Adultos", r["adultos"], cierre.get("adultos") if cierre else None),
        ("Niños", r["ninos"], cierre.get("ninos") if cierre else None),
        ("Cortesías", r["cortesias"], cierre["cortesias"] if cierre else None),
        ("TOTAL RESERVADO", r["total"], None),
        ("Pagadas", e["pagadas"], cierre["pagados"] if cierre else None),
        ("Con anticipo", e["anticipos"], cierre["anticipos"] if cierre else None),
        ("Sin anticipo (+ cortesías)", e["sin_anticipo_con_cortesias"], (cierre["sin_anticipo"] + cierre["cortesias"]) if cierre else None),
    ]
    fila = 5
    inicio_comp = fila
    for nombre, actual, anterior in comparativo:
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=actual).number_format = ENTERO
        if nombre == "TOTAL RESERVADO":
            hoja.cell(row=fila, column=3, value=f"=SUM(C{inicio_comp}:C{fila - 1})").number_format = ENTERO
            for col in range(1, 6):
                hoja.cell(row=fila, column=col).font = Font(bold=True)
        elif anterior is not None:
            hoja.cell(row=fila, column=3, value=anterior).number_format = ENTERO
        else:
            hoja.cell(row=fila, column=3, value="sin dato")
        if nombre == "TOTAL RESERVADO" or anterior is not None:
            hoja.cell(row=fila, column=4, value=f"=B{fila}-C{fila}").number_format = ENTERO
            hoja.cell(row=fila, column=5, value=f'=IF(C{fila}=0,"",D{fila}/C{fila})').number_format = PORCENTAJE
        fila += 1
    fila += 1
    hoja.cell(row=fila, column=1, value="OCUPACIÓN Y DINERO").font = Font(bold=True, color=MORADO)
    fila += 1
    o, d = resumen["ocupacion"], resumen["dinero"]
    fila_cap = fila
    bloque = [
        ("Capacidad total del salón (sillas)", resumen["capacidad_total"], ENTERO),
        ("Sillas reservadas", o["ocupadas"], ENTERO),
        ("Sillas libres", f"=B{fila_cap}-B{fila_cap + 1}", ENTERO),
        ("Ocupación", f'=IF(B{fila_cap}=0,"",B{fila_cap + 1}/B{fila_cap})', PORCENTAJE),
        ("Importe total vendido", d["importe_total"], DINERO),
        ("Cobrado", d["pagado"], DINERO),
        ("Por cobrar", f"=B{fila_cap + 4}-B{fila_cap + 5}", DINERO),
        ("Excedente (pagado de más)", d["excedente"], DINERO),
    ]
    for nombre, valor, formato in bloque:
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=valor).number_format = formato
        fila += 1
    fila += 1
    hoja.cell(row=fila, column=1, value="BOLETOS POR TIPO DE PRECIO").font = Font(bold=True, color=MORADO)
    fila += 1
    for i, h in enumerate(["", "Preventa", "Normal", "Cortesía (sin costo)", "Total"], start=1):
        c = hoja.cell(row=fila, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=GRIS)
    fila += 1
    p, n = resumen["preventa"], resumen["normal"]
    tipo_precio = [
        ("Adultos", p["adultos"]["cantidad"], n["adultos"]["cantidad"], None),
        ("Niños", p["ninos"]["cantidad"], n["ninos"]["cantidad"], None),
        ("Cortesías", None, None, resumen["cortesias_tarifa"]["cantidad"]),
    ]
    ini_tp = fila
    for nombre, pre, nor, cor in tipo_precio:
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=pre if pre is not None else 0).number_format = ENTERO
        hoja.cell(row=fila, column=3, value=nor if nor is not None else 0).number_format = ENTERO
        hoja.cell(row=fila, column=4, value=cor if cor is not None else 0).number_format = ENTERO
        hoja.cell(row=fila, column=5, value=f"=SUM(B{fila}:D{fila})").number_format = ENTERO
        fila += 1
    hoja.cell(row=fila, column=1, value="TOTAL").font = Font(bold=True)
    for col in range(2, 6):
        c = hoja.cell(row=fila, column=col, value=f"=SUM({get_column_letter(col)}{ini_tp}:{get_column_letter(col)}{fila - 1})")
        c.number_format = ENTERO
        c.font = Font(bold=True)
    fila += 2
    hoja.cell(row=fila, column=1, value="DESGLOSE FISCAL (según IVA y propina configurados)").font = Font(bold=True, color=MORADO)
    fila += 1
    for nombre, valor in [("Subtotal", resumen["fiscal"]["subtotal"]), ("IVA", resumen["fiscal"]["iva"]), ("Propina", resumen["fiscal"]["propina"]), ("Total", resumen["fiscal"]["total"])]:
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=valor).number_format = DINERO
        fila += 1

    # ── 3. RESERVACIONES ───────────────────────────────────────────────────
    cols = [
        ("# Reserv.", 9, None), ("Cliente", 30, None), ("Teléfono / cuarto", 18, None), ("Correo", 28, None), ("Estatus", 11, None),
        ("Estado de pago", 14, None), ("Fecha de venta", 13, FECHA), ("Mesa(s)", 12, None), ("Adultos", 9, ENTERO), ("Niños", 8, ENTERO),
        ("Cortesías", 10, ENTERO), ("Importe", 14, DINERO), ("Pagado", 14, DINERO), ("Saldo", 14, DINERO), ("Devuelto", 13, DINERO),
        ("Preventa autorizada", 20, None), ("Motivo de cortesía", 28, None), ("Cortesía autorizada por", 20, None),
        ("Capturó", 16, None), ("Notas", 30, None), ("Motivo de cancelación", 28, None),
    ]
    hoja = _hoja(
        libro, "Reservaciones",
        "Una fila por reservación (activas y canceladas). Importe = suma del precio de sus sillas; Pagado = pagos vigentes; "
        "Saldo = Importe − Pagado (fórmula; negativo = pagó de más). Las canceladas conservan lo pagado hasta que se cancelen sus pagos o se devuelva.",
        cols,
    )
    fila = 5
    for rv in reservaciones:
        activas = [s for s in rv["sillas"] if s["activa"]]
        mesas_r = sorted({s["mesa"] for s in activas})
        valores = [
            rv["id"], rv["cliente"], rv["telefono_o_cuarto"], rv["correo"],
            "Cancelada" if rv["estatus"] == "cancelada" else "Activa",
            "Cancelada" if rv["estatus"] == "cancelada" else ETIQUETA_ESTADO_RESERVACION.get(rv["estado"], rv["estado"]),
            _fecha(rv["fecha_venta"]), ", ".join(str(m) for m in mesas_r),
            sum(1 for s in activas if s["tipo"] == "adulto"), sum(1 for s in activas if s["tipo"] == "nino"), sum(1 for s in activas if s["tipo"] == "cortesia"),
            rv["importe"], rv["pagado"], f"=L{fila}-M{fila}", rv["devuelto"],
            (f"Sí — {rv['preventa_autorizada']['por']}: {rv['preventa_autorizada']['motivo']}" if rv["preventa_autorizada"]["activa"] else ""),
            rv["cortesia"]["motivo"], rv["cortesia"]["autorizada_por"] or ("POR AUTORIZAR" if rv["cortesias_por_autorizar"] else ""),
            rv["capturado_por"], rv["notas"], rv["motivo_cancelacion"],
        ]
        _fila(hoja, fila, valores, [f for _, _, f in cols])
        fila += 1
    _total(hoja, fila, "TOTAL (activas y canceladas)", {9: ENTERO, 10: ENTERO, 11: ENTERO, 12: DINERO, 13: DINERO, 14: DINERO, 15: DINERO}, 5, fila - 1, len(cols))
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 4. SILLAS ──────────────────────────────────────────────────────────
    cols = [
        ("Mesa", 8, ENTERO), ("Zona", 12, None), ("Silla", 8, ENTERO), ("¿Extra (11 o 12)?", 12, None), ("Estado", 14, None), ("# Reserv.", 9, None),
        ("Cliente", 30, None), ("Tipo de boleto", 13, None), ("Precio", 13, DINERO), ("Pagado", 13, DINERO), ("Saldo", 13, DINERO), ("Motivo de bloqueo", 28, None),
    ]
    hoja = _hoja(
        libro, "Sillas",
        "Una fila por silla de todo el salón, también las libres y las bloqueadas. Pagado es lo que le toca a esa silla del dinero de su reservación "
        "(se reparte entre las sillas); Saldo = Precio − Pagado (fórmula). Útil para ver qué mesas quedan disponibles.",
        cols,
    )
    fila = 5
    for m in mesas:
        for s in m["sillas"]:
            ocupada = s["reservacion_id"] is not None
            _fila(
                hoja, fila,
                [
                    m["numero"], m.get("zona") or "", s["numero"], "Sí" if s.get("extra") else "", ETIQUETA_ESTADO_SILLA.get(s["estado"], s["estado"]),
                    s["reservacion_id"] or "", s["cliente"] or "", ETIQUETA_TIPO.get(s["tipo"] or "", ""),
                    s["precio"] if ocupada else None, s["pagado"] if ocupada else None, f"=I{fila}-J{fila}" if ocupada else None,
                    s.get("bloqueo_motivo") or "",
                ],
                [f for _, _, f in cols],
            )
            fila += 1
    _total(hoja, fila, "TOTAL", {9: DINERO, 10: DINERO, 11: DINERO}, 5, fila - 1, len(cols))
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 5. PAGOS ───────────────────────────────────────────────────────────
    donde = {s["id"]: (m["numero"], s["numero"]) for m in mesas for s in m["sillas"]}
    cols = [
        ("# Pago", 8, None), ("# Reserv.", 9, None), ("Cliente", 30, None), ("Fecha del pago", 13, FECHA), ("Importe", 14, DINERO),
        ("Forma de pago", 16, None), ("Referencia", 22, None), ("Capturó", 16, None), ("Vigente / Cancelado", 14, None),
        ("Aplicado a sillas", 34, None), ("# Cobro (varias formas)", 18, None),
    ]
    hoja = _hoja(
        libro, "Pagos",
        "Una fila por pago registrado. Los pagos CANCELADOS se muestran pero no cuentan en los totales. Un cobro hecho con varias formas de pago "
        "(p. ej. efectivo + tarjeta) aparece como varias filas con el mismo '# Cobro'.",
        cols,
    )
    fila = 5
    for rv in reservaciones:
        for pg in rv["pagos"]:
            sillas = ", ".join(f"mesa {donde[i][0]} silla {donde[i][1]}" for i in pg["silla_ids"] if i in donde) or "toda la reservación"
            _fila(
                hoja, fila,
                [pg["id"], rv["id"], rv["cliente"], _fecha(pg["fecha"]), pg["importe"], pg["forma_pago"], pg["referencia"], pg["capturado_por"],
                 "Cancelado" if pg["cancelado"] else "Vigente", sillas, pg.get("cobro_id") or ""],
                [f for _, _, f in cols],
            )
            fila += 1
    hoja.cell(row=fila, column=1, value="TOTAL VIGENTE").font = Font(bold=True)
    c = hoja.cell(row=fila, column=5, value=f'=SUMIF(I5:I{max(fila - 1, 5)},"Vigente",E5:E{max(fila - 1, 5)})')
    c.number_format = DINERO
    c.font = Font(bold=True)
    for col in range(1, len(cols) + 1):
        hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=GRIS)
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 6. COBROS POR DÍA ──────────────────────────────────────────────────
    cols = [
        ("Fecha", 12, FECHA), ("Pago 1", 14, DINERO), ("Pago 2", 14, DINERO), ("Pago 3", 14, DINERO), ("Cobrado", 14, DINERO), ("Devuelto", 13, DINERO),
        ("Neto", 14, DINERO), ("Acumulado", 15, DINERO), (f"Sin IVA (÷{evento['iva_factor']})", 15, DINERO), ("Póliza contabilidad", 15, DINERO),
        ("Diferencia vs póliza", 15, DINERO), ("Vendidos hoy · Adulto", 11, ENTERO), ("Niño", 8, ENTERO), ("Cortesía", 9, ENTERO),
        ("Reservadas (acum.)", 12, ENTERO), ("Pagadas (acum.)", 11, ENTERO), ("Con anticipo (acum.)", 12, ENTERO), ("Sin anticipo (acum.)", 12, ENTERO),
        ("Cortesías (acum.)", 11, ENTERO),
    ]
    hoja = _hoja(
        libro, "Cobros por día",
        "Una fila por día, desde el primer movimiento hasta la fecha del evento. Pago 1/2/3 = primer, segundo y tercer abono de cada reservación cobrado ese día. "
        "Cobrado = Pago1+Pago2+Pago3; Neto = Cobrado − Devuelto; Acumulado y Sin IVA son fórmulas. 'Póliza' es lo que contabilidad registró ese día. "
        "Las columnas (acum.) son sillas acumuladas hasta ese día.",
        cols,
    )
    fila = 5
    for dia in cobros["dias"]:
        _fila(
            hoja, fila,
            [
                _fecha(dia["fecha"]), dia["pago1"], dia["pago2"], dia["pago3"], f"=SUM(B{fila}:D{fila})", dia["devuelto"], f"=E{fila}-F{fila}",
                f"=G{fila}" if fila == 5 else f"=H{fila - 1}+G{fila}", f"=G{fila}/{evento['iva_factor']}" if evento["iva_factor"] else f"=G{fila}",
                dia["poliza"], (f"=I{fila}-J{fila}" if dia["poliza"] is not None else None),
                dia["vendidas"]["adulto"], dia["vendidas"]["nino"], dia["vendidas"]["cortesia"],
                dia["reservas"]["reservados"], dia["reservas"]["pagadas"], dia["reservas"]["anticipos"], dia["reservas"]["sin_anticipo"], dia["reservas"]["cortesias"],
            ],
            [f for _, _, f in cols],
        )
        fila += 1
    _total(hoja, fila, "TOTAL", {2: DINERO, 3: DINERO, 4: DINERO, 5: DINERO, 6: DINERO, 7: DINERO, 9: DINERO, 10: DINERO, 12: ENTERO, 13: ENTERO, 14: ENTERO}, 5, fila - 1, len(cols))
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 7. DEVOLUCIONES ────────────────────────────────────────────────────
    cols = [("# Devolución", 12, None), ("# Reserv.", 9, None), ("Cliente", 30, None), ("Fecha", 13, FECHA), ("Importe", 14, DINERO), ("Forma", 16, None),
            ("Referencia", 20, None), ("Motivo", 34, None), ("Capturó", 16, None), ("Autorizó", 16, None), ("Vigente / Anulada", 14, None)]
    hoja = _hoja(
        libro, "Devoluciones",
        "Dinero que se regresó a un cliente. Solo se devuelve lo que sobra: lo cobrado de una reservación cancelada o el excedente de una activa. "
        "Las devoluciones anuladas se muestran pero no restan.",
        cols,
    )
    fila = 5
    for rv in reservaciones:
        for dv in rv["devoluciones"]:
            _fila(
                hoja, fila,
                [dv["id"], rv["id"], rv["cliente"], _fecha(dv["fecha"]), dv["importe"], dv["forma_pago"], dv["referencia"], dv["motivo"],
                 dv["capturado_por"], dv["autorizado_por"], "Anulada" if dv["cancelada"] else "Vigente"],
                [f for _, _, f in cols],
            )
            fila += 1
    hoja.cell(row=fila, column=1, value="TOTAL VIGENTE").font = Font(bold=True)
    c = hoja.cell(row=fila, column=5, value=f'=SUMIF(K5:K{max(fila - 1, 5)},"Vigente",E5:E{max(fila - 1, 5)})')
    c.number_format = DINERO
    c.font = Font(bold=True)
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 8. COMPARATIVO EDICIÓN ANTERIOR ────────────────────────────────────
    cols = [("Fecha", 13, FECHA), ("Sillas reservadas (acum.)", 14, ENTERO), ("Pagadas (acum.)", 12, ENTERO), ("Con anticipo (acum.)", 12, ENTERO),
            ("Sin anticipo (acum.)", 12, ENTERO), ("Cortesías (acum.)", 12, ENTERO), ("Adultos con pago ese día", 14, ENTERO), ("Niños con pago ese día", 14, ENTERO),
            ("Cobrado de niños ese día", 16, DINERO)]
    hoja = _hoja(
        libro, "Comparativo anterior",
        f"Edición {anio_ant}. Las primeras 6 columnas son las capturas acumuladas que se hacían en el Excel (solo algunos días). Las 3 últimas salen de las listas de boletos con "
        "su fecha de pago: adultos y niños se cuentan el día de su PRIMER pago (anticipo y liquidación cuentan una vez); de los adultos solo se conoce la cantidad. "
        "Las cortesías de esa edición no tienen fecha por día.",
        cols,
    )
    por_dia: dict[str, dict] = {x["fecha"]: x for x in pagos_hist.get("dias", [])}
    fechas = sorted({str(f["fecha"])[:10] for f in historico} | set(por_dia))
    cap = {str(f["fecha"])[:10]: f for f in historico}
    fila = 5
    for f in fechas:
        h = cap.get(f)
        ph = por_dia.get(f, {})
        _fila(
            hoja, fila,
            [_fecha(f), h["reservados"] if h else None, h["pagados"] if h else None, h["anticipos"] if h else None, h["sin_anticipo"] if h else None,
             h["cortesias"] if h else None, (ph.get("adulto") or {}).get("reservados"), (ph.get("nino") or {}).get("reservados"), (ph.get("nino") or {}).get("cobrado")],
            [fm for _, _, fm in cols],
        )
        fila += 1
    _total(hoja, fila, "TOTAL (solo las 3 últimas columnas)", {7: ENTERO, 8: ENTERO, 9: DINERO}, 5, fila - 1, len(cols))
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(fila - 1, 5)}"

    # ── 9. CONFIGURACIÓN ───────────────────────────────────────────────────
    hoja = _hoja(
        libro, "Configuración",
        "Reglas con las que se vendió: precios por tipo de boleto, fecha de corte de preventa, IVA y propina. Solo lectura: se cambian en el sistema.",
        [("Concepto", 34, None), ("Valor", 40, None)],
    )
    filas = [
        ("Evento", evento["nombre"]), ("Año", evento["anio"]), ("Fecha del evento", _fecha(evento["fecha_evento"])),
        ("Corte de preventa (hasta este día inclusive)", _fecha(evento["corte_preventa"])), ("Factor de IVA", evento["iva_factor"]),
        ("Propina", evento["propina_pct"]),
    ]
    for pr in evento["precios"]:
        filas.append((f"Precio {ETIQUETA_TIPO.get(pr['tipo'], pr['tipo'])} · preventa", pr["precio_preventa"]))
        filas.append((f"Precio {ETIQUETA_TIPO.get(pr['tipo'], pr['tipo'])} · normal", pr["precio_normal"]))
    filas.append(("Mensaje de confirmación al cliente", evento["mensaje_confirmacion"]))
    fila = 5
    for a, b in filas:
        hoja.cell(row=fila, column=1, value=a).font = Font(bold=True)
        c = hoja.cell(row=fila, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(b, float) and a.startswith("Precio"):
            c.number_format = DINERO
        elif a == "Propina":
            c.number_format = PORCENTAJE
        elif isinstance(b, (date, datetime)):
            c.number_format = FECHA
        fila += 1

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
