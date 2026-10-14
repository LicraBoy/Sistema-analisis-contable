"""
Productos Vendidos — exportación a Excel de una hoja, ya sea de un solo
grupo o del período completo (todos los grupos juntos, ordenados por
Grupo/Producto). Módulo independiente de estilo (no importa nada de
`costeo_excel_export.py`/`excel_processor.py`, mismo criterio de módulos
autocontenidos que ya sigue el resto de este sistema), aunque reutiliza la
paleta azul/bordes grises que ya usan las demás exportaciones.

A diferencia de Costeo, aquí se escriben valores ya calculados (no fórmulas
vivas) — esta hoja es un reporte de lectura, no un archivo que el usuario
vaya a re-subir editado.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import ProductoVendidoDetalle
from app.services.costeo_recetas import calcular_costeo

_BANNER_FILL = PatternFill(start_color="1864AB", end_color="1864AB", fill_type="solid")
_BANNER_FONT = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5FA8", end_color="2F5FA8", fill_type="solid")
_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
_TOTAL_FILL = PatternFill(start_color="F1F5FB", end_color="F1F5FB", fill_type="solid")
_TOTAL_FONT = Font(name="Segoe UI", size=10, bold=True)
_DATO_FONT = Font(name="Segoe UI", size=9)
_ALERTA_FONT = Font(name="Segoe UI", size=9, bold=True, color="C0392B")

_THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

_FORMATO_MONEDA = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
_FORMATO_PCT = "0.0%"

_COLUMNAS = [
    ("Grupo", 22),
    ("Producto", 32),
    ("Categoría", 14),
    ("Sección", 14),
    ("Encontrado en Catálogo", 12),
    ("% Costo", 10),
    ("Costo", 12),
    ("Venta", 12),
    ("Cant.", 10),
    ("Tot. Costo", 14),
    ("Descuento", 12),
    ("Subtotal", 14),
    ("Impuesto", 12),
    ("TOTAL", 14),
]

_COLUMNAS_MONEDA = {"Costo", "Venta", "Tot. Costo", "Descuento", "Subtotal", "Impuesto", "TOTAL"}

FILA_BANNER = 1
FILA_ENCABEZADO = 3
FILA_INICIO_DATOS = 4


def generar_reporte_productos_vendidos_excel(
    productos: list[ProductoVendidoDetalle],
    config: dict[str, float],
    titulo: str,
) -> bytes:
    """
    Arma un Excel de una sola hoja con `productos` (ya ordenados por quien
    llama), agregando `% Costo` calculado con el Costo/Precio de Venta del
    Catálogo (`costo_catalogo`/`precio_venta_catalogo`) — mismo criterio que
    ya usa el frontend, reutilizando `calcular_costeo` (función pura del
    motor de Costeo) solo por sus campos `pct_costo`/`estado`.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos Vendidos"

    ultima_columna = len(_COLUMNAS)
    ws.merge_cells(start_row=FILA_BANNER, start_column=1, end_row=FILA_BANNER, end_column=ultima_columna)
    celda_banner = ws.cell(row=FILA_BANNER, column=1, value=titulo)
    celda_banner.fill = _BANNER_FILL
    celda_banner.font = _BANNER_FONT
    celda_banner.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[FILA_BANNER].height = 26

    for col, (nombre, ancho) in enumerate(_COLUMNAS, start=1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=col, value=nombre)
        celda.fill = _HEADER_FILL
        celda.font = _HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = f"A{FILA_INICIO_DATOS}"

    iva_factor = config.get("iva_factor", 1.16)
    pct_costo_objetivo = config.get("pct_costo_objetivo", 0.28)

    fila = FILA_INICIO_DATOS
    totales = {nombre: 0.0 for nombre in _COLUMNAS_MONEDA}

    for producto in productos:
        costeo = calcular_costeo(
            costo=producto.costo_catalogo,
            precio_venta=producto.precio_venta_catalogo,
            pct_incremento=0.0,
            iva_factor=iva_factor,
            descuento_pct=0.0,
            pct_costo_objetivo=pct_costo_objetivo,
        )
        pct_costo = costeo["pct_costo"]
        fuera_de_objetivo = costeo["estado"] == "REQ AJUSTE"

        valores = [
            producto.grupo,
            producto.nombre_producto,
            producto.categoria_tipo or "",
            producto.seccion_menu or "",
            "Sí" if producto.producto_catalogo_encontrado else "No",
            pct_costo,
            producto.costo,
            producto.precio_venta,
            producto.cantidad,
            producto.total_costo,
            producto.descuento,
            producto.subtotal,
            producto.impuesto,
            producto.total,
        ]

        for col, (nombre_columna, _ancho) in enumerate(_COLUMNAS, start=1):
            valor = valores[col - 1]
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.border = _THIN_BORDER
            celda.font = _ALERTA_FONT if (nombre_columna == "Encontrado en Catálogo" and valor == "No") else _DATO_FONT
            if nombre_columna in _COLUMNAS_MONEDA:
                celda.number_format = _FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right")
                if isinstance(valor, (int, float)):
                    totales[nombre_columna] += valor
            elif nombre_columna == "% Costo":
                celda.number_format = _FORMATO_PCT
                celda.alignment = Alignment(horizontal="right")
                if pct_costo is not None and fuera_de_objetivo:
                    celda.font = _ALERTA_FONT
            elif nombre_columna == "Cant.":
                celda.alignment = Alignment(horizontal="right")

        fila += 1

    # Fila de totales.
    ws.cell(row=fila, column=1, value="TOTAL").font = _TOTAL_FONT
    ws.cell(row=fila, column=1).fill = _TOTAL_FILL
    for col in range(1, ultima_columna + 1):
        ws.cell(row=fila, column=col).fill = _TOTAL_FILL
        ws.cell(row=fila, column=col).border = _THIN_BORDER
    for col, (nombre_columna, _ancho) in enumerate(_COLUMNAS, start=1):
        if nombre_columna in totales:
            celda = ws.cell(row=fila, column=col, value=round(totales[nombre_columna], 2))
            celda.font = _TOTAL_FONT
            celda.number_format = _FORMATO_MONEDA
            celda.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
