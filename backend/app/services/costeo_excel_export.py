"""
Costeo de Recetas — exportación a Excel con fórmulas vivas (no valores ya
calculados), para que el usuario pueda seguir auditando/editando en Excel
tal como hacía con el archivo original. Módulo independiente, comparte
`CatalogoProducto` solo para leer costo/precio_venta.

Diseño clave: los parámetros de negocio (IVA, % descuento, % costo
objetivo, % incremento por grupo) se escriben como celdas visibles en una
hoja "Parámetros", y TODAS las fórmulas de las demás hojas los referencian
por celda (`Parametros!$B$2`, etc.) en vez de tener el número embebido —
así cambiar un valor en "Parámetros" recalcula todo el libro, igual que
pedía la spec. El % de incremento es por `grupo` (no uno solo por hoja),
consistente con cómo ya funciona el resto del módulo de Costeo.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import CatalogoProducto

# ── Estilo — misma paleta que el resto de exportaciones del sistema
# (services/excel_processor.py: azul 1864AB, "Segoe UI", bordes grises finos).
_BANNER_FILL = PatternFill(start_color="1864AB", end_color="1864AB", fill_type="solid")
_BANNER_FONT = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5FA8", end_color="2F5FA8", fill_type="solid")
_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
_BLOQUE_FILL = PatternFill(start_color="E7EEFA", end_color="E7EEFA", fill_type="solid")
_BLOQUE_FONT = Font(name="Segoe UI", size=10, bold=True, italic=True, underline="single", color="1864AB")
_PROMEDIO_FILL = PatternFill(start_color="F1F5FB", end_color="F1F5FB", fill_type="solid")
_PROMEDIO_FONT = Font(name="Segoe UI", size=9, bold=True)
_GLOBAL_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
_GLOBAL_FONT = Font(name="Segoe UI", size=13, bold=True, color="FF0000")
_DATO_FONT = Font(name="Segoe UI", size=9)
_EDITADO_FONT = Font(name="Segoe UI", size=9, bold=True, color="9C5700")
_PARAM_LABEL_FONT = Font(name="Segoe UI", size=10, bold=True)
_PARAM_VALOR_FILL = PatternFill(start_color="FFF3BF", end_color="FFF3BF", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

_FORMATO_MONEDA = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
_FORMATO_PCT = "0.0%"

# Columnas fijas de cada hoja de sección (1-based) — mismo layout que la spec,
# con F/H/I/M/N reservadas para separación/promedios sin encabezado propio.
COL_COD = 1
COL_PRODUCTO = 2
COL_PVENTA = 3
COL_COSTO = 4
COL_PCT_COSTO = 5
COL_PROMEDIO_PCT_COSTO = 6
COL_PCT_COSTO_DESC = 7
COL_PROMEDIO_PCT_COSTO_DESC = 8
COL_INCREMENTO = 10
COL_PRECIO_SUGERIDO = 11
COL_PCT_COSTO_SUGERIDO = 12
COL_PCT_COSTO_SUGERIDO_DESC = 15
COL_INCREMENTO_EFECTIVO = 18
ULTIMA_COLUMNA_HOJA = 18  # R — para fusionar banners/anchos

FILA_BANNER = 1
FILA_ENCABEZADO = 4
FILA_INICIO_DATOS = 5

# Filas fijas de la hoja "Parámetros".
PARAM_FILA_IVA = 2
PARAM_FILA_DESCUENTO = 3
PARAM_FILA_OBJETIVO = 4
PARAM_FILA_INICIO_GRUPOS = 7


def _aplicar_borde(ws: Worksheet, fila: int, col_inicio: int, col_fin: int) -> None:
    for col in range(col_inicio, col_fin + 1):
        ws.cell(row=fila, column=col).border = _THIN_BORDER


def _escribir_hoja_parametros(
    wb: Workbook,
    config: dict[str, float],
    incrementos_por_grupo: dict[str, float],
    grupos_usados: list[str],
) -> dict[str, int]:
    """Escribe la hoja 'Parámetros' y regresa {grupo: fila} para que las hojas de sección referencien su propio % de incremento."""
    ws = wb.create_sheet("Parámetros")
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    banner = ws.cell(row=1, column=1, value="Parámetros de Costeo")
    banner.fill = _BANNER_FILL
    banner.font = _BANNER_FONT
    banner.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    filas_globales = [
        (PARAM_FILA_IVA, "IVA (factor)", config["iva_factor"], "0.00"),
        (PARAM_FILA_DESCUENTO, "% Descuento / promoción", config["descuento_pct"], _FORMATO_PCT),
        (PARAM_FILA_OBJETIVO, "% Costo objetivo (OK / REQ AJUSTE)", config["pct_costo_objetivo"], _FORMATO_PCT),
    ]
    for fila, etiqueta, valor, formato in filas_globales:
        celda_label = ws.cell(row=fila, column=1, value=etiqueta)
        celda_label.font = _PARAM_LABEL_FONT
        celda_valor = ws.cell(row=fila, column=2, value=valor)
        celda_valor.font = _PARAM_LABEL_FONT
        celda_valor.fill = _PARAM_VALOR_FILL
        celda_valor.number_format = formato
        celda_valor.alignment = Alignment(horizontal="center")
        _aplicar_borde(ws, fila, 1, 2)

    ws.cell(row=PARAM_FILA_INICIO_GRUPOS - 1, column=1, value="Grupo").font = _HEADER_FONT
    ws.cell(row=PARAM_FILA_INICIO_GRUPOS - 1, column=2, value="% Incremento").font = _HEADER_FONT
    for col in (1, 2):
        celda = ws.cell(row=PARAM_FILA_INICIO_GRUPOS - 1, column=col)
        celda.fill = _HEADER_FILL

    fila_por_grupo: dict[str, int] = {}
    fila = PARAM_FILA_INICIO_GRUPOS
    for grupo in sorted(grupos_usados):
        ws.cell(row=fila, column=1, value=grupo).font = _DATO_FONT
        celda_pct = ws.cell(row=fila, column=2, value=incrementos_por_grupo.get(grupo, config["incremento_default"]))
        celda_pct.number_format = _FORMATO_PCT
        celda_pct.font = _DATO_FONT
        celda_pct.alignment = Alignment(horizontal="center")
        _aplicar_borde(ws, fila, 1, 2)
        fila_por_grupo[grupo] = fila
        fila += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    return fila_por_grupo


def _escribir_encabezados_hoja(ws: Worksheet, nombre_seccion: str) -> None:
    ws.merge_cells(start_row=FILA_BANNER, start_column=1, end_row=FILA_BANNER, end_column=ULTIMA_COLUMNA_HOJA)
    banner = ws.cell(row=FILA_BANNER, column=1, value=f"Costeo de Recetas — {nombre_seccion}")
    banner.fill = _BANNER_FILL
    banner.font = _BANNER_FONT
    banner.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[FILA_BANNER].height = 26

    encabezados = {
        COL_COD: "COD.",
        COL_PRODUCTO: "Producto",
        COL_PVENTA: "P. Venta Act",
        COL_COSTO: "Costo",
        COL_PCT_COSTO: "% Costo",
        COL_PCT_COSTO_DESC: "% Costo (Desc. 5%)",
        COL_INCREMENTO: "Increm.",
        COL_PRECIO_SUGERIDO: "Precio Sugerido",
        COL_PCT_COSTO_SUGERIDO: "% Costo (Sugerido)",
        COL_PCT_COSTO_SUGERIDO_DESC: "% Costo Sugerido (Desc. 5%)",
    }
    for col in range(1, ULTIMA_COLUMNA_HOJA + 1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=col, value=encabezados.get(col, ""))
        celda.fill = _HEADER_FILL
        celda.font = _HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[FILA_ENCABEZADO].height = 30

    anchos = {
        COL_COD: 8, COL_PRODUCTO: 30, COL_PVENTA: 12, COL_COSTO: 12, COL_PCT_COSTO: 10,
        COL_PROMEDIO_PCT_COSTO: 10, COL_PCT_COSTO_DESC: 12, COL_PROMEDIO_PCT_COSTO_DESC: 10,
        9: 3, COL_INCREMENTO: 10, COL_PRECIO_SUGERIDO: 14, COL_PCT_COSTO_SUGERIDO: 12,
        13: 3, 14: 3, COL_PCT_COSTO_SUGERIDO_DESC: 14, 16: 3, 17: 3, COL_INCREMENTO_EFECTIVO: 14,
    }
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho

    ws.freeze_panes = f"C{FILA_INICIO_DATOS}"
    ws.sheet_view.showGridLines = False


def _escribir_hoja_seccion(
    wb: Workbook,
    nombre_seccion: str,
    productos: list[CatalogoProducto],
    fila_por_grupo_parametros: dict[str, int],
    overrides_precio_sugerido: dict[str, float],
) -> dict[str, dict]:
    """
    Escribe la hoja de una sección (Desayuno, etc.), agrupada por `grupo` en
    bloques con fila de promedio. Regresa, por grupo, la fila de su fila de
    promedio (para que RESUMEN la referencie) — la hoja RESUMEN opera a nivel
    de grupo/subcategoría, igual que el RESUMEN del Excel original.
    """
    ws = wb.create_sheet(nombre_seccion[:31])
    _escribir_encabezados_hoja(ws, nombre_seccion)

    productos_por_grupo: dict[str, list[CatalogoProducto]] = {}
    for p in productos:
        productos_por_grupo.setdefault(p.grupo, []).append(p)

    info_por_grupo: dict[str, dict] = {}
    fila = FILA_INICIO_DATOS
    primera_fila_datos: int | None = None
    ultima_fila_datos = FILA_INICIO_DATOS - 1

    for grupo in sorted(productos_por_grupo.keys()):
        productos_grupo = sorted(productos_por_grupo[grupo], key=lambda p: p.nombre_producto)

        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ULTIMA_COLUMNA_HOJA)
        celda_bloque = ws.cell(row=fila, column=1, value=grupo)
        celda_bloque.fill = _BLOQUE_FILL
        celda_bloque.font = _BLOQUE_FONT
        fila += 1

        fila_inicio_bloque = fila
        for producto in productos_grupo:
            if producto.costo is None or producto.precio_venta is None or producto.precio_venta == 0:
                # Sin costo/precio no hay costeo posible — se listan los datos
                # base y se dejan las fórmulas fuera (Excel mostraría #DIV/0!).
                ws.cell(row=fila, column=COL_COD, value=producto.id)
                ws.cell(row=fila, column=COL_PRODUCTO, value=producto.nombre_producto)
                if producto.precio_venta is not None:
                    ws.cell(row=fila, column=COL_PVENTA, value=producto.precio_venta)
                if producto.costo is not None:
                    ws.cell(row=fila, column=COL_COSTO, value=producto.costo)
                for col in range(1, ULTIMA_COLUMNA_HOJA + 1):
                    ws.cell(row=fila, column=col).font = _DATO_FONT
                fila += 1
                continue

            c, p_col, d_col, e_col, g_col, j_col, k_col, l_col, o_col = (
                get_column_letter(COL_COD), get_column_letter(COL_PVENTA), get_column_letter(COL_COSTO),
                get_column_letter(COL_PCT_COSTO), get_column_letter(COL_PCT_COSTO_DESC),
                get_column_letter(COL_INCREMENTO), get_column_letter(COL_PRECIO_SUGERIDO),
                get_column_letter(COL_PCT_COSTO_SUGERIDO), get_column_letter(COL_PCT_COSTO_SUGERIDO_DESC),
            )
            r = fila
            fila_incremento_grupo = fila_por_grupo_parametros.get(grupo)

            ws.cell(row=r, column=COL_COD, value=producto.id)
            ws.cell(row=r, column=COL_PRODUCTO, value=producto.nombre_producto)
            ws.cell(row=r, column=COL_PVENTA, value=producto.precio_venta)
            ws.cell(row=r, column=COL_COSTO, value=producto.costo)
            ws.cell(row=r, column=COL_PCT_COSTO, value=f"={d_col}{r}/({p_col}{r}/Parámetros!$B${PARAM_FILA_IVA})")
            ws.cell(
                row=r,
                column=COL_PCT_COSTO_DESC,
                value=f"={d_col}{r}/(({p_col}{r}*(1-Parámetros!$B${PARAM_FILA_DESCUENTO}))/Parámetros!$B${PARAM_FILA_IVA})",
            )

            override = overrides_precio_sugerido.get(producto.nombre_producto)
            celda_precio_sugerido = ws.cell(row=r, column=COL_PRECIO_SUGERIDO)
            if override is not None:
                celda_precio_sugerido.value = override
                celda_precio_sugerido.font = _EDITADO_FONT
            elif fila_incremento_grupo is not None:
                celda_precio_sugerido.value = f"=ROUND({p_col}{r}*(1+Parámetros!$B${fila_incremento_grupo}),0)"

            ws.cell(row=r, column=COL_INCREMENTO, value=f"=({k_col}{r}/{p_col}{r})-1")
            ws.cell(row=r, column=COL_PCT_COSTO_SUGERIDO, value=f"={d_col}{r}/({k_col}{r}/Parámetros!$B${PARAM_FILA_IVA})")
            ws.cell(
                row=r,
                column=COL_PCT_COSTO_SUGERIDO_DESC,
                value=f"={d_col}{r}/(({k_col}{r}*(1-Parámetros!$B${PARAM_FILA_DESCUENTO}))/Parámetros!$B${PARAM_FILA_IVA})",
            )

            for col in (COL_PVENTA, COL_COSTO, COL_PRECIO_SUGERIDO):
                ws.cell(row=r, column=col).number_format = _FORMATO_MONEDA
            for col in (COL_PCT_COSTO, COL_PCT_COSTO_DESC, COL_INCREMENTO, COL_PCT_COSTO_SUGERIDO, COL_PCT_COSTO_SUGERIDO_DESC):
                ws.cell(row=r, column=col).number_format = _FORMATO_PCT
            for col in range(1, ULTIMA_COLUMNA_HOJA + 1):
                celda = ws.cell(row=r, column=col)
                if col not in (COL_PRECIO_SUGERIDO,) or override is None:
                    celda.font = _DATO_FONT
            fila += 1

        fila_fin_bloque = fila - 1
        if primera_fila_datos is None:
            primera_fila_datos = fila_inicio_bloque
        ultima_fila_datos = fila_fin_bloque

        e_col = get_column_letter(COL_PCT_COSTO)
        g_col = get_column_letter(COL_PCT_COSTO_DESC)
        k_col = get_column_letter(COL_PRECIO_SUGERIDO)
        p_col = get_column_letter(COL_PVENTA)
        fila_promedio = fila
        ws.cell(row=fila_promedio, column=COL_PRODUCTO, value=f"Promedio {grupo}")
        ws.cell(row=fila_promedio, column=COL_PROMEDIO_PCT_COSTO, value=f"=AVERAGE({e_col}{fila_inicio_bloque}:{e_col}{fila_fin_bloque})")
        ws.cell(
            row=fila_promedio,
            column=COL_PROMEDIO_PCT_COSTO_DESC,
            value=f"=AVERAGE({g_col}{fila_inicio_bloque}:{g_col}{fila_fin_bloque})",
        )
        ws.cell(
            row=fila_promedio,
            column=COL_INCREMENTO_EFECTIVO,
            value=f"=SUM({k_col}{fila_inicio_bloque}:{k_col}{fila_fin_bloque})/SUM({p_col}{fila_inicio_bloque}:{p_col}{fila_fin_bloque})-1",
        )
        for col in range(1, ULTIMA_COLUMNA_HOJA + 1):
            celda = ws.cell(row=fila_promedio, column=col)
            celda.fill = _PROMEDIO_FILL
            celda.font = _PROMEDIO_FONT
            celda.alignment = Alignment(horizontal="center")
        ws.cell(row=fila_promedio, column=COL_PROMEDIO_PCT_COSTO).number_format = _FORMATO_PCT
        ws.cell(row=fila_promedio, column=COL_PROMEDIO_PCT_COSTO_DESC).number_format = _FORMATO_PCT
        ws.cell(row=fila_promedio, column=COL_INCREMENTO_EFECTIVO).number_format = _FORMATO_PCT

        info_por_grupo[grupo] = {"hoja": ws.title, "fila_promedio": fila_promedio}
        fila = fila_promedio + 2  # espacio en blanco entre bloques

    if primera_fila_datos is not None:
        # Fila de totales del pie de la hoja — reutiliza posiciones de
        # columna ya existentes (D/E/J/L/R) igual que la fila "Promedio" de
        # cada grupo reutiliza F/H/R, en vez de abrir columnas nuevas:
        # D = total Venta Actual, E = % Costo Global Actual (misma columna
        # que usa cada fila de producto para su % Costo), J = total Precio
        # Sugerido, L = % Costo Global Sugerido (misma columna que % Costo
        # Sugerido por fila, dividido contra el total de precio sugerido —
        # no contra venta actual), R = % incremento efectivo global (total
        # sugerido / total venta - 1). Fórmulas confirmadas contra el Excel
        # original del usuario.
        fila_global = fila
        c_col = get_column_letter(COL_PVENTA)
        d_col = get_column_letter(COL_COSTO)
        j_col = get_column_letter(COL_INCREMENTO)
        k_col = get_column_letter(COL_PRECIO_SUGERIDO)

        ws.merge_cells(start_row=fila_global, start_column=1, end_row=fila_global, end_column=COL_PVENTA - 1)
        ws.cell(row=fila_global, column=1, value=f"TOTALES — {nombre_seccion.upper()}")

        celda_total_venta = ws.cell(
            row=fila_global,
            column=COL_COSTO,
            value=f"=SUM({c_col}{primera_fila_datos}:{c_col}{ultima_fila_datos})",
        )
        celda_total_venta.number_format = _FORMATO_MONEDA

        celda_pct_global_actual = ws.cell(
            row=fila_global,
            column=COL_PCT_COSTO,
            value=(
                f"=SUM({d_col}{primera_fila_datos}:{d_col}{ultima_fila_datos})"
                f"/(SUM({c_col}{primera_fila_datos}:{c_col}{ultima_fila_datos})/Parámetros!$B${PARAM_FILA_IVA})"
            ),
        )
        celda_pct_global_actual.number_format = _FORMATO_PCT

        celda_total_sugerido = ws.cell(
            row=fila_global,
            column=COL_INCREMENTO,
            value=f"=SUM({k_col}{primera_fila_datos}:{k_col}{ultima_fila_datos})",
        )
        celda_total_sugerido.number_format = _FORMATO_MONEDA

        celda_pct_global_sugerido = ws.cell(
            row=fila_global,
            column=COL_PCT_COSTO_SUGERIDO,
            value=(
                f"=SUM({d_col}{primera_fila_datos}:{d_col}{ultima_fila_datos})"
                f"/(SUM({k_col}{primera_fila_datos}:{k_col}{ultima_fila_datos})/Parámetros!$B${PARAM_FILA_IVA})"
            ),
        )
        celda_pct_global_sugerido.number_format = _FORMATO_PCT

        celda_incremento_global = ws.cell(
            row=fila_global,
            column=COL_INCREMENTO_EFECTIVO,
            value=f"=+{j_col}{fila_global}/{d_col}{fila_global}-1",
        )
        celda_incremento_global.number_format = _FORMATO_PCT

        for col in (1, COL_COSTO, COL_PCT_COSTO, COL_INCREMENTO, COL_PCT_COSTO_SUGERIDO, COL_INCREMENTO_EFECTIVO):
            celda = ws.cell(row=fila_global, column=col)
            celda.fill = _GLOBAL_FILL
            celda.font = _GLOBAL_FONT
            celda.alignment = Alignment(horizontal="center")
        ws.cell(row=fila_global, column=1).alignment = Alignment(horizontal="left")
        ws.row_dimensions[fila_global].height = 22

        info_por_grupo["__global__"] = {
            "hoja": ws.title,
            "rango_costo": (primera_fila_datos, ultima_fila_datos, d_col),
            "rango_precio": (primera_fila_datos, ultima_fila_datos, c_col),
        }

    return info_por_grupo


def _escribir_hoja_resumen(wb: Workbook, info_por_seccion: dict[str, dict[str, dict]]) -> None:
    """
    Una fila por grupo/subcategoría de todas las secciones exportadas
    (equivalente a la hoja RESUMEN del Excel original), más una fila GLOBAL
    con el % de costo ponderado de todo lo exportado.
    """
    ws = wb.create_sheet("RESUMEN")
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    banner = ws.cell(row=1, column=1, value="RESUMEN — % de Costo por Categoría")
    banner.fill = _BANNER_FILL
    banner.font = _BANNER_FONT
    ws.row_dimensions[1].height = 26

    encabezados = ["", "Categoría", "", "% Costo Actual", "% Costo (Desc. 5%)", "", "Observación"]
    for col, texto in enumerate(encabezados, start=1):
        if not texto:
            continue
        celda = ws.cell(row=3, column=col, value=texto)
        celda.fill = _HEADER_FILL
        celda.font = _HEADER_FONT
        celda.alignment = Alignment(horizontal="center")

    fila = 4
    rangos_costo: list[tuple[str, int, int, str]] = []
    rangos_precio: list[tuple[str, int, int, str]] = []
    for info_grupos in info_por_seccion.values():
        global_info = info_grupos.get("__global__")
        if global_info:
            rangos_costo.append((global_info["hoja"], *global_info["rango_costo"]))
            rangos_precio.append((global_info["hoja"], *global_info["rango_precio"]))
        for grupo, info in sorted(info_grupos.items()):
            if grupo == "__global__":
                continue
            hoja = info["hoja"]
            f = info["fila_promedio"]
            f_col = get_column_letter(COL_PROMEDIO_PCT_COSTO)
            h_col = get_column_letter(COL_PROMEDIO_PCT_COSTO_DESC)

            ws.cell(row=fila, column=2, value=f"{grupo} ({hoja})")
            celda_d = ws.cell(row=fila, column=4, value=f"='{hoja}'!{f_col}{f}")
            celda_e = ws.cell(row=fila, column=5, value=f"='{hoja}'!{h_col}{f}")
            celda_d.number_format = _FORMATO_PCT
            celda_e.number_format = _FORMATO_PCT
            ws.cell(
                row=fila,
                column=7,
                value=f'=IF(D{fila}<=Parámetros!$B${PARAM_FILA_OBJETIVO},"OK","REQ AJUSTE")',
            )
            for col in range(1, 8):
                ws.cell(row=fila, column=col).font = _DATO_FONT
                ws.cell(row=fila, column=col).border = _THIN_BORDER
            fila += 1

    fila += 1
    ws.cell(row=fila, column=2, value="GLOBAL (todo lo exportado)").font = _GLOBAL_FONT
    if rangos_costo:
        suma_costo = "+".join(f"SUM('{hoja}'!{col}{ini}:{col}{fin})" for hoja, ini, fin, col in rangos_costo)
        suma_precio = "+".join(f"SUM('{hoja}'!{col}{ini}:{col}{fin})" for hoja, ini, fin, col in rangos_precio)
        celda_global = ws.cell(
            row=fila,
            column=4,
            value=f"=({suma_costo})/(({suma_precio})/Parámetros!$B${PARAM_FILA_IVA})",
        )
        celda_global.number_format = _FORMATO_PCT
        celda_global.font = _GLOBAL_FONT
        for col in range(1, 5):
            ws.cell(row=fila, column=col).fill = _GLOBAL_FILL

    fila += 2
    ws.cell(row=fila, column=2, value="Ajuste general sugerido:").font = _PARAM_LABEL_FONT
    ws.cell(row=fila, column=4, value="Revisar categorías marcadas REQ AJUSTE").font = _DATO_FONT

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 3
    ws.column_dimensions["G"].width = 16


def generar_reporte_costeo_excel(
    productos_por_seccion: dict[str, list[CatalogoProducto]],
    config: dict[str, float],
    incrementos_por_grupo: dict[str, float],
    overrides_precio_sugerido: dict[str, float],
) -> bytes:
    """
    Genera el libro completo: hoja "Parámetros" + una hoja por sección de
    `productos_por_seccion` (en el orden dado) + hoja "RESUMEN" al final.
    Todas las columnas calculadas se escriben como fórmulas de Excel vivas
    (excepto el precio sugerido de un producto con ajuste manual, que se
    escribe como valor literal — mismo criterio que Facturación usa para
    celdas editadas a mano).
    """
    wb = Workbook()
    wb.remove(wb.active)

    grupos_usados = sorted({p.grupo for productos in productos_por_seccion.values() for p in productos})
    fila_por_grupo_parametros = _escribir_hoja_parametros(wb, config, incrementos_por_grupo, grupos_usados)

    info_por_seccion: dict[str, dict[str, dict]] = {}
    for nombre_seccion, productos in productos_por_seccion.items():
        info_por_seccion[nombre_seccion] = _escribir_hoja_seccion(
            wb, nombre_seccion, productos, fila_por_grupo_parametros, overrides_precio_sugerido
        )

    _escribir_hoja_resumen(wb, info_por_seccion)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
