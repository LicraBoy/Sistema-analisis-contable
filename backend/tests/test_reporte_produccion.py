"""Tests de `services/reporte_produccion.py` (cadena de saldos, entradas automáticas de cocina) y del router."""

from datetime import date

import pytest
from sqlalchemy import select

from app.api import reporte_produccion as api
from app.models import ReporteProduccionCaptura, ReporteProduccionRenglon
from app.services.reporte_produccion import (
    PRECIOS_KG_SEMILLA,
    agrupar_por_seccion,
    RENDIMIENTOS_SEMILLA,
    calcular_cadena_renglon,
    calcular_control_limpieza,
    calcular_inventario_bruto,
    completar_rendimientos,
    construir_clave_cruce,
    construir_entradas_automaticas,
    estado_diferencia,
    factor_de_pesajes,
    factores_por_producto,
    filas_semilla,
    gramaje_desde_texto,
    gramaje_efectivo_kg,
    merma_kg,
    renglones_sin_contraparte,
    resolver_crudos,
    salidas_auto_crudo,
)

D1, D2, D3, D5 = date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 3), date(2026, 10, 5)


def _renglon(id, hoja, producto, presentacion=None, activo=True):
    return ReporteProduccionRenglon(
        id=id, hoja=hoja, seccion="X", producto=producto, presentacion=presentacion, unidad="PORCION",
        orden=id, activo=activo, clave_cruce=construir_clave_cruce(producto, presentacion),
    )


def _cap(fecha, inicial=None, entradas=None, salidas=None, conteo=None, renglon_id=1):
    return ReporteProduccionCaptura(
        fecha=fecha, renglon_id=renglon_id, inicial=inicial, entradas=entradas, salidas=salidas, conteo_fisico=conteo
    )


# ── Cadena de saldos ───────────────────────────────────────────────────


def test_cadena_primer_dia_es_apertura_y_calcula_final():
    (d1,) = calcular_cadena_renglon([_cap(D1, inicial=20, entradas=10, salidas=10)])
    assert d1["es_apertura"] is True
    assert d1["final_calculado"] == 20 and d1["existencia_final"] == 20


def test_cadena_arrastra_el_final_como_inicial_del_dia_siguiente():
    d1, d2, d3 = calcular_cadena_renglon([
        _cap(D3, entradas=1, salidas=6),  # desordenado a propósito
        _cap(D1, inicial=20, entradas=10, salidas=10),
        _cap(D2, entradas=5, salidas=3),
    ])
    assert (d1["existencia_final"], d2["inicial"], d2["existencia_final"]) == (20, 20, 22)
    assert (d3["inicial"], d3["final_calculado"]) == (22, 17)
    assert d2["es_apertura"] is False and d3["es_apertura"] is False


def test_inicial_manual_corrige_el_arrastre_y_se_propaga():
    d1, d2, d3 = calcular_cadena_renglon([
        _cap(D1, inicial=5),
        _cap(D2, inicial=20, entradas=1),  # el día anterior estuvo mal capturado: se corrige el inicial
        _cap(D3, salidas=1),
    ])
    assert d2["inicial_arrastrado"] == 5 and d2["inicial_manual"] == 20
    assert (d2["inicial"], d2["final_calculado"], d3["inicial"], d3["final_calculado"]) == (20, 21, 21, 20)
    assert d1["es_apertura"] is True and d2["es_apertura"] is False


def test_cadena_dias_sin_captura_se_saltan():
    d1, d5 = calcular_cadena_renglon([_cap(D1, inicial=10, salidas=4), _cap(D5, entradas=2)])
    assert d5["inicial"] == 6 and d5["final_calculado"] == 8


def test_cadena_corregir_dia_1_recalcula_los_siguientes():
    antes = calcular_cadena_renglon([_cap(D1, inicial=10, salidas=4), _cap(D2, entradas=2), _cap(D3, salidas=1)])
    despues = calcular_cadena_renglon([_cap(D1, inicial=10, salidas=2), _cap(D2, entradas=2), _cap(D3, salidas=1)])
    assert antes[-1]["existencia_final"] == 7
    assert despues[-1]["existencia_final"] == 9


def test_conteo_fisico_reemplaza_al_calculado_y_da_diferencia():
    d1, d2 = calcular_cadena_renglon([_cap(D1, inicial=10, salidas=2, conteo=7), _cap(D2, entradas=1)])
    assert d1["final_calculado"] == 8 and d1["diferencia"] == -1
    assert estado_diferencia(d1["diferencia"]) == "faltante"
    assert d2["inicial"] == 7 and d2["final_calculado"] == 8  # el conteo corrige el arrastre


def test_sin_conteo_no_hay_diferencia_y_cero_es_conteo_real():
    (d1,) = calcular_cadena_renglon([_cap(D1, inicial=3)])
    assert d1["diferencia"] is None and estado_diferencia(None) == "sin_conteo"
    (d,) = calcular_cadena_renglon([_cap(D1, inicial=3, conteo=0)])
    assert d["diferencia"] == -3  # conteo 0 capturado ≠ vacío


def test_final_negativo_marca_alerta():
    (d1,) = calcular_cadena_renglon([_cap(D1, inicial=2, salidas=5)])
    assert d1["final_calculado"] == -3 and d1["alerta_negativo"] is True


def test_estado_diferencia_sobrante():
    assert estado_diferencia(0.001) == "cuadra"
    assert estado_diferencia(2) == "sobrante"


# ── Catálogo / clave de cruce ──────────────────────────────────────────


def test_clave_cruce_normaliza_acentos_y_espacios():
    assert construir_clave_cruce("Camarón  U-15", " Tampiqueña ") == "CAMARON U-15|TAMPIQUENA"
    assert construir_clave_cruce("CARNITAS", None) == "CARNITAS|"


def test_semilla_reproduce_renglones_del_formato():
    porc, coc = filas_semilla("porcionador"), filas_semilla("cocina")
    assert len(porc) == 45 and len(coc) == 32
    assert not any(f["producto"].startswith("9") for f in coc)  # errata "9PULPA" corregida
    assert any(f["producto"] == "PESCADO AHUMADO (MARLIN)" for f in coc)  # paréntesis cerrado


# ── Entradas automáticas de cocina ─────────────────────────────────────


def test_entradas_cocina_son_las_salidas_del_porcionador_por_clave():
    renglones = [
        _renglon(1, "porcionador", "FILETE DE RES 1 KG", "Mignon"),
        _renglon(2, "cocina", "FILETE DE RES 1 KG", "Mignon"),
        _renglon(3, "cocina", "CORTE RIB EYE"),  # sin equivalente en porcionador
        _renglon(4, "porcionador", "POLLO CANAL 1 KG", "Pollo BBQ", activo=False),
    ]
    capturas = [_cap(D1, salidas=20, renglon_id=1)]
    auto = construir_entradas_automaticas(renglones, capturas)
    assert auto(renglones[1], D1) == 20
    assert auto(renglones[1], D2) == 0  # equivalente pero sin salidas ese día
    assert auto(renglones[2], D1) is None  # captura a mano
    assert auto(renglones[0], D1) is None  # solo aplica a cocina


def test_entradas_cocina_suman_renglones_porcionador_con_misma_clave():
    renglones = [
        _renglon(1, "porcionador", "CALAMAR", "Porcion"),
        _renglon(2, "porcionador", "CALAMAR", "Porcion"),
        _renglon(3, "cocina", "CALAMAR", "Porcion"),
    ]
    capturas = [_cap(D1, salidas=4, renglon_id=1), _cap(D1, salidas=6, renglon_id=2)]
    assert construir_entradas_automaticas(renglones, capturas)(renglones[2], D1) == 10


def test_renglones_sin_contraparte_lista_ambos_lados():
    renglones = [
        _renglon(1, "porcionador", "CARNITAS"),
        _renglon(2, "cocina", "CARNITAS"),
        _renglon(3, "porcionador", "PULPA DE JAIBA 1 KG"),
        _renglon(4, "cocina", "CORTE T-BONE"),
    ]
    res = {(r["producto"], r["hoja"]) for r in renglones_sin_contraparte(renglones)}
    assert res == {("PULPA DE JAIBA 1 KG", "porcionador"), ("CORTE T-BONE", "cocina")}


# ── Router ─────────────────────────────────────────────────────────────


async def _buscar(db, hoja, producto, presentacion=None):
    q = select(ReporteProduccionRenglon).where(ReporteProduccionRenglon.hoja == hoja, ReporteProduccionRenglon.producto == producto)
    if presentacion is not None:
        q = q.where(ReporteProduccionRenglon.presentacion == presentacion)
    return (await db.execute(q)).scalars().first()


async def _buscar_crudo(db, producto):
    """La presentación en KG sin variante de un producto (el crudo que se limpia)."""
    q = select(ReporteProduccionRenglon).where(
        ReporteProduccionRenglon.hoja == "porcionador",
        ReporteProduccionRenglon.producto == producto,
        ReporteProduccionRenglon.unidad == "KG",
        ReporteProduccionRenglon.presentacion.is_(None),
    )
    return (await db.execute(q)).scalars().first()


def _fila(res, renglon_id):
    return next(r for r in res["renglones"] if r["id"] == renglon_id)


@pytest.mark.anyio
async def test_router_dia1_apertura_y_dia2_solo_entradas_y_salidas(db_session):
    assert (await api.sembrar_catalogo(db=db_session, current_user={}))["creados"] == 45 + 32
    assert (await api.sembrar_catalogo(db=db_session, current_user={}))["creados"] == 0  # idempotente
    mignon = await _buscar(db_session, "porcionador", "FILETE DE RES 1 KG", "Mignon")

    dia1 = await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={})
    assert _fila(dia1, mignon.id)["es_apertura"] is True  # nada previo: se captura la apertura

    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, inicial=20, entradas=10, salidas=10)]),
        db=db_session, current_user={},
    )
    f1 = _fila(await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert (f1["inicial"], f1["final_calculado"], f1["es_apertura"]) == (20, 20, True)

    # día 2: sin capturar nada ya se ve el inicial; solo se capturan entradas y salidas
    previa = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert previa["inicial"] == 20 and previa["es_apertura"] is False and previa["capturado"] is False

    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, entradas=5, salidas=3)]),
        db=db_session, current_user={},
    )
    res2 = await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={})
    f2 = _fila(res2, mignon.id)
    assert (f2["inicial"], f2["final_calculado"], f2["existencia_final"]) == (20, 22, 22)
    assert res2["fecha_dia_anterior"] == "2026-10-01"

    # un inicial manual en el día 2 (falla del día anterior) manda sobre el arrastre
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, inicial=30, entradas=5, salidas=3)]),
        db=db_session, current_user={},
    )
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert (f2["inicial"], f2["inicial_manual"], f2["inicial_arrastrado"], f2["existencia_final"]) == (30, 30, 20, 32)
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, entradas=5, salidas=3)]),
        db=db_session, current_user={},
    )

    # corregir una salida del día 1 recalcula el día 2
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, inicial=20, entradas=10, salidas=8)]),
        db=db_session, current_user={},
    )
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert (f2["inicial"], f2["existencia_final"]) == (22, 24)

    # conteo físico: diferencia contra lo calculado
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, entradas=5, salidas=3, conteo_fisico=21)]),
        db=db_session, current_user={},
    )
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert (f2["diferencia"], f2["estado_diferencia"], f2["existencia_final"]) == (-3, "faltante", 21)

    # vaciar todo borra la captura
    r = await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id)]), db=db_session, current_user={}
    )
    assert r == {"guardados": 0, "borrados": 1, "completados": 0}


@pytest.mark.anyio
async def test_router_cocina_recibe_entradas_del_porcionador_y_captura_ventas(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    p = await _buscar(db_session, "porcionador", "FILETE DE RES 1 KG", "Mignon")
    c = await _buscar(db_session, "cocina", "FILETE DE RES 1 KG", "Mignon")
    rib = await _buscar(db_session, "cocina", "CORTE RIB EYE")

    # porcionador manda 10 a cocina el día 1 y 4 el día 2
    for d, salidas in ((D1, 10), (D2, 4)):
        await api.guardar_captura(
            api.CapturaIn(fecha=d, renglones=[api.CapturaRenglonIn(renglon_id=p.id, inicial=50 if d == D1 else None, salidas=salidas)]),
            db=db_session, current_user={},
        )

    # cocina día 1: las entradas ya vienen solas como sugerencia; solo captura ventas
    antes = _fila(await api.obtener_dia(fecha=D1, hoja="cocina", db=db_session, current_user={}), c.id)
    assert antes["entradas"] == 10 and antes["entradas_es_auto"] is True and antes["entradas_auto"] == 10
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[
            api.CapturaRenglonIn(renglon_id=c.id, inicial=0, salidas=7),
            api.CapturaRenglonIn(renglon_id=rib.id, inicial=3, entradas=2, salidas=1),  # sin equivalente: entradas a mano
        ]),
        db=db_session, current_user={},
    )
    res1 = await api.obtener_dia(fecha=D1, hoja="cocina", db=db_session, current_user={})
    f_c, f_rib = _fila(res1, c.id), _fila(res1, rib.id)
    assert (f_c["entradas"], f_c["salidas"], f_c["final_calculado"]) == (10, 7, 3)
    assert f_rib["entradas_es_auto"] is False and f_rib["final_calculado"] == 4

    # día 2: inicial = 3 (final de ayer), entradas = 4 del porcionador aunque cocina aún no guarde ese día
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="cocina", db=db_session, current_user={}), c.id)
    assert (f2["inicial"], f2["entradas"], f2["final_calculado"]) == (3, 4, 7)

    # cocina puede corregir las entradas a mano (llegó distinto a lo que salió): su valor manda
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=c.id, entradas=6, salidas=1)]),
        db=db_session, current_user={},
    )
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="cocina", db=db_session, current_user={}), c.id)
    assert (f2["entradas"], f2["entradas_manual"], f2["entradas_auto"], f2["final_calculado"]) == (6, 6, 4, 8)
    # sin entradas manuales (null) vuelve a la sugerencia del porcionador
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, renglones=[api.CapturaRenglonIn(renglon_id=c.id, salidas=1)]), db=db_session, current_user={}
    )
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="cocina", db=db_session, current_user={}), c.id)
    assert (f2["entradas"], f2["entradas_manual"]) == (4, None)


@pytest.mark.anyio
async def test_router_borrar_dia_solo_borra_ese_dia(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    p = await _buscar(db_session, "porcionador", "FILETE DE RES 1 KG", "Mignon")
    for d in (D1, D2):
        await api.guardar_captura(
            api.CapturaIn(fecha=d, renglones=[api.CapturaRenglonIn(renglon_id=p.id, inicial=5, entradas=1)]),
            db=db_session, current_user={},
        )
    await api.guardar_dia(api.DiaIn(fecha=D1, notas="prueba"), db=db_session, current_user={})

    assert await api.borrar_dia(fecha=D1, db=db_session, current_user={}) == {"borrados": 1}
    assert [f["fecha"] for f in await api.listar_fechas(db=db_session, current_user={})] == ["2026-10-02"]
    assert (await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={}))["dia"]["notas"] is None
    assert len(await api.listar_catalogo(hoja=None, incluir_inactivos=False, db=db_session, current_user={})) == 45 + 32


@pytest.mark.anyio
async def test_router_completar_dia_registra_todos_los_productos_sin_fijar_valores(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    mignon = await _buscar(db_session, "porcionador", "FILETE DE RES 1 KG", "Mignon")
    otro = await _buscar(db_session, "porcionador", "CARNITAS")

    # día 1: apertura de dos productos, y se completa la hoja (45 productos con fila)
    r1 = await api.guardar_captura(
        api.CapturaIn(
            fecha=D1, completar_hoja="porcionador",
            renglones=[
                api.CapturaRenglonIn(renglon_id=mignon.id, inicial=20, entradas=10, salidas=10),
                api.CapturaRenglonIn(renglon_id=otro.id, inicial=7),
            ],
        ),
        db=db_session, current_user={},
    )
    assert r1["completados"] == 45 - 2

    # día 2: solo se edita Mignon; "completar" registra los otros 44, que NO tienen inicial/final fijados
    r2 = await api.guardar_captura(
        api.CapturaIn(fecha=D2, completar_hoja="porcionador", renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, entradas=5, salidas=3)]),
        db=db_session, current_user={},
    )
    assert r2["completados"] == 44
    dia2 = await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={})
    assert all(r["capturado"] for r in dia2["renglones"])  # el día quedó registrado para todos
    f_otro = _fila(dia2, otro.id)
    assert (f_otro["inicial"], f_otro["existencia_final"], f_otro["inicial_manual"], f_otro["conteo_fisico"]) == (7, 7, None, None)
    assert f_otro["entradas_manual"] is None and f_otro["salidas"] is None  # entradas/salidas siguen en blanco
    assert _fila(dia2, mignon.id)["existencia_final"] == 22

    # corregir el día 1 sigue arrastrándose a los productos que no se tocaron en el día 2
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=otro.id, inicial=9)]), db=db_session, current_user={}
    )
    f_otro = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), otro.id)
    assert (f_otro["inicial"], f_otro["existencia_final"]) == (9, 9)

    # vaciar una celda de un producto ya registrado no lo borra del día (se recrea sin movimiento)
    await api.guardar_captura(
        api.CapturaIn(fecha=D2, completar_hoja="porcionador", renglones=[api.CapturaRenglonIn(renglon_id=mignon.id)]),
        db=db_session, current_user={},
    )
    f_mig = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), mignon.id)
    assert f_mig["capturado"] is True and f_mig["entradas_manual"] is None

    # completar es idempotente
    r3 = await api.guardar_captura(api.CapturaIn(fecha=D2, completar_hoja="porcionador", renglones=[]), db=db_session, current_user={})
    assert r3["completados"] == 0


@pytest.mark.anyio
async def test_router_cruce_lista_renglones_sin_contraparte(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    res = await api.obtener_cruce(fecha=D1, db=db_session, current_user={})
    sin = {(r["producto"], r["hoja"]) for r in res["sin_contraparte"]}
    assert ("CORTE RIB EYE", "cocina") in sin and ("PULPA DE JAIBA 1 KG (ALMACEN)", "porcionador") in sin


@pytest.mark.anyio
async def test_router_renglon_inexistente_400(db_session):
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        await api.guardar_captura(
            api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=9999, entradas=1)]),
            db=db_session, current_user={},
        )
    assert exc.value.status_code == 400


# ── Inventario en kg bruto y en pesos (hoja PORCIONADOR) ───────────────────────────────────────────


def _rend(id, producto, presentacion=None, unidad="PORCION", gramaje=None, bruto=None, limpio=None, hoja="porcionador"):
    r = _renglon(id, hoja, producto, presentacion)
    r.unidad, r.gramaje_kg, r.prueba_bruto_kg, r.prueba_limpio_kg = unidad, gramaje, bruto, limpio
    return r


def test_inventario_bruto_reproduce_el_caso_de_la_hoja():
    # CAMARON COCTELERO: 104 porciones de 110 g y 95 de 130 g; pruebas 10→6.27 y 8→5.28; $180/kg
    r1 = _rend(1, "CAMARON COCTELERO", "Cocido", gramaje=0.11, bruto=10, limpio=6.27)
    r2 = _rend(2, "CAMARON COCTELERO", "Fajitas", gramaje=0.13, bruto=8, limpio=5.28)
    inv = calcular_inventario_bruto([r1, r2], {1: 104, 2: 95}, {"CAMARON COCTELERO": 180.0})
    f1 = inv["productos"][0]["filas"][0]
    assert f1["kg_limpio"] == 11.44 and f1["factor"] == pytest.approx(1.5949, abs=1e-3)
    assert f1["kg_bruto"] == pytest.approx(18.246, abs=1e-3)
    assert inv["total_kg_bruto"] == pytest.approx(36.958, abs=2e-3)  # J7 de la hoja
    assert inv["total_valor"] == pytest.approx(6652.39, abs=0.5)  # L7 de la hoja
    assert inv["advertencias"] == []


def test_inventario_bruto_unidad_kg_sin_prueba_ya_es_bruto():
    entero = _rend(1, "PULPA BOLA", "Pieza entera", unidad="KG")
    inv = calcular_inventario_bruto([entero], {1: 5.5}, {"PULPA BOLA": 100.0})
    fila = inv["productos"][0]["filas"][0]
    assert (fila["kg_limpio"], fila["kg_bruto"], fila["factor"]) == (5.5, 5.5, None)
    assert inv["total_valor"] == 550.0


def test_inventario_bruto_bulto_en_kg_con_prueba_si_aplica_factor():
    # sirloin: 8 bultos de 2 kg, prueba 8.1→6.51 → factor 1.244
    bulto = _rend(1, "SIRLOIN", "Bulto", unidad="KG", gramaje=2.0, bruto=8.1, limpio=6.51)
    fila = calcular_inventario_bruto([bulto], {1: 8}, {})["productos"][0]["filas"][0]
    assert fila["kg_limpio"] == 16 and fila["kg_bruto"] == pytest.approx(19.908, abs=1e-3)


def test_inventario_bruto_faltantes_se_advierten_y_no_suman():
    sin_factor = _rend(1, "DORADO", "Ceviche", gramaje=0.15)  # porción sin prueba: no se inventa factor 1
    sin_gramaje = _rend(2, "POLLO", "Parrilla", bruto=1.2, limpio=1.0)
    sin_precio = _rend(3, "JUREL", "Empanizar", gramaje=0.18, bruto=2, limpio=1)
    inv = calcular_inventario_bruto([sin_factor, sin_gramaje, sin_precio], {1: 10, 2: 4, 3: 10}, {"DORADO": 180.0})
    motivos = {f["renglon_id"]: f["falta"] for p in inv["productos"] for f in p["filas"]}
    assert motivos == {1: "sin factor de rendimiento", 2: "sin gramaje", 3: None}
    assert inv["total_kg_bruto"] == pytest.approx(3.6, abs=1e-3)  # solo JUREL suma
    assert inv["total_valor"] == 0.0  # JUREL no tiene precio
    assert any("JUREL" in a and "sin precio" in a for a in inv["advertencias"])
    assert any("DORADO" in a and "sin factor" in a for a in inv["advertencias"])


def test_inventario_bruto_sin_existencia_no_genera_advertencias():
    r = _rend(1, "DORADO", "Ceviche")  # sin datos, pero también sin inventario
    inv = calcular_inventario_bruto([r], {1: None}, {})
    assert inv["advertencias"] == [] and inv["total_kg_bruto"] == 0


def test_inventario_bruto_cocina_hereda_gramaje_y_factor_de_porcionador():
    porc = _rend(1, "FILETE DE RES", "Mignon", gramaje=0.185, bruto=12.76, limpio=9.65)
    coc = _rend(2, "FILETE DE RES", "Mignon", hoja="cocina")
    inv = calcular_inventario_bruto([coc], {2: 10}, {"FILETE DE RES": 240.0}, catalogo_referencia=[porc])
    fila = inv["productos"][0]["filas"][0]
    assert fila["gramaje_kg"] == 0.185 and fila["kg_bruto"] == pytest.approx(10 * 0.185 * 12.76 / 9.65, abs=1e-3)


def test_completar_rendimientos_solo_llena_vacios():
    nuevo = _rend(1, "CAMARON COCTELERO 41/50 1 KG", "Cocido")
    editado = _rend(2, "CAMARON COCTELERO 41/50 1 KG", "Fajitas, ceviche, ajillo, etc.", gramaje=0.5)  # el usuario ya lo cambió
    ajeno = _rend(3, "ALGO NUEVO", None)
    assert completar_rendimientos([nuevo, editado, ajeno]) == 2
    assert (nuevo.gramaje_kg, nuevo.prueba_bruto_kg, nuevo.prueba_limpio_kg) == (0.11, 10.0, 6.27)
    assert (editado.gramaje_kg, editado.prueba_bruto_kg) == (0.5, 8.0)  # conserva su gramaje, rellena lo vacío
    assert ajeno.gramaje_kg is None
    assert completar_rendimientos([nuevo, editado, ajeno]) == 0  # idempotente


def test_semillas_cubren_renglones_del_formato():
    claves = {(f["producto"], f["presentacion"]) for h in ("porcionador", "cocina") for f in filas_semilla(h)}
    assert set(RENDIMIENTOS_SEMILLA) <= claves  # toda semilla apunta a un renglón real del catálogo
    productos = {f["producto"] for h in ("porcionador", "cocina") for f in filas_semilla(h)}
    assert set(PRECIOS_KG_SEMILLA) <= productos


@pytest.mark.anyio
async def test_router_inventario_bruto_precios_y_siembra(db_session):
    res = await api.sembrar_catalogo(db=db_session, current_user={})
    assert res["creados"] == 45 + 32 and res["rendimientos"] > 0 and res["precios"] > 0
    res2 = await api.sembrar_catalogo(db=db_session, current_user={})
    assert (res2["creados"], res2["rendimientos"], res2["precios"]) == (0, 0, 0)  # idempotente

    # el catálogo es editable: un renglón renombrado/desactivado por el usuario NO se vuelve a crear al sembrar
    filetitos = await _buscar(db_session, "porcionador", "PECHUGA DE POLLO 1 KG", "Filetitos")
    filetitos.presentacion = "X MEDIAS PIEZAS"
    filetitos.clave_cruce = construir_clave_cruce(filetitos.producto, filetitos.presentacion)
    await db_session.flush()
    assert (await api.sembrar_catalogo(db=db_session, current_user={}))["creados"] == 0
    assert len(await api.listar_catalogo(hoja="porcionador", incluir_inactivos=True, db=db_session, current_user={})) == 45

    cocido = await _buscar(db_session, "porcionador", "CAMARON COCTELERO 41/50 1 KG", "Cocido")
    fajitas = await _buscar(db_session, "porcionador", "CAMARON COCTELERO 41/50 1 KG", "Fajitas, ceviche, ajillo, etc.")
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[
            api.CapturaRenglonIn(renglon_id=cocido.id, inicial=104),
            api.CapturaRenglonIn(renglon_id=fajitas.id, inicial=95),
        ]),
        db=db_session, current_user={},
    )
    inv = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    coctelero = next(p for p in inv["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")
    assert coctelero["kg_bruto_total"] == pytest.approx(36.958, abs=2e-3)
    assert coctelero["valor"] == pytest.approx(6652.39, abs=0.5)

    # día siguiente sin captura: el inventario es el saldo que arrastra
    inv2 = await api.obtener_inventario_bruto(fecha=D2, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(p for p in inv2["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")["kg_bruto_total"] == coctelero["kg_bruto_total"]

    # editar el precio cambia el valor, y la siembra posterior no lo pisa
    await api.guardar_precio_producto_base(
        api.PrecioProductoIn(producto="CAMARON COCTELERO 41/50 1 KG", precio_kg=200), db=db_session, current_user={}
    )
    inv3 = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(p for p in inv3["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")["precio_kg"] == 200
    await api.sembrar_catalogo(db=db_session, current_user={})
    listado = await api.listar_productos_base(db=db_session, current_user={})
    assert next(p for p in listado if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")["precio_kg"] == 200

    from fastapi import HTTPException

    with pytest.raises(HTTPException):
        await api.guardar_precio_producto_base(api.PrecioProductoIn(producto="X", precio_kg=-1), db=db_session, current_user={})

    # exportar a Excel con fórmulas vivas
    resp = await api.exportar_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert resp.media_type.endswith("spreadsheetml.sheet")


_LIBRO = next(
    (p for p in __import__("glob").glob(
        str(__import__("pathlib").Path(__file__).resolve().parents[2] / "docs" / "reporte*" / "INVENTARIO COCINA*.xlsx")
    )),
    None,
)


@pytest.mark.skipif(_LIBRO is None, reason="no está el libro de inventario de muestra en docs/")
def test_inventario_bruto_coincide_con_la_hoja_porcionador_real():
    """Con las porciones de la hoja y las semillas del sistema, los kg bruto y los pesos coinciden con los de la hoja."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(_LIBRO, data_only=True)["PORCIONADOR"]

    def renglones(producto, filas):
        out = []
        for i, (pres, fila_hoja) in enumerate(filas, start=1):
            g, b, l = RENDIMIENTOS_SEMILLA[(producto, pres)]
            out.append((_rend(i, producto, pres, gramaje=g, bruto=b, limpio=l), float(ws.cell(row=fila_hoja, column=3).value)))
        return out

    casos = [
        # (producto, [(presentación, fila de la hoja con las porciones)], celda del total kg bruto, celda del valor)
        ("CAMARON COCTELERO 41/50 1 KG", [("Cocido", 8), ("Fajitas, ceviche, ajillo, etc.", 9)], "J7", "L7"),
        ("FILETE DE RES 1 KG", [("Mignon", 36), ("Mignoncito", 37), ("Tampiqueña", 38), ("Fajitas", 39)], "J35", "L35"),
        ("PULPA BOLA LA REVUELTA 1 KG", [("Hamburguesa", 45)], "J44", "L44"),
    ]
    for producto, filas, celda_kg, celda_valor in casos:
        rs = renglones(producto, filas)
        inv = calcular_inventario_bruto([r for r, _ in rs], {r.id: c for r, c in rs}, {producto: PRECIOS_KG_SEMILLA[producto]})
        assert inv["total_kg_bruto"] == pytest.approx(float(ws[celda_kg].value), abs=0.01), producto
        assert inv["total_valor"] == pytest.approx(float(ws[celda_valor].value), abs=1.0), producto


# ── Pesaje de limpieza: el factor real sale de lo que se pesa cada día ─────────────────────────────


class _Pesaje:
    def __init__(self, fecha, producto, bruto, limpio):
        self.fecha, self.producto, self.bruto_kg, self.limpio_kg = fecha, producto, bruto, limpio


FILETE = "FILETE DE RES 1 KG"


def test_pesaje_factor_y_merma_del_ejemplo_del_filete():
    ps = [_Pesaje(D1, FILETE, 9.285, 7.550)]
    assert factor_de_pesajes(ps) == pytest.approx(1.2298, abs=1e-4)
    assert merma_kg(9.285, 7.550) == 1.735  # la merma la calcula el sistema (bruto − limpio)


def test_pesaje_varios_del_dia_se_suman_y_sin_datos_no_hay_factor():
    ps = [_Pesaje(D1, FILETE, 9.0, 7.0), _Pesaje(D1, FILETE, 6.0, 5.0)]
    assert factor_de_pesajes(ps) == pytest.approx(15.0 / 12.0)
    assert factor_de_pesajes([]) is None
    assert factor_de_pesajes([_Pesaje(D1, FILETE, 5.0, 0.0)]) is None  # limpio 0: no se inventa


def test_factor_vigente_es_el_del_ultimo_pesaje_sin_mirar_el_futuro():
    ps = [
        _Pesaje(D1, FILETE, 10.0, 8.0),  # 1.25
        _Pesaje(D2, FILETE, 9.0, 6.0),  # 1.5
        _Pesaje(D2, FILETE, 9.0, 6.0),
        _Pesaje(D5, FILETE, 20.0, 10.0),  # 2.0 (futuro respecto a D3)
    ]
    f = factores_por_producto(ps, D3)[FILETE]
    assert f["factor"] == pytest.approx(1.5) and f["fecha"] == "2026-10-02"
    assert (f["bruto_kg"], f["limpio_kg"]) == (18.0, 12.0)  # los dos pesajes del último día
    assert factores_por_producto(ps, D1)[FILETE]["factor"] == pytest.approx(1.25)
    assert factores_por_producto(ps, date(2026, 9, 30)) == {}  # antes del primer pesaje no hay factor


def test_inventario_usa_el_factor_del_pesaje_sobre_la_prueba_del_catalogo():
    mignon = _rend(1, FILETE, "Mignon", gramaje=0.185, bruto=12.76, limpio=9.65)  # prueba vieja: 1.322
    factores = factores_por_producto([_Pesaje(D1, FILETE, 9.285, 7.550)], D1)
    inv = calcular_inventario_bruto([mignon], {1: 21}, {FILETE: 240.0}, factores_pesaje=factores)
    fila = inv["productos"][0]["filas"][0]
    assert fila["kg_limpio"] == pytest.approx(3.885)
    assert fila["factor"] == pytest.approx(1.2298, abs=1e-4)
    assert fila["kg_bruto"] == pytest.approx(4.7778, abs=1e-3)  # 3.885 × 1.2298
    assert (fila["prueba_bruto_kg"], fila["prueba_limpio_kg"]) == (9.285, 7.55)  # el Excel mantiene factor = F/G
    prod = inv["productos"][0]
    assert (prod["factor_origen"], prod["factor_fecha"]) == ("pesaje", "2026-10-01")


def test_inventario_sin_pesaje_cae_a_la_prueba_del_catalogo_y_lo_dice():
    mignon = _rend(1, FILETE, "Mignon", gramaje=0.185, bruto=12.76, limpio=9.65)
    prod = calcular_inventario_bruto([mignon], {1: 21}, {})["productos"][0]
    assert prod["filas"][0]["factor"] == pytest.approx(1.3223, abs=1e-3)
    assert (prod["factor_origen"], prod["factor_fecha"]) == ("catalogo", None)


def test_inventario_pieza_en_kg_sin_prueba_sigue_siendo_bruto_aunque_haya_pesaje():
    entero = _rend(1, FILETE, None, unidad="KG")
    mignon = _rend(2, FILETE, "Mignon", gramaje=0.185)  # sin prueba propia: depende del pesaje
    factores = factores_por_producto([_Pesaje(D1, FILETE, 9.285, 7.550)], D1)
    inv = calcular_inventario_bruto([entero, mignon], {1: 5, 2: 10}, {}, factores_pesaje=factores)
    f_entero, f_mignon = inv["productos"][0]["filas"]
    assert (f_entero["kg_bruto"], f_entero["factor"]) == (5, None)  # ya es bruto: no se le aplica factor
    assert f_mignon["kg_bruto"] == pytest.approx(10 * 0.185 * 9.285 / 7.55, abs=1e-3)


def _control(pesajes, entradas):
    rs = [
        _rend(1, FILETE, "Mignon", gramaje=0.185),
        _rend(2, FILETE, "Mignoncito", gramaje=0.13),
        _rend(3, FILETE, "Tampiqueña", gramaje=0.135),
        _rend(4, FILETE, "Fajitas", gramaje=0.15),
        _rend(5, FILETE, None, unidad="KG"),  # recepción en KG: no es producción
        _rend(6, FILETE, "Filete res corte"),  # sin gramaje
    ]
    return calcular_control_limpieza(pesajes, rs, entradas)


def test_control_cuadra_dentro_de_la_tolerancia():
    # 15·0.185 + 19·0.13 + 9·0.135 + 9·0.15 = 7.81 kg de porciones contra 7.55 limpios (−3.4 %)
    c = _control([_Pesaje(D1, FILETE, 9.285, 7.550)], {1: 15, 2: 19, 3: 9, 4: 9, 5: 4})
    assert c["kg_porciones"] == pytest.approx(7.81) and c["merma_kg"] == 1.735
    assert c["diferencia_kg"] == pytest.approx(-0.26) and c["estado"] == "cuadra"
    assert c["factor"] == pytest.approx(1.2298, abs=1e-4)
    assert len(c["porciones"]) == 4  # la recepción en KG no cuenta como porciones


def test_control_marca_faltan_porciones_y_exceden_limpio():
    # solo 10 mignon (1.85 kg) de 7.55 kg limpios: faltan porciones
    faltan = _control([_Pesaje(D1, FILETE, 9.285, 7.550)], {1: 10})
    assert faltan["estado"] == "faltan_porciones" and faltan["diferencia_kg"] == pytest.approx(5.7)
    # 60 mignon (11.1 kg) con solo 7.55 kg limpios: no alcanzan
    exceden = _control([_Pesaje(D1, FILETE, 9.285, 7.550)], {1: 60})
    assert exceden["estado"] == "exceden_limpio" and exceden["diferencia_pct"] < -5


def test_control_sin_pesaje_y_renglones_sin_gramaje_se_avisan():
    c = _control([], {1: 10, 6: 4})
    assert c["estado"] == "sin_pesaje" and c["diferencia_kg"] is None and c["merma_kg"] is None
    assert c["sin_gramaje"] == ["Filete res corte"]  # produjo algo sin gramaje: no se cuenta, se avisa


@pytest.mark.anyio
async def test_router_pesajes_crud_validaciones_y_control(db_session):
    from fastapi import HTTPException

    await api.sembrar_catalogo(db=db_session, current_user={})
    mignon = await _buscar(db_session, "porcionador", FILETE, "Mignon")
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=mignon.id, inicial=0, entradas=21)]),
        db=db_session, current_user={},
    )

    # validaciones
    for bruto, limpio in ((0, 1), (5, 0), (5, 6)):
        with pytest.raises(HTTPException) as exc:
            await api.crear_pesaje(api.PesajeIn(fecha=D1, producto=FILETE, bruto_kg=bruto, limpio_kg=limpio), db=db_session, current_user={})
        assert exc.value.status_code == 400
    with pytest.raises(HTTPException):
        await api.crear_pesaje(api.PesajeIn(fecha=D1, producto="NO EXISTE", bruto_kg=2, limpio_kg=1), db=db_session, current_user={})

    creado = await api.crear_pesaje(
        api.PesajeIn(fecha=D1, producto=FILETE, bruto_kg=9.285, limpio_kg=7.55), db=db_session, current_user={}
    )
    assert creado["merma_kg"] == 1.735 and creado["factor"] == pytest.approx(1.2298, abs=1e-4)

    dia = await api.listar_pesajes(fecha=D1, db=db_session, current_user={})
    filete = next(p for p in dia["productos"] if p["producto"] == FILETE)
    assert filete["activo_hoy"] is True and len(filete["pesajes"]) == 1
    assert filete["control"]["kg_porciones"] == pytest.approx(21 * 0.185)  # sembrado con gramaje 0.185
    assert all(p["producto"] != "PULPA DE JAIBA 1 KG (ALMACEN)" for p in dia["productos"])  # solo KG: no se pesa

    # el inventario usa ese factor y dice de dónde sale
    inv = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    prod = next(p for p in inv["productos"] if p["producto"] == FILETE)
    assert (prod["factor_origen"], prod["factor_fecha"]) == ("pesaje", "2026-10-01")
    fila = next(f for f in prod["filas"] if f["renglon_id"] == mignon.id)
    assert fila["factor"] == pytest.approx(1.2298, abs=1e-4) and fila["kg_bruto"] == pytest.approx(21 * 0.185 * 9.285 / 7.55, abs=1e-3)

    # un pesaje más reciente cambia el factor del día siguiente; uno futuro no afecta a un día anterior
    await api.crear_pesaje(api.PesajeIn(fecha=D2, producto=FILETE, bruto_kg=10, limpio_kg=5), db=db_session, current_user={})
    inv2 = await api.obtener_inventario_bruto(fecha=D2, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(f for f in next(p for p in inv2["productos"] if p["producto"] == FILETE)["filas"] if f["renglon_id"] == mignon.id)["factor"] == 2.0
    inv_d1 = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(p for p in inv_d1["productos"] if p["producto"] == FILETE)["factor_fecha"] == "2026-10-01"

    # editar y borrar
    editado = await api.editar_pesaje(creado["id"], api.PesajePatch(limpio_kg=8.0), db=db_session, current_user={})
    assert editado["merma_kg"] == pytest.approx(1.285)
    with pytest.raises(HTTPException):
        await api.editar_pesaje(creado["id"], api.PesajePatch(limpio_kg=99), db=db_session, current_user={})
    assert (await api.borrar_pesaje(creado["id"], db=db_session, current_user={})) == {"ok": True}
    with pytest.raises(HTTPException) as exc:
        await api.borrar_pesaje(creado["id"], db=db_session, current_user={})
    assert exc.value.status_code == 404
    # sin pesaje de ese día vuelve a la prueba del catálogo
    inv3 = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(p for p in inv3["productos"] if p["producto"] == FILETE)["factor_origen"] == "catalogo"


# ── Crudo que se descuenta con el pesaje y base "producción" ───────────────────────────────────────


def _crudo(id, producto, presentacion=None, unidad="KG", es_crudo=None, activo=True):
    r = _renglon(id, "porcionador", producto, presentacion, activo=activo)
    r.unidad, r.es_crudo = unidad, es_crudo
    r.gramaje_kg = r.prueba_bruto_kg = r.prueba_limpio_kg = None
    return r


def test_crudo_es_la_unica_presentacion_en_kg_del_producto():
    kg = _crudo(1, FILETE)
    mignon = _crudo(2, FILETE, "Mignon", unidad="PORCION")
    assert resolver_crudos([kg, mignon]) == {1}


def test_crudo_con_varias_filas_en_kg_no_se_elige_solo_salvo_marca_explicita():
    a, b = _crudo(1, "POLLO", "Caja KG"), _crudo(2, "POLLO", None)
    assert resolver_crudos([a, b]) == set()  # ambiguo: no se adivina
    a.es_crudo = True
    assert resolver_crudos([a, b]) == {1}  # el usuario lo marcó
    b.es_crudo = True
    assert resolver_crudos([a, b]) == {1, 2}


def test_crudo_marcado_false_se_descarta_y_otro_candidato_unico_pasa_a_serlo():
    a, b = _crudo(1, "POLLO", "Caja KG", es_crudo=False), _crudo(2, "POLLO", None)
    assert resolver_crudos([a, b]) == {2}
    assert resolver_crudos([_crudo(3, "SOLO", None, es_crudo=False)]) == set()


def test_crudo_no_aplica_a_cocina_ni_a_inactivos():
    cocina = _crudo(1, FILETE)
    cocina.hoja = "cocina"
    assert resolver_crudos([cocina]) == set()
    assert resolver_crudos([_crudo(2, FILETE, activo=False)]) == set()


def test_salidas_auto_crudo_suma_el_bruto_pesado_por_dia():
    kg = _crudo(1, FILETE)
    ps = [_Pesaje(D1, FILETE, 9.0, 7.0), _Pesaje(D1, FILETE, 1.5, 1.2), _Pesaje(D2, FILETE, 4.0, 3.0), _Pesaje(D1, "OTRO", 5, 4)]
    assert salidas_auto_crudo([kg], ps) == {(1, D1): 10.5, (1, D2): 4.0}


def test_salidas_auto_crudo_con_dos_crudos_marcados_no_cuenta_doble():
    a, b = _crudo(1, "POLLO", "Caja KG", es_crudo=True), _crudo(2, "POLLO", None, es_crudo=True)
    assert salidas_auto_crudo([a, b], [_Pesaje(D1, "POLLO", 8.0, 6.0)]) == {(1, D1): 8.0}


@pytest.mark.anyio
async def test_router_pesaje_descuenta_el_crudo_y_la_edicion_manual_manda(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    crudo = await _buscar_crudo(db_session, FILETE)
    assert crudo.unidad == "KG"

    # el caso real: crudo 4 + 4 = 8 kg; se pesan 9.45 kg bruto para limpiar
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=crudo.id, inicial=4, entradas=4)]),
        db=db_session, current_user={},
    )
    antes = _fila(await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={}), crudo.id)
    assert antes["salidas_es_auto"] is True and antes["salidas_auto"] is None and antes["final_calculado"] == 8

    await api.crear_pesaje(api.PesajeIn(fecha=D1, producto=FILETE, bruto_kg=9.45, limpio_kg=7.55), db=db_session, current_user={})
    f = _fila(await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={}), crudo.id)
    assert f["salidas_auto"] == 9.45 and f["salidas"] is None  # lo capturado a mano sigue vacío
    assert f["final_calculado"] == pytest.approx(-1.45) and f["alerta_negativo"] is True  # pesó más de lo que había
    assert f["diferencia"] is None  # ya no hay un "Ajuste −8" falso

    # una salida capturada a mano manda sobre la automática
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=crudo.id, inicial=4, entradas=4, salidas=8)]),
        db=db_session, current_user={},
    )
    f = _fila(await api.obtener_dia(fecha=D1, hoja="porcionador", db=db_session, current_user={}), crudo.id)
    assert (f["salidas"], f["final_calculado"], f["alerta_negativo"]) == (8, 0, False)

    # el saldo del crudo sigue al día siguiente
    f2 = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), crudo.id)
    assert f2["inicial"] == 0


@pytest.mark.anyio
async def test_router_pesaje_descuenta_aunque_el_crudo_no_se_haya_guardado_ese_dia(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    crudo = await _buscar_crudo(db_session, FILETE)
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=crudo.id, inicial=20)]), db=db_session, current_user={}
    )
    await api.crear_pesaje(api.PesajeIn(fecha=D2, producto=FILETE, bruto_kg=9, limpio_kg=7), db=db_session, current_user={})
    f = _fila(await api.obtener_dia(fecha=D2, hoja="porcionador", db=db_session, current_user={}), crudo.id)
    assert (f["inicial"], f["salidas_auto"], f["existencia_final"]) == (20, 9, 11)


@pytest.mark.anyio
async def test_router_inventario_base_produccion_solo_usa_entradas_procesadas(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    mignon = await _buscar(db_session, "porcionador", FILETE, "Mignon")
    crudo = await _buscar_crudo(db_session, FILETE)
    camaron = await _buscar(db_session, "porcionador", "CAMARON COCTELERO 41/50 1 KG", "Cocido")
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[
            api.CapturaRenglonIn(renglon_id=mignon.id, inicial=50, entradas=16),  # inventario grande, producción 16
            api.CapturaRenglonIn(renglon_id=crudo.id, inicial=4, entradas=4),  # recepción de crudo: NO es producción
            api.CapturaRenglonIn(renglon_id=camaron.id, inicial=104),  # solo existencia, sin producción hoy
        ]),
        db=db_session, current_user={},
    )
    await api.crear_pesaje(api.PesajeIn(fecha=D1, producto=FILETE, bruto_kg=9.45, limpio_kg=7.55), db=db_session, current_user={})

    prod = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="produccion", db=db_session, current_user={})
    assert prod["base"] == "produccion"
    filete = next(p for p in prod["productos"] if p["producto"] == FILETE)
    fila_mignon = next(f for f in filete["filas"] if f["renglon_id"] == mignon.id)
    assert fila_mignon["cantidad"] == 16  # solo la producción, no las 66 que hay
    assert fila_mignon["kg_limpio"] == pytest.approx(16 * 0.185)
    assert fila_mignon["kg_bruto"] == pytest.approx(16 * 0.185 * 9.45 / 7.55, abs=1e-3)
    assert next(f for f in filete["filas"] if f["renglon_id"] == crudo.id)["kg_bruto"] is None  # el crudo no cuenta como producción
    assert filete["kg_bruto_total"] == pytest.approx(16 * 0.185 * 9.45 / 7.55, abs=1e-3)
    coctelero = next(p for p in prod["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")
    assert coctelero["kg_bruto_total"] == 0  # sin producción: no se calcula factor ni valor

    # la base "existencia" sigue calculando todo lo que hay
    exis = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    assert next(p for p in exis["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")["kg_bruto_total"] > 0

    from fastapi import HTTPException

    with pytest.raises(HTTPException):
        await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="otra", db=db_session, current_user={})
    resp = await api.exportar_inventario_bruto(fecha=D1, hoja="porcionador", base="produccion", db=db_session, current_user={})
    assert "produccion_bruto_porcionador" in resp.headers["content-disposition"]


@pytest.mark.anyio
async def test_router_crear_renglon_guarda_rendimiento_y_es_crudo(db_session):
    nuevo = await api.crear_renglon(
        api.RenglonIn(
            hoja="porcionador", seccion="RES", producto="LOMO", presentacion="Medallón", unidad="PORCION",
            gramaje_kg=0.2, prueba_bruto_kg=10, prueba_limpio_kg=8, es_crudo=False,
        ),
        db=db_session, current_user={},
    )
    assert (nuevo["gramaje_kg"], nuevo["prueba_bruto_kg"], nuevo["prueba_limpio_kg"], nuevo["es_crudo"]) == (0.2, 10, 8, False)
    editado = await api.editar_renglon(nuevo["id"], api.RenglonPatch(es_crudo=None), db=db_session, current_user={})
    assert editado["es_crudo"] is None  # vuelve a la regla automática


# ── Existencia total por grupo (sección) ───────────────────────────────────────────────────────────


def _con_seccion(r, seccion):
    r.seccion = seccion
    return r


def test_inventario_incluye_la_seccion_de_cada_producto_y_agrupa_en_orden_del_catalogo():
    c1 = _con_seccion(_rend(1, "CAMARON COCTELERO", "Cocido", gramaje=0.11, bruto=10, limpio=6.27), "CAMARÓN")
    c2 = _con_seccion(_rend(2, "CAMARON U-15", "P/Imperiales", gramaje=0.14, bruto=6, limpio=5.74), "CAMARÓN")
    r1 = _con_seccion(_rend(3, FILETE, "Mignon", gramaje=0.185, bruto=12.76, limpio=9.65), "RES")
    inv = calcular_inventario_bruto(
        [c1, c2, r1], {1: 78, 2: 25, 3: 10}, {"CAMARON COCTELERO": 180.0, "CAMARON U-15": 280.0, FILETE: 240.0}
    )
    assert [p["seccion"] for p in inv["productos"]] == ["CAMARÓN", "CAMARÓN", "RES"]
    assert [g["seccion"] for g in inv["grupos"]] == ["CAMARÓN", "RES"]  # orden del catálogo
    camaron = inv["grupos"][0]
    assert [p["producto"] for p in camaron["productos"]] == ["CAMARON COCTELERO", "CAMARON U-15"]
    # la suma de los grupos coincide con los totales generales
    assert sum(g["kg_bruto_total"] for g in inv["grupos"]) == pytest.approx(inv["total_kg_bruto"], abs=1e-3)
    assert sum(g["valor"] for g in inv["grupos"]) == pytest.approx(inv["total_valor"], abs=0.02)
    # camarón coctelero: 78 × 0.11 = 8.58 kg limpio × (10/6.27) = 13.68 kg bruto × $180
    coctelero = camaron["productos"][0]
    assert coctelero["kg_bruto_total"] == pytest.approx(78 * 0.11 * 10 / 6.27, abs=1e-3)
    assert coctelero["valor"] == pytest.approx(78 * 0.11 * 10 / 6.27 * 180, abs=0.05)


def test_grupo_con_producto_sin_precio_lo_avisa_y_no_suma_su_valor():
    a = _con_seccion(_rend(1, "DORADO", "Empanizar", gramaje=0.18, bruto=11.06, limpio=8.35), "PESCADO")
    b = _con_seccion(_rend(2, "JUREL", "Empanizar", gramaje=0.18, bruto=2, limpio=1), "PESCADO")
    inv = calcular_inventario_bruto([a, b], {1: 10, 2: 10}, {"DORADO": 180.0})  # JUREL sin precio
    g = inv["grupos"][0]
    assert g["sin_precio"] is True
    assert g["valor"] == pytest.approx(inv["productos"][0]["valor"])  # solo suma el que tiene precio
    assert g["kg_bruto_total"] > inv["productos"][0]["kg_bruto_total"]  # pero el bruto sí cuenta los dos


def test_agrupar_por_seccion_sin_productos():
    assert agrupar_por_seccion([]) == []


@pytest.mark.anyio
async def test_router_existencia_total_por_grupo_y_excel_con_subtotales(db_session):
    from io import BytesIO

    import openpyxl

    await api.sembrar_catalogo(db=db_session, current_user={})
    cocido = await _buscar(db_session, "porcionador", "CAMARON COCTELERO 41/50 1 KG", "Cocido")
    fajitas = await _buscar(db_session, "porcionador", "CAMARON COCTELERO 41/50 1 KG", "Fajitas, ceviche, ajillo, etc.")
    mignon = await _buscar(db_session, "porcionador", FILETE, "Mignon")
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[
            api.CapturaRenglonIn(renglon_id=cocido.id, inicial=78),
            api.CapturaRenglonIn(renglon_id=fajitas.id, inicial=70),
            api.CapturaRenglonIn(renglon_id=mignon.id, inicial=21),
        ]),
        db=db_session, current_user={},
    )
    inv = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    nombres = [g["seccion"] for g in inv["grupos"]]
    assert nombres[0] == "CAMARÓN" and "RES" in nombres and len(nombres) == len(set(nombres))
    camaron = next(g for g in inv["grupos"] if g["seccion"] == "CAMARÓN")
    coctelero = next(p for p in camaron["productos"] if p["producto"] == "CAMARON COCTELERO 41/50 1 KG")
    esperado_kg = (78 * 0.11) * (10 / 6.27) + (70 * 0.13) * (8 / 5.28)  # el ejemplo de la hoja
    assert coctelero["kg_bruto_total"] == pytest.approx(esperado_kg, abs=1e-3)
    assert coctelero["valor"] == pytest.approx(esperado_kg * 180, abs=0.1)  # $180/kg sembrado
    assert sum(g["kg_bruto_total"] for g in inv["grupos"]) == pytest.approx(inv["total_kg_bruto"], abs=1e-3)
    assert sum(g["valor"] for g in inv["grupos"]) == pytest.approx(inv["total_valor"], abs=0.05)

    # el Excel trae una fila por grupo con su subtotal (fórmula) y el total general suma los subtotales
    resp = await api.exportar_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    cuerpo = b"".join([chunk async for chunk in resp.body_iterator])
    ws = openpyxl.load_workbook(BytesIO(cuerpo)).active
    filas_grupo = {str(ws.cell(row=r, column=1).value): r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value in nombres}
    assert set(filas_grupo) == set(nombres)
    sub = str(ws.cell(row=filas_grupo["CAMARÓN"], column=9).value)
    assert sub.startswith("=I") and "+" in sub  # suma de los productos del grupo
    total = str(ws.cell(row=ws.max_row, column=9).value)
    assert total.count("+") == len(nombres) - 1 and all(f"I{r}" in total for r in filas_grupo.values())


# ── Gramaje: el número en kg, o el que dice el texto impreso ──────────────────────────────────────


@pytest.mark.parametrize(
    "texto,esperado",
    [
        ("180g", 0.18), ("150 G", 0.15), ("440g", 0.44), ("200 gr", 0.2), ("2.00 KG", 2.0), ("1 kg", 1.0), ("0,5 kg", 0.5),
        ("130-135g", None),  # rango: no se adivina
        ("KG", None), ("2", None), ("", None), (None, None), ("abc", None),
        ("1.300 G", None),  # punto de miles ("1,300 g"): menos de 10 g no es una porción
    ],
)
def test_gramaje_desde_texto(texto, esperado):
    assert gramaje_desde_texto(texto) == esperado


def _con_texto(id, producto, presentacion, unidad, texto, gramaje_kg=None):
    r = _rend(id, producto, presentacion, unidad=unidad, gramaje=gramaje_kg)
    r.gramaje = texto
    return r


def test_gramaje_efectivo_usa_el_numero_primero_luego_el_texto_y_nunca_en_kg():
    assert gramaje_efectivo_kg(_con_texto(1, "X", "a", "PORCION", "180g")) == 0.18  # solo texto
    assert gramaje_efectivo_kg(_con_texto(2, "X", "a", "PORCION", "180g", gramaje_kg=0.185)) == 0.185  # el número manda
    assert gramaje_efectivo_kg(_con_texto(3, "X", "a", "1/2 PZA", "440g")) == 0.44
    # en KG la cantidad ya es peso: "2 KG" es el tamaño del paquete y no se multiplica
    assert gramaje_efectivo_kg(_con_texto(4, "X", "Enmaquetado", "KG", "2 KG")) is None
    assert gramaje_efectivo_kg(_con_texto(5, "X", "Bulto", "KG", "2 KG", gramaje_kg=2.0)) == 2.0  # salvo número explícito


def test_inventario_toma_el_gramaje_del_texto_cuando_falta_el_numero():
    jurel = _con_texto(1, "JUREL", "Empanizar", "PORCION", "180g")
    jurel.prueba_bruto_kg, jurel.prueba_limpio_kg = 2, 1
    inv = calcular_inventario_bruto([jurel], {1: 50}, {"JUREL": 100.0})
    fila = inv["productos"][0]["filas"][0]
    assert (fila["gramaje_kg"], fila["falta"]) == (0.18, None)
    assert fila["kg_limpio"] == pytest.approx(9.0) and fila["kg_bruto"] == pytest.approx(18.0)
    assert inv["advertencias"] == [] and inv["total_valor"] == pytest.approx(1800.0)


def test_inventario_en_kg_no_multiplica_por_el_texto_del_paquete():
    envase = _con_texto(1, "CAMARON", "Enmaquetado", "KG", "2 KG")
    fila = calcular_inventario_bruto([envase], {1: 5}, {})["productos"][0]["filas"][0]
    assert (fila["gramaje_kg"], fila["kg_bruto"]) == (1.0, 5)  # 5 kg son 5 kg, no 10


def test_control_de_limpieza_tambien_lee_el_gramaje_del_texto():
    r = _con_texto(1, FILETE, "Filete res corte", "PORCION", "400g")
    c = calcular_control_limpieza([_Pesaje(D1, FILETE, 9.285, 7.55)], [r], {1: 6})
    assert c["sin_gramaje"] == [] and c["kg_porciones"] == pytest.approx(2.4)


@pytest.mark.anyio
async def test_router_catalogo_informa_el_gramaje_efectivo(db_session):
    nuevo = await api.crear_renglon(
        api.RenglonIn(hoja="porcionador", seccion="RES", producto="LOMO", presentacion="Medallón", gramaje="200g", unidad="PORCION"),
        db=db_session, current_user={},
    )
    assert nuevo["gramaje_kg"] is None and nuevo["gramaje_kg_efectivo"] == 0.2


# ── Limpio intermedio (filetitos): pollo ya limpio que entra y sale ───────────────────────────────

PECHUGA = "PECHUGA DE POLLO 1 KG"


def _pechuga():
    dedos = _rend(1, PECHUGA, "Dedos de pollo", gramaje=0.25, bruto=26.33, limpio=12.14)
    fajitas = _rend(2, PECHUGA, "Fajitas", gramaje=0.13, bruto=26.33, limpio=12.14)
    filetitos = _crudo(3, PECHUGA, "FILETITOS")  # KG, sin gramaje ni prueba propios
    filetitos.es_limpio = True
    return dedos, fajitas, filetitos


def test_limpio_intermedio_en_existencia_se_pasa_a_bruto_con_el_factor_del_producto():
    dedos, fajitas, filetitos = _pechuga()
    inv = calcular_inventario_bruto([dedos, fajitas, filetitos], {1: 0, 2: 0, 3: 3.07}, {PECHUGA: 81.0})
    fila = inv["productos"][0]["filas"][2]
    # 3.07 kg limpios × (26.33 / 12.14) = 6.66 kg bruto: NO se toman 1 a 1 como si fueran crudos
    assert fila["limpio_intermedio"] is True and fila["kg_limpio"] == pytest.approx(3.07)
    assert fila["kg_bruto"] == pytest.approx(3.07 * 26.33 / 12.14, abs=1e-3)


def test_limpio_intermedio_usa_el_factor_del_pesaje_si_lo_hay():
    dedos, fajitas, filetitos = _pechuga()
    factores = factores_por_producto([_Pesaje(D1, PECHUGA, 20.0, 10.0)], D1)
    inv = calcular_inventario_bruto([dedos, fajitas, filetitos], {1: 0, 2: 0, 3: 3.07}, {}, factores_pesaje=factores)
    fila = inv["productos"][0]["filas"][2]
    assert fila["factor"] == 2.0 and fila["kg_bruto"] == pytest.approx(6.14)


def test_filetitos_sin_ningun_factor_avisa_y_no_se_toma_como_bruto():
    filetitos = _crudo(1, "POLLO SIN PRUEBA", "FILETITOS")
    filetitos.es_limpio = True
    inv = calcular_inventario_bruto([filetitos], {1: 3.0}, {})
    assert inv["productos"][0]["filas"][0]["falta"] == "sin factor de rendimiento"
    assert inv["total_kg_bruto"] == 0


def test_produccion_con_filetitos_entra_y_sale_no_cuenta_dos_veces():
    dedos, fajitas, filetitos = _pechuga()
    # hoy: 14 dedos (3.5 kg) hechos con los 3.07 kg de filetitos que sobraron hoy: entradas 3.07 y salidas 3.07
    neto = 3.07 - 3.07
    inv = calcular_inventario_bruto([dedos, fajitas, filetitos], {1: 14, 2: 0, 3: neto}, {PECHUGA: 81.0})
    assert inv["productos"][0]["kg_bruto_total"] == pytest.approx(14 * 0.25 * 26.33 / 12.14, abs=1e-3)


def test_produccion_filetitos_usado_de_dias_anteriores_se_descuenta_del_bruto():
    dedos, fajitas, filetitos = _pechuga()
    # los 14 dedos (3.5 kg) salieron de filetitos de ayer (salidas 3.07, sin entradas hoy): solo 0.43 kg son nuevos
    inv = calcular_inventario_bruto([dedos, fajitas, filetitos], {1: 14, 2: 0, 3: -3.07}, {PECHUGA: 81.0})
    assert inv["productos"][0]["kg_bruto_total"] == pytest.approx((3.5 - 3.07) * 26.33 / 12.14, abs=1e-3)
    assert inv["productos"][0]["filas"][2]["kg_bruto"] < 0  # el neto negativo se ve


def test_control_suma_el_limpio_intermedio_neto_al_de_las_porciones():
    dedos, fajitas, filetitos = _pechuga()
    pesajes = [_Pesaje(D1, PECHUGA, 26.33, 12.14)]
    # porciones directas: 20 dedos (5.0) + 15 fajitas (1.95); sobraron 3.07 kg de filetitos y no se usaron: 5.0+1.95+3.07 = 10.02
    c = calcular_control_limpieza(pesajes, [dedos, fajitas, filetitos], {1: 20, 2: 15, 3: 3.07}, {})
    assert c["kg_porciones"] == pytest.approx(10.02)
    assert any(p.get("limpio_intermedio") for p in c["porciones"])
    # si ese sobrante se usó para dedos el mismo día, entra y sale: no suma
    c2 = calcular_control_limpieza(pesajes, [dedos, fajitas, filetitos], {1: 20, 2: 15, 3: 3.07}, {3: 3.07})
    assert c2["kg_porciones"] == pytest.approx(5.0 + 1.95)
    assert not any(p.get("limpio_intermedio") for p in c2["porciones"])


def test_el_sobrante_limpio_nunca_se_elige_como_crudo():
    caja = _crudo(1, PECHUGA, "Caja KG")
    filetitos = _crudo(2, PECHUGA, "FILETITOS")
    filetitos.es_limpio = True
    assert resolver_crudos([caja, filetitos]) == {1}  # el filetitos no cuenta como candidato: la caja es el único crudo


@pytest.mark.anyio
async def test_router_filetitos_marcado_limpio_netea_la_produccion(db_session):
    await api.sembrar_catalogo(db=db_session, current_user={})
    dedos = await _buscar(db_session, "porcionador", PECHUGA, "Dedos de pollo")
    filetitos = await _buscar(db_session, "porcionador", PECHUGA, "Filetitos")
    ed = await api.editar_renglon(filetitos.id, api.RenglonPatch(unidad="KG", es_limpio=True), db=db_session, current_user={})
    assert ed["es_limpio"] is True
    await api.crear_pesaje(api.PesajeIn(fecha=D1, producto=PECHUGA, bruto_kg=26.33, limpio_kg=12.14), db=db_session, current_user={})
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[
            api.CapturaRenglonIn(renglon_id=dedos.id, entradas=14),
            api.CapturaRenglonIn(renglon_id=filetitos.id, entradas=3.07, salidas=3.07),
        ]),
        db=db_session, current_user={},
    )
    prod = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="produccion", db=db_session, current_user={})
    pech = next(p for p in prod["productos"] if p["producto"] == PECHUGA)
    assert pech["kg_bruto_total"] == pytest.approx(14 * 0.25 * 26.33 / 12.14, abs=1e-3)  # entra y sale: solo cuentan los dedos
    dia = await api.listar_pesajes(fecha=D1, db=db_session, current_user={})
    control = next(p for p in dia["productos"] if p["producto"] == PECHUGA)["control"]
    assert control["kg_porciones"] == pytest.approx(3.5)

    # si el sobrante NO se usó (queda en existencia), cuenta como producido y en existencia se pasa a bruto
    await api.guardar_captura(
        api.CapturaIn(fecha=D1, renglones=[api.CapturaRenglonIn(renglon_id=filetitos.id, entradas=3.07)]), db=db_session, current_user={}
    )
    prod2 = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="produccion", db=db_session, current_user={})
    assert next(p for p in prod2["productos"] if p["producto"] == PECHUGA)["kg_bruto_total"] == pytest.approx(
        (3.5 + 3.07) * 26.33 / 12.14, abs=1e-3
    )
    exis = await api.obtener_inventario_bruto(fecha=D1, hoja="porcionador", base="existencia", db=db_session, current_user={})
    fila = next(f for p in exis["productos"] if p["producto"] == PECHUGA for f in p["filas"] if f["renglon_id"] == filetitos.id)
    assert fila["kg_bruto"] == pytest.approx(3.07 * 26.33 / 12.14, abs=1e-3)
