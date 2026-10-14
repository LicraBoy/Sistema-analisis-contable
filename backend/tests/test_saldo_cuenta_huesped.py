"""
Tests de Saldo Cuenta Huésped: parseo del RCP170 de Novohit (sintético y, si
existe la carpeta `saldo cuenta huesped/`, contra el archivo real de
septiembre 2026), filtro de garantía, fórmulas del desglose y round-trip del
router (subir → capturar IVA/propina → re-subir conservando lo capturado).
"""

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.db.session import get_db
from app.main import app
from app.services.saldo_cuenta_huesped import (
    aplanar_lineas,
    calcular_desglose,
    es_excel_editado,
    filtrar_con_garantia,
    generar_excel_saldo_cuenta_huesped,
    leer_excel_editado,
    parsear_rcp170,
)

RCP170_REAL = Path(__file__).resolve().parents[2] / "saldo cuenta huesped" / "RCP170.xls"


def _rcp170_sintetico() -> bytes:
    """Mismo layout de columnas (B=1, J=9, L=11 ...) que el reporte real."""
    wb = Workbook()
    ws = wb.active

    def poner(fila: int, **celdas):
        for col, valor in celdas.items():
            ws.cell(row=fila, column=int(col[1:]) + 1, value=valor)

    poner(6, c1="Hab.", c9="In", c11="Huésped", c32="Garantía", c37="Saldo")
    poner(8, c1="B203", c9="ORXU9", c11="GALLEGOS, BERTHA", c16="2 / 0", c20="12/09/2026", c23="26/09/2026",
          c25="0.00", c32="CORPORATIVO", c37="6,780.40")
    poner(9, c7="Agencia", c11="GALLEGOS BERTHA", c21="Principal", c32="6,780.40")
    poner(11, c1="B406", c9="QHFQU", c11="COBLE, JILL*1566*", c32="CORPORATIVO", c37="11,816.16")
    poner(12, c7="Agencia", c11="COBLE JILL", c21="Principal", c32="8,151.06")
    poner(13, c21="BATTS JOAN*1567*", c32="3,665.10")
    poner(15, c1="A505", c9="WkIn", c11="ESCOBAR, OMAR", c32=None, c37="0.00")  # sin garantía
    poner(16, c7="Agencia", c11="TORRES MAZATLAN", c21="Principal", c32="0.00")
    poner(18, c1="#Habitaciones", c13=3.0, c26="Saldo de Habitaciones", c34="18,596.56")
    poner(20, c1="623", c9="Varios", c11="IGNACIO", c20="05/09/2026", c23="23/09/2026", c37="5,135.40")
    poner(22, c1="#Folios de Varios", c13=1.0, c26="Saldo de Folios de Varios", c34="5,135.40")
    poner(24, c26="Saldo Total", c34="23,731.96")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parsear_rcp170_sintetico():
    r = parsear_rcp170(_rcp170_sintetico())
    assert [h["habitacion"] for h in r["habitaciones"]] == ["B203", "B406", "A505"]
    assert [h["habitacion"] for h in r["varios"]] == ["623"]
    assert r["saldo_habitaciones"] == 18596.56 and r["saldo_varios"] == 5135.40 and r["saldo_total"] == 23731.96
    b406 = r["habitaciones"][1]
    assert b406["saldo"] == 11816.16
    assert [(l["linea"], l["monto"]) for l in b406["lineas"]] == [("Principal", 8151.06), ("BATTS JOAN*1567*", 3665.10)]
    assert b406["lineas"][0]["agencia"] == "COBLE JILL" and b406["lineas"][1]["agencia"] == ""


def test_filtro_garantia_y_aplanado():
    r = parsear_rcp170(_rcp170_sintetico())
    assert [h["habitacion"] for h in filtrar_con_garantia(r["habitaciones"])] == ["B203", "B406"]
    filas = aplanar_lineas(r["habitaciones"], r["varios"])
    assert [f["habitacion"] for f in filas] == ["B203", "B406", "B406", "623"]
    assert [f["es_primera_linea"] for f in filas] == [True, True, False, True]
    assert sum(f["monto"] for f in filas) == pytest.approx(6780.40 + 8151.06 + 3665.10 + 5135.40)


def test_parsear_rcp170_rechaza_archivo_ajeno():
    wb = Workbook()
    wb.active["A1"] = "otra cosa"
    buffer = BytesIO()
    wb.save(buffer)
    with pytest.raises(ValueError):
        parsear_rcp170(buffer.getvalue())


def test_calcular_desglose_con_valores_capturados():
    # B203 en "SALDO CUENTA HUESPED LUIS SEPTIEMBRE": los 3 montos ahora se capturan, no se derivan
    d = calcular_desglose(6780.40, subtotal=5465.30, iva=790.10, propina=525)
    assert d["gravado"] == pytest.approx(4938.125) and d["tasa0"] == pytest.approx(527.175)  # derivados
    assert d["subtotal"] == 5465.30
    assert d["total2"] == pytest.approx(6780.40) and d["dif"] == pytest.approx(0)
    # sin capturar nada: DIF vale 0 (no hay contra qué comparar)
    assert calcular_desglose(100.0)["dif"] == 0
    # captura parcial: DIF avisa lo que falta
    assert calcular_desglose(100.0, iva=16.0)["dif"] == pytest.approx(84.0)


def test_excel_captura_manual_y_formulas_solo_en_total_y_dif():
    filas = aplanar_lineas(*(lambda r: (r["habitaciones"], r["varios"]))(parsear_rcp170(_rcp170_sintetico())))
    for f in filas:
        f.update(subtotal=3.0, iva=4.0, propina=5.0, ieps=0.0, ish=0.0)
    ws = load_workbook(BytesIO(generar_excel_saldo_cuenta_huesped(filas, "2026-09"))).active
    assert ws["L4"].value == 6780.40  # TOTAL = monto directo, sin columna "Monto" aparte
    assert (ws["M4"].value, ws["N4"].value, ws["O4"].value) == ("=P4/0.16", "=O4-M4", 3.0)
    assert ws["T4"].value == "=O4+P4+Q4+R4+S4"
    assert ws["U4"].value == "=IF(SUM(M4:S4)=0,0,L4-T4)"
    assert ws.cell(row=4 + len(filas), column=12).value == f"=SUM(L4:L{3 + len(filas)})"


@pytest.mark.skipif(not RCP170_REAL.exists(), reason="carpeta de muestra 'saldo cuenta huesped/' no presente")
def test_rcp170_real_septiembre_cuadra_con_el_reporte():
    r = parsear_rcp170(RCP170_REAL.read_bytes())
    assert len(r["habitaciones"]) == 29 and len(r["varios"]) == 2
    assert r["saldo_total"] == pytest.approx(175841.81)
    # las líneas (incluidas las que caen tras un salto de página) suman el saldo de cada habitación
    for h in r["habitaciones"]:
        assert sum(l["monto"] for l in h["lineas"]) == pytest.approx(h["saldo"]), h["habitacion"]
    filas = aplanar_lineas(r["habitaciones"], r["varios"])
    assert sum(f["monto"] for f in filas) == pytest.approx(175841.81)  # las 2 sin garantía saldan en 0


@pytest.mark.anyio
async def test_router_subir_capturar_y_resubir_conserva_captura(client, db_session):
    async def _get_db():
        yield db_session

    app.dependency_overrides[get_db] = _get_db
    try:
        archivo = {"file": ("RCP170.xls", _rcp170_sintetico(), "application/vnd.ms-excel")}
        r = await client.post("/saldo-cuenta-huesped/2026-09/upload", files=archivo)
        assert r.status_code == 200, r.text
        cuerpo = r.json()
        assert len(cuerpo["filas"]) == 4 and cuerpo["excluidas_sin_garantia"] == ["A505"]

        linea = cuerpo["filas"][0]  # B203
        p = await client.patch(
            f"/saldo-cuenta-huesped/2026-09/lineas/{linea['id']}", json={"subtotal": 5465.30, "iva": 790.10, "propina": 525}
        )
        assert p.status_code == 200 and p.json()["dif"] == pytest.approx(0)

        r2 = await client.post(
            "/saldo-cuenta-huesped/2026-09/upload",
            files={"file": ("RCP170.xls", _rcp170_sintetico(), "application/vnd.ms-excel")},
        )
        assert r2.json()["capturas_conservadas"] == 4
        primera = r2.json()["filas"][0]
        assert (primera["iva"], primera["propina"], primera["subtotal"], primera["gravado"]) == (790.10, 525, 5465.30, 4938.125)

        exp = await client.get("/saldo-cuenta-huesped/2026-09/export")
        assert exp.status_code == 200 and exp.content[:2] == b"PK"

        assert (await client.patch("/saldo-cuenta-huesped/2026-09/lineas/99999", json={"iva": 1})).status_code == 404
        assert (await client.get("/saldo-cuenta-huesped/2026-13")).status_code == 400
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_lineas_en_cero_se_omiten_y_la_primera_que_queda_abre_el_grupo():
    habitaciones = [
        {"habitacion": "X1", "codigo": "c", "huesped": "H", "adultos_ninos": "", "llegada": "", "salida": "",
         "tarifa": 0.0, "garantia": "CORPORATIVO", "saldo": 50.0, "tipo": "habitacion",
         "lineas": [{"agencia": "A", "linea": "Principal", "monto": 0.0}, {"agencia": "", "linea": "SUB*1*", "monto": 50.0}]},
        {"habitacion": "X2", "codigo": "c", "huesped": "H", "adultos_ninos": "", "llegada": "", "salida": "",
         "tarifa": 0.0, "garantia": "CORPORATIVO", "saldo": 0.0, "tipo": "habitacion",
         "lineas": [{"agencia": "A", "linea": "Principal", "monto": 0.0}]},
    ]
    filas = aplanar_lineas(habitaciones, [])
    assert [(f["habitacion"], f["linea_nombre"], f["es_primera_linea"]) for f in filas] == [("X1", "SUB*1*", True)]


def test_varios_en_cero_se_conserva():
    varios = [{"habitacion": "700", "codigo": "", "huesped": "V", "adultos_ninos": "", "llegada": "", "salida": "",
               "tarifa": 0.0, "garantia": "", "saldo": 0.0, "tipo": "varios", "lineas": []}]
    filas = aplanar_lineas([], varios)
    assert [(f["habitacion"], f["monto"]) for f in filas] == [("700", 0.0)]


@pytest.mark.anyio
async def test_router_oculta_lineas_en_cero_ya_guardadas_pero_no_varios(client, db_session):
    from app.models import SaldoCuentaHuespedLinea

    def fila(orden, hab, tipo, monto, primera):
        return SaldoCuentaHuespedLinea(
            periodo="2026-10", orden=orden, tipo=tipo, habitacion=hab, codigo="", huesped="H",
            garantia="CORPORATIVO" if tipo == "habitacion" else "", saldo_habitacion=monto,
            es_primera_linea=primera, linea_nombre="L", monto=monto,
        )

    db_session.add_all([
        fila(1, "B304", "habitacion", 0.0, True),      # solo línea en 0 -> se oculta
        fila(2, "B406", "habitacion", 0.0, True),      # 1ª en 0, la 2ª abre el grupo
        fila(3, "B406", "habitacion", 50.0, False),
        fila(4, "623", "varios", 0.0, True),           # Varios en 0 se conserva
    ])
    await db_session.flush()

    async def _get_db():
        yield db_session

    app.dependency_overrides[get_db] = _get_db
    try:
        filas = (await client.get("/saldo-cuenta-huesped/2026-10")).json()["filas"]
        assert [(f["habitacion"], f["es_primera_linea"]) for f in filas] == [("B406", True), ("623", True)]
        exp = await client.get("/saldo-cuenta-huesped/2026-10/export")
        ws = load_workbook(BytesIO(exp.content)).active
        assert ws["A4"].value == "B406" and ws["A5"].value == "623" and ws["A6"].value is None
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_excel_formato_de_trabajo():
    filas = aplanar_lineas(*(lambda r: (r["habitaciones"], r["varios"]))(parsear_rcp170(_rcp170_sintetico())))
    for f in filas:
        f.update(gravado=0.0, tasa0=0.0, subtotal=0.0, iva=0.0, propina=0.0, ieps=0.0, ish=0.0)
    ws = load_workbook(BytesIO(generar_excel_saldo_cuenta_huesped(filas, "2026-09"))).active
    assert ws.freeze_panes is None  # sin paneles inmovilizados (pedido explícito)
    assert ws.auto_filter.ref == f"A3:U{3 + len(filas)}"
    assert "A1:U1" in [str(r) for r in ws.merged_cells.ranges]
    assert "Saldo en Cuenta Huésped" in ws["A1"].value and "2026-09" in ws["A1"].value
    validaciones = ws.data_validations.dataValidation
    assert len(validaciones) == 1 and "O4:S" in str(validaciones[0].sqref)
    assert ws.page_setup.orientation == "landscape"


def _excel_de_prueba(filas: list[dict], periodo: str = "2026-09") -> bytes:
    return generar_excel_saldo_cuenta_huesped(filas, periodo)


def test_excel_editado_round_trip():
    filas = aplanar_lineas(*(lambda r: (r["habitaciones"], r["varios"]))(parsear_rcp170(_rcp170_sintetico())))
    for f in filas:
        f.update(subtotal=0.0, iva=0.0, propina=0.0, ieps=0.0, ish=0.0)
    filas[0].update(subtotal=5465.30, iva=790.10, propina=525.0)
    leido = leer_excel_editado(_excel_de_prueba(filas, "2026-09"))
    assert leido["periodo_archivo"] == "2026-09"
    assert [f["habitacion"] for f in leido["filas"]] == [f["habitacion"] for f in filas]
    primera = leido["filas"][0]
    assert (primera["monto"], primera["subtotal"], primera["iva"], primera["propina"]) == (6780.40, 5465.30, 790.10, 525.0)
    assert es_excel_editado(_excel_de_prueba(filas)) and not es_excel_editado(_rcp170_sintetico())


@pytest.mark.skipif(not (Path(__file__).resolve().parents[2] / "docs" / "SALDO CUENTA HUESPED SEPTIEMBRE 2026..xlsx").exists(),
                    reason="archivo de muestra editado no presente")
def test_excel_editado_real_del_usuario():
    ruta = Path(__file__).resolve().parents[2] / "docs" / "SALDO CUENTA HUESPED SEPTIEMBRE 2026..xlsx"
    contenido = ruta.read_bytes()
    assert es_excel_editado(contenido)
    leido = leer_excel_editado(contenido)
    assert leido["periodo_archivo"] == "2026-10"
    assert len(leido["filas"]) == 52
    a201 = leido["filas"][0]
    assert (a201["habitacion"], a201["monto"], a201["iva"], a201["propina"]) == ("A201", 18313.4, 2140.86, 2375.0)


@pytest.mark.anyio
async def test_router_excel_editado_pide_confirmar_periodo(client, db_session):
    async def _get_db():
        yield db_session

    app.dependency_overrides[get_db] = _get_db
    try:
        filas = aplanar_lineas(*(lambda r: (r["habitaciones"], r["varios"]))(parsear_rcp170(_rcp170_sintetico())))
        for f in filas:
            f.update(subtotal=0.0, iva=0.0, propina=0.0, ieps=0.0, ish=0.0)
        excel = _excel_de_prueba(filas, "2026-10")
        url = "/saldo-cuenta-huesped/2026-09/upload"
        aviso = await client.post(url, files={"file": ("x.xlsx", excel, "application/vnd.ms-excel")})
        assert aviso.json() == {"requiere_confirmacion": True, "periodo_archivo": "2026-10", "periodo_seleccionado": "2026-09"}
        ok = await client.post(url + "?confirmar_periodo=true", files={"file": ("x.xlsx", excel, "application/vnd.ms-excel")})
        assert ok.status_code == 200 and ok.json()["origen"] == "excel_editado" and len(ok.json()["filas"]) == len(filas)
    finally:
        app.dependency_overrides.pop(get_db, None)
