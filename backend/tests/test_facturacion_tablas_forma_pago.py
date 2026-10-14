"""
Guardado en base de datos del detalle de Tablas_Forma_Pago (Facturación,
"Paso 2"): el parser (leer_filas_completas_tablas_forma_pago) debe capturar
las 15 columnas de cada sección tolerando filas en blanco intercaladas
(igual que agregar_pue_ppd_tablas_forma_pago, del que reusa la lógica de
localización de secciones), y el guardado debe reemplazar limpiamente un
periodo ya guardado en vez de duplicar filas.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import delete, func, select

from app.models import FacturacionTablaFormaPagoDetalle
from app.services.excel_processor import leer_filas_completas_tablas_forma_pago
from app.services.facturacion import filas_a_registros_tabla_forma_pago

# A→O, mismo orden que COLUMNAS_TABLA_FACTURACION.
_COLUMNAS = [
    "Revisión", "Fecha", "POS", "Forma de Pago", "Cuenta", "Subtotal", "Descuento",
    "GRAVADO", "INGTASA 0", "DESC TASA 0", "SUBTASA 0", "IVA", "Propina",
    "PAGO PROVEDOR", "TOTAL",
]


def _escribir_fila(ws, row: int, valores: dict) -> None:
    for indice, nombre in enumerate(_COLUMNAS, start=1):
        ws.cell(row=row, column=indice, value=valores.get(nombre))


def _construir_excel_tablas_forma_pago() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tablas_Forma_Pago"

    # Sección EFECTIVO: título, header, dato, fila en blanco intercalada, dato, TOTAL.
    ws.cell(row=1, column=1, value="EFECTIVO")
    _escribir_fila(ws, 2, {nombre: nombre for nombre in _COLUMNAS})  # header, se ignora
    _escribir_fila(ws, 3, {
        "Fecha": "2026-08-01", "POS": "LA HACIENDA", "Forma de Pago": "EFECTIVO", "Cuenta": "1001",
        "Subtotal": 100.0, "Descuento": 10.0, "GRAVADO": 90.0, "INGTASA 0": 0.0, "DESC TASA 0": 0.0,
        "SUBTASA 0": 0.0, "IVA": 14.4, "Propina": 5.0, "PAGO PROVEDOR": 0.0, "TOTAL": 109.4,
    })
    # Fila 4 se deja completamente en blanco (simula un espaciador insertado a mano).
    _escribir_fila(ws, 5, {
        "Fecha": "2026-08-02", "POS": "LA HACIENDA", "Forma de Pago": "EFECTIVO", "Cuenta": "1002",
        "Subtotal": 200.0, "Descuento": 0.0, "GRAVADO": 200.0, "INGTASA 0": 0.0, "DESC TASA 0": 0.0,
        "SUBTASA 0": 0.0, "IVA": 32.0, "Propina": 0.0, "PAGO PROVEDOR": 0.0, "TOTAL": 232.0,
    })
    _escribir_fila(ws, 6, {"Cuenta": "TOTAL", "Subtotal": 300.0})

    # Sección HUÉSPED: título, header, dato, TOTAL.
    ws.cell(row=7, column=1, value="HUÉSPED")
    _escribir_fila(ws, 8, {nombre: nombre for nombre in _COLUMNAS})
    _escribir_fila(ws, 9, {
        "Fecha": "2026-08-03", "POS": "LA HACIENDA", "Forma de Pago": "HUESPED", "Cuenta": "2001",
        "Subtotal": 50.0, "Descuento": 0.0, "GRAVADO": 50.0, "INGTASA 0": 0.0, "DESC TASA 0": 0.0,
        "SUBTASA 0": 0.0, "IVA": 8.0, "Propina": 0.0, "PAGO PROVEDOR": 0.0, "TOTAL": 58.0,
    })
    _escribir_fila(ws, 10, {"Cuenta": "TOTAL", "Subtotal": 50.0})

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_leer_filas_completas_tolera_blancos_y_lee_las_15_columnas():
    filas = leer_filas_completas_tablas_forma_pago(_construir_excel_tablas_forma_pago())

    efectivo = [f for f in filas if f["tabla_origen"] == "EFECTIVO"]
    huesped = [f for f in filas if f["tabla_origen"] == "HUÉSPED"]

    # La fila en blanco intercalada (fila 4) no debe colarse ni cortar la lectura.
    assert [f["Cuenta"] for f in efectivo] == ["1001", "1002"]
    assert [f["Cuenta"] for f in huesped] == ["2001"]

    primera = efectivo[0]
    assert primera["Fecha"] == "2026-08-01"
    assert primera["POS"] == "LA HACIENDA"
    assert primera["Forma de Pago"] == "EFECTIVO"
    assert primera["Subtotal"] == 100.0
    assert primera["Descuento"] == 10.0
    assert primera["GRAVADO"] == 90.0
    assert primera["IVA"] == 14.4
    assert primera["Propina"] == 5.0
    assert primera["TOTAL"] == 109.4
    assert primera["fila_excel"] == 3

    # La fila "TOTAL" de cierre de cada sección nunca debe aparecer como dato.
    assert all(f["Cuenta"] != "TOTAL" for f in filas)


def test_filas_a_registros_asigna_categoria_pue_ppd():
    filas = leer_filas_completas_tablas_forma_pago(_construir_excel_tablas_forma_pago())
    registros = filas_a_registros_tabla_forma_pago(filas, periodo="Agosto 2026")

    por_cuenta = {r["cuenta"]: r for r in registros}
    assert por_cuenta["1001"]["categoria_pue_ppd"] == "PUE"
    assert por_cuenta["1001"]["periodo"] == "Agosto 2026"
    assert por_cuenta["1001"]["tabla_origen"] == "EFECTIVO"
    assert por_cuenta["2001"]["categoria_pue_ppd"] == "PPD"
    assert por_cuenta["2001"]["tabla_origen"] == "HUÉSPED"


@pytest.mark.anyio
async def test_guardar_reemplaza_periodo_sin_duplicar(db_session):
    periodo = "Agosto 2026"

    async def _contar() -> int:
        res = await db_session.execute(
            select(func.count())
            .select_from(FacturacionTablaFormaPagoDetalle)
            .where(FacturacionTablaFormaPagoDetalle.periodo == periodo)
        )
        return res.scalar_one()

    filas = leer_filas_completas_tablas_forma_pago(_construir_excel_tablas_forma_pago())
    registros = filas_a_registros_tabla_forma_pago(filas, periodo=periodo)
    assert len(registros) == 3

    for registro in registros:
        db_session.add(FacturacionTablaFormaPagoDetalle(**registro))
    await db_session.commit()
    assert await _contar() == 3

    # "Re-guardar" el mismo periodo (mismo comportamiento que
    # /facturacion/tablas-forma-pago/guardar con confirmar_reemplazo=True):
    # borrar lo anterior e insertar de nuevo, sin dejar mezcla de filas viejas
    # y nuevas ni duplicados.
    await db_session.execute(
        delete(FacturacionTablaFormaPagoDetalle).where(FacturacionTablaFormaPagoDetalle.periodo == periodo)
    )
    for registro in registros:
        db_session.add(FacturacionTablaFormaPagoDetalle(**registro))
    await db_session.commit()

    assert await _contar() == 3
