"""
Flujo completo de la Fiesta contra la API real (BD en memoria): reservar, doble venta, cobro con 2 formas de pago,
corregir/cancelar pagos, cancelar reservación, totales de Cobros y auditoría de cada paso. QA previo al lanzamiento.
"""

from datetime import date

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.db.base import Base
from app.db.session import get_db
from app.main import app

pytestmark = pytest.mark.anyio

PREFIJO = "/fiesta-fin-anio"
HOY = date.today().isoformat()
CONTACTO = {"telefono_o_cuarto": "669 123 4567", "correo": "cliente@correo.com"}


@pytest.fixture
async def api(monkeypatch):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool, connect_args={"check_same_thread": False})
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    fabrica = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _get_db():
        async with fabrica() as s:
            try:
                yield s
                await s.commit()
            except Exception:
                await s.rollback()
                raise

    monkeypatch.setattr(settings, "auth_enabled", False)
    app.dependency_overrides[get_db] = _get_db
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        assert (await c.post(f"{PREFIJO}/sembrar")).status_code in (200, 201)
        precios = [
            {"tipo": "adulto", "precio_preventa": 2000, "precio_normal": 2450},
            {"tipo": "nino", "precio_preventa": 1000, "precio_normal": 1200},
            {"tipo": "cortesia", "precio_preventa": 0, "precio_normal": 0},
        ]
        assert (await c.put(f"{PREFIJO}/precios", json=precios)).status_code == 200
        # corte de preventa en el futuro para que el precio sea el de preventa
        evento = (await c.get(f"{PREFIJO}/evento")).json()
        await c.put(f"{PREFIJO}/evento", json={"corte_preventa": date(date.today().year + 1, 1, 1).isoformat()})
        yield c, evento
    app.dependency_overrides.pop(get_db, None)
    await engine.dispose()


async def _sillas_libres(c, n):
    mesas = (await c.get(f"{PREFIJO}/mesas")).json()["mesas"]
    libres = [s["id"] for m in mesas for s in m["sillas"] if s["estado"] == "libre" and not s.get("extra")]
    return libres[:n]


async def _reservar(c, cliente, tipos, **extra):
    ids = await _sillas_libres(c, len(tipos))
    r = await c.post(
        f"{PREFIJO}/reservaciones",
        json={"cliente": cliente, **CONTACTO, "sillas": [{"silla_id": i, "tipo": t} for i, t in zip(ids, tipos)], "fecha_venta": HOY, **extra},
    )
    return r, ids


async def test_flujo_completo_con_auditoria(api):
    c, _ = api

    # 1) reservar 2 adultos + 1 niño: precio de preventa congelado por silla
    r, ids = await _reservar(c, "Luis Jose", ["adulto", "adulto", "nino"])
    assert r.status_code == 201, r.text
    res = r.json()
    rid = res["id"]
    assert res["importe"] == pytest.approx(2000 + 2000 + 1000)

    # 2) doble venta de la misma silla → 409 con mensaje claro
    r2 = await c.post(f"{PREFIJO}/reservaciones", json={"cliente": "Otra", **CONTACTO, "sillas": [{"silla_id": ids[0], "tipo": "adulto"}], "fecha_venta": HOY})
    assert r2.status_code == 409 and "silla" in r2.json()["detail"].lower()

    # 3) cobro con 2 formas de pago aplicado a las 3 sillas → UN solo "Pago 1" por silla
    r3 = await c.post(
        f"{PREFIJO}/reservaciones/{rid}/pagos/multiple",
        json={"pagos": [{"importe": 837.5, "forma_pago": "Efectivo", "fecha": HOY}, {"importe": 1000, "forma_pago": "Tarjeta", "fecha": HOY}]},
    )
    assert r3.status_code == 201, r3.text
    det = r3.json()
    assert det["pagado"] == pytest.approx(1837.5)
    for silla in det["sillas"]:
        assert len(silla["pagos"]) == 1, "las dos formas de un cobro deben ser UN solo Pago N"
        assert {p["forma_pago"] for p in silla["pagos"][0]["partes"]} == {"Efectivo", "Tarjeta"}
    assert sum(s["pagado"] for s in det["sillas"]) == pytest.approx(1837.5)

    # 4) Cobros: el total cuadra, por forma de pago separa efectivo y tarjeta, y el día muestra el cobro completo en Pago 1
    cobros = (await c.get(f"{PREFIJO}/cobros")).json()
    assert cobros["total"] == pytest.approx(1837.5)
    por_forma = {f["forma_pago"]: f["importe"] for f in cobros["por_forma_pago"]}
    assert por_forma.get("Efectivo") == pytest.approx(837.5) and por_forma.get("Tarjeta") == pytest.approx(1000)
    dia = next(d for d in cobros["dias"] if d["fecha"] == HOY)
    assert dia["pago1"] == pytest.approx(1837.5) and dia["pago2"] == 0
    assert dia["vendidas"] == {"adulto": 2, "nino": 1, "cortesia": 0}

    # 5) corregir un pago y cancelarlo
    pago_ef = next(p for p in det["pagos"] if p["forma_pago"] == "Efectivo")
    r5 = await c.patch(f"{PREFIJO}/pagos/{pago_ef['id']}", json={"importe": 800})
    assert r5.status_code == 200 and r5.json()["pagado"] == pytest.approx(1800)
    r6 = await c.delete(f"{PREFIJO}/pagos/{pago_ef['id']}")
    assert r6.status_code == 200 and r6.json()["pagado"] == pytest.approx(1000)

    # 6) cancelar la reservación libera las sillas
    r7 = await c.delete(f"{PREFIJO}/reservaciones/{rid}", params={"motivo": "prueba QA"})
    assert r7.status_code == 200
    mesas = (await c.get(f"{PREFIJO}/mesas")).json()["mesas"]
    estado = {s["id"]: s["estado"] for m in mesas for s in m["sillas"]}
    assert all(estado[i] == "libre" for i in ids)
    assert (await c.delete(f"{PREFIJO}/reservaciones/{rid}")).status_code == 409  # ya cancelada

    # 7) la auditoría registró cada paso con su reservación y cliente
    aud = (await c.get(f"{PREFIJO}/auditoria", params={"por_pagina": 200})).json()
    acciones = [f["accion"] for f in aud["filas"] if f["reservacion_id"] == rid]
    for esperado in ("crear", "pago", "editar_pago", "cancelar_pago", "cancelar"):
        assert esperado in acciones, f"falta {esperado} en {acciones}"
    assert all(f["cliente"] == "Luis Jose" for f in aud["filas"] if f["reservacion_id"] == rid)
    assert aud["resumen"]["por_grupo"].get("cobros", 0) >= 3
    ed = next(f for f in aud["filas"] if f["accion"] == "editar_pago")
    assert ed["antes"]["importe"] == pytest.approx(837.5) and ed["despues"]["importe"] == pytest.approx(800)
    # filtros
    solo_cancel = (await c.get(f"{PREFIJO}/auditoria", params={"grupo": "cancelaciones"})).json()
    assert solo_cancel["total"] >= 2 and all(f["grupo"] == "cancelaciones" for f in solo_cancel["filas"])
    por_res = (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": rid})).json()
    assert por_res["total"] == len(acciones) or por_res["total"] >= 5


async def test_validaciones_no_dan_500(api):
    c, _ = api
    r, ids = await _reservar(c, "Validar", ["adulto"])
    rid = r.json()["id"]
    # importes inválidos
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 0})).status_code == 422
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": -5})).status_code == 422
    # reservación, pago y silla inexistentes
    assert (await c.post(f"{PREFIJO}/reservaciones/99999/pagos", json={"importe": 10})).status_code == 404
    assert (await c.patch(f"{PREFIJO}/pagos/99999", json={"importe": 10})).status_code == 404
    assert (await c.delete(f"{PREFIJO}/pagos/99999")).status_code == 404
    bad = await c.post(f"{PREFIJO}/reservaciones", json={"cliente": "X", **CONTACTO, "sillas": [{"silla_id": 999999, "tipo": "adulto"}], "fecha_venta": HOY})
    assert bad.status_code in (400, 404, 409, 422)
    # reservación sin sillas / tipo inválido / correo inválido
    assert (await c.post(f"{PREFIJO}/reservaciones", json={"cliente": "X", **CONTACTO, "sillas": []})).status_code == 422
    assert (await c.post(f"{PREFIJO}/reservaciones", json={"cliente": "X", **CONTACTO, "sillas": [{"silla_id": ids[0], "tipo": "vip"}]})).status_code == 422
    assert (await c.post(f"{PREFIJO}/reservaciones", json={"cliente": "X", "telefono_o_cuarto": "6691234567", "correo": "no-es-correo", "sillas": [{"silla_id": 1, "tipo": "adulto"}]})).status_code == 422
    # abonar a una reservación cancelada
    await c.delete(f"{PREFIJO}/reservaciones/{rid}")
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 10})).status_code == 409
    # filtros de auditoría inválidos
    assert (await c.get(f"{PREFIJO}/auditoria", params={"grupo": "inventado"})).status_code == 400
    assert (await c.get(f"{PREFIJO}/auditoria", params={"pagina": 0})).status_code == 422


async def test_humo_de_todas_las_lecturas_de_la_fiesta(api):
    c, _ = api
    r, _ = await _reservar(c, "Humo", ["adulto", "cortesia"], cortesia_motivo="Invitado del hotel")
    rid = r.json()["id"]
    await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 500, "forma_pago": "Efectivo"})
    rutas = [
        "/evento", "/mesas", "/reservaciones", f"/reservaciones/{rid}", "/reservaciones?incluir_canceladas=true",
        "/resumen", "/cobros", "/historico", "/bitacora", "/auditoria", "/auditoria/export",
    ]
    fallos = []
    for ruta in rutas:
        resp = await c.get(PREFIJO + ruta)
        if resp.status_code >= 500 or resp.status_code in (401, 403):
            fallos.append(f"{ruta} → {resp.status_code}")
    assert not fallos, fallos


async def test_sobrepago_no_rompe_los_totales(api):
    c, _ = api
    r, ids = await _reservar(c, "Sobrepago", ["adulto"])  # 2000
    rid = r.json()["id"]
    p = await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 2500, "forma_pago": "Efectivo", "fecha": HOY})
    assert p.status_code == 201
    det = p.json()
    assert det["pagado"] == pytest.approx(2500)
    assert det["sillas"][0]["saldo"] == 0 and det["sillas"][0]["estado_pago"] == "pagada"
    assert det["saldo"] >= 0  # el saldo nunca sale negativo
    cobros = (await c.get(f"{PREFIJO}/cobros")).json()
    assert cobros["total"] == pytest.approx(2500)  # el dinero cobrado es real aunque sobre


async def test_varias_cajas_reservando_la_misma_silla_solo_una_gana(tmp_path, monkeypatch):
    """Con una BD en archivo y conexiones separadas (como en producción): una caja gana, las demás reciben 409, sin 500."""
    import asyncio

    from app.db.session import configurar_sqlite

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'concurrencia.sqlite3'}")
    configurar_sqlite(engine)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    fabrica = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async def _get_db():
        async with fabrica() as s:
            try:
                yield s
                await s.commit()
            except Exception:
                await s.rollback()
                raise

    monkeypatch.setattr(settings, "auth_enabled", False)
    app.dependency_overrides[get_db] = _get_db
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            assert (await c.post(f"{PREFIJO}/sembrar")).status_code in (200, 201)
            await c.put(f"{PREFIJO}/precios", json=[
                {"tipo": "adulto", "precio_preventa": 2000, "precio_normal": 2450},
                {"tipo": "nino", "precio_preventa": 1000, "precio_normal": 1200},
                {"tipo": "cortesia", "precio_preventa": 0, "precio_normal": 0},
            ])
            ids = await _sillas_libres(c, 1)
            cuerpo = lambda n: {"cliente": n, **CONTACTO, "sillas": [{"silla_id": ids[0], "tipo": "adulto"}], "fecha_venta": HOY}
            respuestas = await asyncio.gather(*[c.post(f"{PREFIJO}/reservaciones", json=cuerpo(f"Caja {i}")) for i in range(6)])
            codigos = sorted(r.status_code for r in respuestas)
            assert codigos == [201, 409, 409, 409, 409, 409], codigos
            activas = [x for x in (await c.get(f"{PREFIJO}/reservaciones")).json()["reservaciones"] if x["estatus"] == "activa"]
            assert len(activas) == 1
    finally:
        app.dependency_overrides.pop(get_db, None)
        await engine.dispose()


async def test_precio_cambia_de_preventa_a_normal_solo_en_ventas_futuras(api):
    c, _ = api
    r1, _ = await _reservar(c, "Preventa", ["adulto"])  # corte en el futuro → preventa 2000
    assert r1.json()["importe"] == pytest.approx(2000)
    # el corte pasa: las ventas nuevas cuestan el precio normal; la ya vendida conserva el suyo
    await c.put(f"{PREFIJO}/evento", json={"corte_preventa": "2020-01-01"})
    r2, _ = await _reservar(c, "Normal", ["adulto"])
    assert r2.json()["importe"] == pytest.approx(2450)
    det1 = (await c.get(f"{PREFIJO}/reservaciones/{r1.json()['id']}")).json()
    assert det1["importe"] == pytest.approx(2000)


async def test_cortesia_cuesta_cero_y_no_pide_pago(api):
    c, _ = api
    r, _ = await _reservar(c, "Cortesia", ["cortesia", "cortesia"], cortesia_motivo="Patrocinador del evento")
    assert r.status_code == 201
    assert r.json()["cortesia"]["motivo"] == "Patrocinador del evento"
    assert r.json()["importe"] == 0
    assert r.json()["estado"] == "cortesia"


async def test_reservar_exige_nombre_telefono_y_correo(api):
    c, _ = api
    ids = await _sillas_libres(c, 1)
    base = {"sillas": [{"silla_id": ids[0], "tipo": "adulto"}], "fecha_venta": HOY}
    completo = {"cliente": "Maria Lopez", "telefono_o_cuarto": "669 123 4567", "correo": "maria@correo.com"}
    # cada dato faltante, vacío o inválido → 422 y NO se aparta la silla
    invalidos = [
        {**completo, "cliente": ""}, {**completo, "cliente": "   "}, {**completo, "cliente": "A"},
        {k: v for k, v in completo.items() if k != "telefono_o_cuarto"}, {**completo, "telefono_o_cuarto": ""},
        {**completo, "telefono_o_cuarto": "12345"}, {**completo, "telefono_o_cuarto": "abc"},
        {k: v for k, v in completo.items() if k != "correo"}, {**completo, "correo": ""}, {**completo, "correo": "no-es-correo"},
    ]
    for datos in invalidos:
        r = await c.post(f"{PREFIJO}/reservaciones", json={**base, **datos})
        assert r.status_code == 422, (datos, r.status_code, r.text[:120])
    assert (await _sillas_libres(c, 1)) == ids  # la silla sigue libre
    ok = await c.post(f"{PREFIJO}/reservaciones", json={**base, **completo})
    assert ok.status_code == 201
    det = ok.json()
    assert det["cliente"] == "Maria Lopez" and det["telefono_o_cuarto"] == "669 123 4567" and det["correo"] == "maria@correo.com"
    # al editar no se pueden dejar en blanco
    rid = det["id"]
    assert (await c.patch(f"{PREFIJO}/reservaciones/{rid}", json={"correo": ""})).status_code == 422
    assert (await c.patch(f"{PREFIJO}/reservaciones/{rid}", json={"telefono_o_cuarto": "  "})).status_code == 422
    assert (await c.patch(f"{PREFIJO}/reservaciones/{rid}", json={"notas": "solo notas"})).status_code == 200


async def test_resumen_cuenta_pagadas_y_anticipos_por_silla(api):
    """Una silla pagada completa dentro de una reservación con otras sillas sin pagar debe contar como PAGADA en el Resumen."""
    c, _ = api
    r, ids = await _reservar(c, "Resumen por silla", ["adulto", "adulto", "adulto"])  # 2000 c/u
    rid = r.json()["id"]
    p1 = await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 2000, "forma_pago": "Efectivo", "fecha": HOY, "silla_ids": [ids[0]]})
    assert p1.status_code == 201
    p2 = await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 500, "forma_pago": "Efectivo", "fecha": HOY, "silla_ids": [ids[1]]})
    assert p2.status_code == 201
    estado = (await c.get(f"{PREFIJO}/resumen")).json()["estado"]
    assert (estado["pagadas"], estado["anticipos"], estado["sin_anticipo"]) == (1, 1, 1)
    # coincide con "cómo vamos al día" de Cobros (que ya contaba por silla)
    dia = next(d for d in (await c.get(f"{PREFIJO}/cobros")).json()["dias"] if d["fecha"] == HOY)
    assert (dia["reservas"]["pagadas"], dia["reservas"]["anticipos"], dia["reservas"]["sin_anticipo"]) == (1, 1, 1)


async def test_cortesia_sin_motivo_se_rechaza_y_no_aparta_la_silla(api):
    c, _ = api
    ids = await _sillas_libres(c, 1)
    sin_motivo = await _reservar(c, "Regalo", ["cortesia"])
    assert sin_motivo[0].status_code == 422 and "motivo" in sin_motivo[0].json()["detail"].lower()
    corto = await _reservar(c, "Regalo", ["cortesia"], cortesia_motivo="ab")
    assert corto[0].status_code == 422
    assert (await _sillas_libres(c, 1)) == ids  # nada quedó apartado
    ok = await _reservar(c, "Regalo", ["adulto", "cortesia"], cortesia_motivo="Cortesía de gerencia")
    assert ok[0].status_code == 201
    # agregar una cortesía después también exige motivo
    rid = ok[0].json()["id"]
    libre = (await _sillas_libres(c, 1))[0]
    r1 = await c.patch(f"{PREFIJO}/reservaciones/{rid}", json={"sillas_agregar": [{"silla_id": libre, "tipo": "cortesia"}]})
    assert r1.status_code == 422
    r2 = await c.patch(f"{PREFIJO}/reservaciones/{rid}", json={"sillas_agregar": [{"silla_id": libre, "tipo": "cortesia"}], "cortesia_motivo": "Otra cortesía"})
    assert r2.status_code == 200 and "Otra cortesía" in r2.json()["cortesia"]["motivo"]
    # y queda en la auditoría con su motivo
    aud = (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": rid})).json()
    assert any((f["despues"] or {}).get("cortesia_motivo") for f in aud["filas"])


async def test_cortesia_de_cajera_queda_por_autorizar_y_el_supervisor_la_resuelve(api, monkeypatch):
    from app.api import fiesta_fin_anio as modulo

    c, _ = api
    monkeypatch.setattr(modulo, "_es_supervisor", lambda: False)  # quien captura es una cajera
    r, _ = await _reservar(c, "Invitado", ["adulto", "cortesia"], cortesia_motivo="Invitado del gerente")
    assert r.status_code == 201
    rid = r.json()["id"]
    assert r.json()["cortesias_por_autorizar"] == 1 and r.json()["cortesia"]["autorizada_por"] == ""
    assert sum(1 for s in r.json()["sillas"] if s["por_autorizar"]) == 1

    pendientes = (await c.get(f"{PREFIJO}/solicitudes-cortesia")).json()["solicitudes"]
    assert len(pendientes) == 1 and pendientes[0]["motivo"] == "Invitado del gerente" and pendientes[0]["cliente"] == "Invitado"

    ok = await c.post(f"{PREFIJO}/solicitudes-cortesia/{pendientes[0]['id']}/aprobar")
    assert ok.status_code == 200 and ok.json()["cortesias_por_autorizar"] == 0
    assert (await c.get(f"{PREFIJO}/solicitudes-cortesia")).json()["total"] == 0
    # ya resuelta: no se puede resolver otra vez (otro supervisor se adelantó)
    assert (await c.post(f"{PREFIJO}/solicitudes-cortesia/{pendientes[0]['id']}/rechazar")).status_code == 409

    # rechazar libera las sillas y, si era todo lo que tenía la reservación, la cancela
    libres_antes = await _sillas_libres(c, 1)
    r2, _ = await _reservar(c, "Solo cortesia", ["cortesia"], cortesia_motivo="Por compromiso")
    sol = (await c.get(f"{PREFIJO}/solicitudes-cortesia")).json()["solicitudes"][0]
    no = await c.post(f"{PREFIJO}/solicitudes-cortesia/{sol['id']}/rechazar", json={"nota": "No aplica"})
    assert no.status_code == 200 and no.json()["estatus"] == "cancelada"
    assert await _sillas_libres(c, 1) == libres_antes
    aud = (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": rid})).json()
    assert {f["accion"] for f in aud["filas"]} >= {"cortesia_solicitada", "cortesia_aprobada"}


async def test_solicitud_de_cortesia_desaparece_si_la_cajera_cancela_la_reservacion(api, monkeypatch):
    from app.api import fiesta_fin_anio as modulo

    c, _ = api
    monkeypatch.setattr(modulo, "_es_supervisor", lambda: False)
    r, _ = await _reservar(c, "Se arrepiente", ["cortesia"], cortesia_motivo="Prueba")
    assert (await c.get(f"{PREFIJO}/solicitudes-cortesia")).json()["total"] == 1
    assert (await c.delete(f"{PREFIJO}/reservaciones/{r.json()['id']}", params={"motivo": "ya no"})).status_code == 200
    assert (await c.get(f"{PREFIJO}/solicitudes-cortesia")).json()["total"] == 0


async def test_devoluciones_solo_de_lo_que_sobra_y_restan_de_cobros(api):
    c, _ = api
    r, _ = await _reservar(c, "Se arrepintio", ["adulto", "adulto"])  # 2 x 2000
    rid = r.json()["id"]
    await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 1500, "forma_pago": "Efectivo", "fecha": HOY})
    # activa y sin excedente: no hay nada que devolver
    r0 = await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 100, "motivo": "prueba"})
    assert r0.status_code == 409
    # cancelar la reservación deja los 1500 por devolver
    assert (await c.delete(f"{PREFIJO}/reservaciones/{rid}", params={"motivo": "cliente desistió"})).status_code == 200
    asi = (await c.get(f"{PREFIJO}/reservaciones/{rid}")).json()
    assert asi["por_devolver"] == pytest.approx(1500) and asi["devuelto"] == 0
    # validaciones
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 0, "motivo": "x1234"})).status_code == 422
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 100})).status_code == 422  # sin motivo
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 1600, "motivo": "de más"})).status_code == 409
    # devolución parcial y luego el resto
    d1 = await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 1000, "forma_pago": "Efectivo", "fecha": HOY, "motivo": "reembolso parcial"})
    assert d1.status_code == 201
    assert d1.json()["devuelto"] == pytest.approx(1000) and d1.json()["por_devolver"] == pytest.approx(500)
    cobros = (await c.get(f"{PREFIJO}/cobros")).json()
    assert cobros["total"] == pytest.approx(1500) and cobros["devuelto"] == pytest.approx(1000) and cobros["neto"] == pytest.approx(500)
    dia = next(d for d in cobros["dias"] if d["fecha"] == HOY)
    assert dia["devuelto"] == pytest.approx(1000) and dia["neto"] == pytest.approx(500)
    resumen = (await c.get(f"{PREFIJO}/resumen")).json()["dinero"]
    assert resumen["devuelto"] == pytest.approx(1000) and resumen["neto"] == pytest.approx(500)
    # anular una devolución la regresa a "por devolver"
    did = d1.json()["devoluciones"][0]["id"]
    anul = await c.delete(f"{PREFIJO}/devoluciones/{did}")
    assert anul.status_code == 200 and anul.json()["por_devolver"] == pytest.approx(1500)
    assert (await c.delete(f"{PREFIJO}/devoluciones/{did}")).status_code == 409
    assert (await c.delete(f"{PREFIJO}/devoluciones/99999")).status_code == 404
    # auditoría
    acciones = [f["accion"] for f in (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": rid})).json()["filas"]]
    assert "devolucion" in acciones and "cancelar_devolucion" in acciones


async def test_devolucion_del_excedente_de_una_reservacion_activa(api):
    c, _ = api
    r, _ = await _reservar(c, "Sobrepago", ["adulto"])  # 2000
    rid = r.json()["id"]
    await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 2300, "forma_pago": "Efectivo", "fecha": HOY})
    det = (await c.get(f"{PREFIJO}/reservaciones/{rid}")).json()
    assert det["excedente"] == pytest.approx(300) and det["por_devolver"] == pytest.approx(300)
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 301, "motivo": "de más"})).status_code == 409
    ok = await c.post(f"{PREFIJO}/reservaciones/{rid}/devoluciones", json={"importe": 300, "motivo": "cobro duplicado"})
    assert ok.status_code == 201 and ok.json()["por_devolver"] == 0


async def test_confirmacion_al_cliente_queda_en_la_auditoria(api):
    c, _ = api
    r, _ = await _reservar(c, "Confirmar", ["adulto"])
    rid = r.json()["id"]
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/confirmacion", json={"canal": "whatsapp"})).status_code == 200
    assert (await c.post(f"{PREFIJO}/reservaciones/{rid}/confirmacion", json={"canal": "paloma"})).status_code == 422
    filas = (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": rid})).json()["filas"]
    conf = [f for f in filas if f["accion"] == "confirmacion"]
    assert conf and conf[0]["despues"] == {"canal": "whatsapp"} and conf[0]["etiqueta"].startswith("Confirmación")


async def test_mensaje_de_confirmacion_se_configura_en_el_evento(api):
    c, _ = api
    assert (await c.get(f"{PREFIJO}/evento")).json()["mensaje_confirmacion"] == ""
    r = await c.put(f"{PREFIJO}/evento", json={"mensaje_confirmacion": "Depósito a la cuenta 1234. Gracias."})
    assert r.status_code == 200 and r.json()["mensaje_confirmacion"] == "Depósito a la cuenta 1234. Gracias."


async def test_historico_de_pagos_de_ninos_2025_por_dia(api):
    c, _ = api
    r = await c.get(f"{PREFIJO}/historico/pagos", params={"anio": 2025})
    assert r.status_code == 200
    d = r.json()
    por_fecha = {x["fecha"]: x for x in d["dias"]}
    assert d["tipos"]["nino"]["boletos"] > 0 and d["tipos"]["nino"]["con_dinero"] is True
    # adultos 2025: solo la cantidad por día (sin dinero)
    assert d["tipos"]["adulto"]["boletos"] == 547 and d["tipos"]["adulto"]["con_dinero"] is False and d["tipos"]["adulto"]["cobrado"] == 0
    assert por_fecha["2025-12-10"]["adulto"]["reservados"] == 123 and por_fecha["2025-12-09"]["adulto"]["reservados"] == 90
    # 05/11: un anticipo de 272.50; el 10/12 entran las 3 liquidaciones de 817.50 además de los de contado
    assert por_fecha["2025-11-05"]["nino"]["cobrado"] == 272.50 and por_fecha["2025-11-05"]["nino"]["anticipos"] == 1
    assert por_fecha["2025-12-10"]["nino"]["liquidaciones"] == 3
    assert por_fecha["2025-12-10"]["nino"]["cobrado"] == round(por_fecha["2025-12-10"]["nino"]["completos"] * 1090 + 3 * 817.50, 2)
    # año sin datos: respuesta vacía, no error
    vacio = (await c.get(f"{PREFIJO}/historico/pagos", params={"anio": 2019})).json()
    assert vacio["dias"] == [] and vacio["tipos"]["nino"]["disponible"] is False


async def test_eliminar_canceladas_solo_sin_dinero_y_queda_en_la_auditoria(api):
    c, _ = api
    # activa: no se puede eliminar
    activa, _ = await _reservar(c, "Activa", ["adulto"])
    r0 = await c.delete(f"{PREFIJO}/reservaciones/{activa.json()['id']}/definitivo")
    assert r0.status_code == 409 and "canceladas" in r0.json()["detail"]
    # cancelada CON pago vigente: no se puede (el dinero seguiría en Cobros)
    con_pago, _ = await _reservar(c, "Con pago", ["adulto"])
    cid = con_pago.json()["id"]
    await c.post(f"{PREFIJO}/reservaciones/{cid}/pagos", json={"importe": 500, "forma_pago": "Efectivo", "fecha": HOY})
    await c.delete(f"{PREFIJO}/reservaciones/{cid}", params={"motivo": "prueba"})
    r1 = await c.delete(f"{PREFIJO}/reservaciones/{cid}/definitivo")
    assert r1.status_code == 409 and "pago" in r1.json()["detail"].lower()
    # cancelada sin dinero: se elimina y desaparece de la lista
    sin_dinero, _ = await _reservar(c, "Sin dinero", ["adulto"])
    sid = sin_dinero.json()["id"]
    await c.delete(f"{PREFIJO}/reservaciones/{sid}", params={"motivo": "ya no"})
    assert (await c.delete(f"{PREFIJO}/reservaciones/{sid}/definitivo")).status_code == 200
    assert (await c.get(f"{PREFIJO}/reservaciones/{sid}")).status_code == 404
    ids = {x["id"] for x in (await c.get(f"{PREFIJO}/reservaciones", params={"incluir_canceladas": True})).json()["reservaciones"]}
    assert sid not in ids and cid in ids
    aud = (await c.get(f"{PREFIJO}/auditoria", params={"reservacion_id": sid})).json()
    assert any(f["accion"] == "eliminar" and f["antes"]["cliente"] == "Sin dinero" for f in aud["filas"])
    # en lote: borra las que no tienen dinero y avisa de las que sí
    otra, _ = await _reservar(c, "Otra cancelada", ["nino"])
    await c.delete(f"{PREFIJO}/reservaciones/{otra.json()['id']}", params={"motivo": "x"})
    lote = (await c.post(f"{PREFIJO}/reservaciones/eliminar-canceladas")).json()
    assert lote["eliminadas"] == 1 and [x["id"] for x in lote["con_dinero"]] == [cid]
    # las sillas de lo eliminado quedan libres y la activa no se tocó
    assert (await c.get(f"{PREFIJO}/reservaciones/{activa.json()['id']}")).status_code == 200


async def test_exportar_excel_completo_trae_todas_las_hojas_con_formulas(api):
    import io
    from openpyxl import load_workbook

    c, _ = api
    r, _ = await _reservar(c, "Cliente Excel", ["adulto", "nino"])
    rid = r.json()["id"]
    await c.post(f"{PREFIJO}/reservaciones/{rid}/pagos", json={"importe": 800, "forma_pago": "Efectivo", "fecha": HOY})
    otra, _ = await _reservar(c, "Cancelada Excel", ["adulto"])
    await c.delete(f"{PREFIJO}/reservaciones/{otra.json()['id']}", params={"motivo": "prueba"})
    resp = await c.get(f"{PREFIJO}/exportar-excel")
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    assert "respaldo" in resp.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(resp.content))
    assert libro.sheetnames == ["LEEME", "Resumen", "Reservaciones", "Sillas", "Pagos", "Cobros por día", "Devoluciones", "Comparativo anterior", "Configuración"]
    # cada hoja explica qué es (fila 2) y tiene encabezados (fila 4)
    for nombre in libro.sheetnames[1:]:
        assert libro[nombre]["A1"].value and libro[nombre]["A2"].value and libro[nombre]["A4"].value, nombre
    res = libro["Reservaciones"]
    nombres = [res.cell(row=f, column=2).value for f in range(5, res.max_row + 1)]
    assert "Cliente Excel" in nombres and "Cancelada Excel" in nombres
    fila = 5 + nombres.index("Cliente Excel")
    assert res.cell(row=fila, column=14).value == f"=L{fila}-M{fila}"  # saldo es fórmula viva
    assert res.cell(row=fila, column=13).value == 800
    assert str(res.cell(row=res.max_row, column=12).value).startswith("=SUM(")
    sillas = libro["Sillas"]
    assert sillas.max_row > 700  # todas las sillas del salón, también las libres
    pagos = libro["Pagos"]
    assert pagos.cell(row=5, column=5).value == 800 and pagos.cell(row=5, column=9).value == "Vigente"
    # queda en la auditoría
    aud = (await c.get(f"{PREFIJO}/auditoria", params={"accion": "exportar"})).json()
    assert any(f["accion"] == "exportar" for f in aud["filas"])
